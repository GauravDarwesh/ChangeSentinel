"""Behavioural tests for dashboard-lib.js, run through Node.

The redline, scan-cycle and LED-field logic decide what a reviewer sees, so
they are exercised directly rather than only pinned as source strings.
"""
import json
import shutil
import subprocess
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
NODE = shutil.which("node")


def run_js(body: str):
    """Evaluate `body` with the library bound to `L`; it must print one JSON value."""
    script = "const L = require(" + json.dumps(str(ROOT / "dashboard-lib.js")) + ");\n" + body
    result = subprocess.run([NODE, "-e", script], capture_output=True, text=True, timeout=30, check=True)
    return json.loads(result.stdout)


@unittest.skipUnless(NODE, "node is required for dashboard library tests")
class TestRedline(unittest.TestCase):
    DIFF = "\n".join([
        "--- before",
        "+++ after",
        "@@ -10,3 +10,3 @@",
        " Submission",
        "-Reports are due 30 calendar days after the as-of date.",
        "+Reports are due 15 calendar days after the as-of date.",
        " Questions",
        "@@ -40,2 +40,3 @@",
        " a",
        "+New <b>line</b>",
        " b",
    ])

    def test_a_one_word_edit_is_marked_inside_the_sentence(self):
        out = run_js("console.log(JSON.stringify(L.renderRedline(" + json.dumps(self.DIFF) + ")));")
        self.assertIn("Reports are due <del>30</del><ins>15</ins> calendar days", out["html"])
        self.assertEqual(out["stats"], {"added": 1, "removed": 0, "edited": 1})
        self.assertEqual(out["summary"], "1 line edited, 1 line added")

    def test_scraped_markup_is_escaped(self):
        out = run_js("console.log(JSON.stringify(L.renderRedline(" + json.dumps(self.DIFF) + ").html));")
        self.assertIn("New &lt;b&gt;line&lt;/b&gt;", out)
        self.assertNotIn("<b>", out)

    def test_unchanged_stretches_are_collapsed_between_hunks(self):
        out = run_js("console.log(JSON.stringify(L.renderRedline(" + json.dumps(self.DIFF) + ").html));")
        self.assertIn("9 unchanged lines before this change", out)
        self.assertIn("27 unchanged lines", out)

    def test_removed_lines_carry_no_current_line_number(self):
        diff = "--- before\n+++ after\n@@ -5,3 +5,1 @@\n keep\n-gone one\n-gone two"
        rows = run_js("console.log(JSON.stringify(L.renderRedline(" + json.dumps(diff) + ").html));")
        self.assertIn('<span class="rl-no" aria-hidden="true"></span><span class="rl-text"><del aria-label="Removed line">gone one</del>', rows)
        self.assertIn('<span class="rl-no" aria-hidden="true">5</span><span class="rl-text">keep', rows)

    def test_a_rewritten_line_is_shown_as_removed_and_added(self):
        diff = "--- before\n+++ after\n@@ -1,1 +1,1 @@\n-Alpha beta gamma delta\n+Completely different wording here"
        out = run_js("console.log(JSON.stringify(L.renderRedline(" + json.dumps(diff) + ").stats));")
        self.assertEqual(out, {"added": 1, "removed": 1, "edited": 0})

    def test_a_truncated_diff_says_so(self):
        diff = self.DIFF + "\n... diff truncated after 250 lines"
        out = run_js("console.log(JSON.stringify(L.renderRedline(" + json.dumps(diff) + ")));")
        self.assertTrue(out["truncated"])
        self.assertIn("stops after 250 lines", out["html"])

    def test_an_empty_diff_explains_itself(self):
        out = run_js('console.log(JSON.stringify(L.renderRedline("")));')
        self.assertIn("No line-level text differences", out["html"])

    def test_the_real_fr_3033p_evidence_renders_as_four_removed_lines(self):
        evidence = ROOT / "data/evidence/1c4171c28c233d33130ce0a35bc8a04167020afadf1f11385d118c6b80e748e1.json"
        if not evidence.exists():
            self.skipTest("historical evidence file not present")
        diff = json.loads(evidence.read_text(encoding="utf-8"))["diff"]
        out = run_js("console.log(JSON.stringify(L.renderRedline(" + json.dumps(diff) + ").summary));")
        self.assertEqual(out, "4 lines removed")


@unittest.skipUnless(NODE, "node is required for dashboard library tests")
class TestScanSchedule(unittest.TestCase):
    def test_next_scan_follows_the_utc_six_hour_schedule(self):
        out = run_js('const c = L.scanCycle(new Date("2026-10-20T07:30:00Z")); console.log(JSON.stringify({next: c.next.toISOString(), progress: c.progress, remaining: c.remainingMinutes}));')
        self.assertEqual(out, {"next": "2026-10-20T12:00:00.000Z", "progress": 0.25, "remaining": 270})

    def test_exactly_on_a_slot_the_next_scan_is_six_hours_away(self):
        out = run_js('console.log(JSON.stringify(L.nextScheduledScan(new Date("2026-10-20T06:00:00Z")).toISOString()));')
        self.assertEqual(out, "2026-10-20T12:00:00.000Z")

    def test_durations_are_written_for_people(self):
        out = run_js("console.log(JSON.stringify([L.formatDuration(45), L.formatDuration(120), L.formatDuration(226)]));")
        self.assertEqual(out, ["45 min", "2 h", "3 h 46 min"])


@unittest.skipUnless(NODE, "node is required for dashboard library tests")
class TestScanColumns(unittest.TestCase):
    HISTORY = [
        {"run_id": "a", "generated_at": "2026-10-20T04:32:57Z", "source_id": "s1"},
        {"run_id": "a", "generated_at": "2026-10-20T04:33:13Z", "source_id": "s2"},
        {"run_id": "m", "generated_at": "2026-10-20T14:50:52Z", "source_id": "s1", "trigger": "manual"},
        {"run_id": "b", "generated_at": "2026-10-20T21:31:03Z", "source_id": "s1", "trigger": "scheduled"},
        {"run_id": "b", "generated_at": "2026-10-20T21:31:02Z", "source_id": "s2", "trigger": "scheduled"},
    ]

    def columns(self, history=None, runs=None, limit=0):
        return run_js("console.log(JSON.stringify(L.runColumns(" + json.dumps(history or self.HISTORY) + ", {workflowRuns: " +
                      json.dumps(runs or []) + ", limit: " + str(limit) + "})));")

    def test_every_saved_scan_is_one_column_including_manual_runs(self):
        cols = self.columns()
        self.assertEqual([c["id"] for c in cols], ["a", "m", "b"])
        self.assertEqual([c["trigger"] for c in cols], ["unknown", "manual", "scheduled"])
        self.assertEqual(sorted(cols[0]["sources"]), ["s1", "s2"])
        self.assertEqual(sorted(cols[1]["sources"]), ["s1"])
        # Each column owns the events saved since the previous scan.
        self.assertEqual(cols[2]["prevAt"], cols[1]["at"])

    def test_older_history_is_labelled_from_the_actions_api(self):
        runs = [{"id": 7, "event": "workflow_dispatch", "status": "completed", "conclusion": "success",
                 "created_at": "2026-10-20T04:31:27Z", "run_started_at": "2026-10-20T04:31:30Z", "updated_at": "2026-10-20T04:35:00Z",
                 "actor": "GauravDarwesh"}]
        cols = self.columns(runs=runs)
        self.assertEqual((cols[0]["trigger"], cols[0]["workflowRunId"]), ("manual", "7"))

    def test_continuation_dispatches_are_not_counted_as_manual(self):
        runs = [{"id": 8, "event": "workflow_dispatch", "actor": "github-actions[bot]", "created_at": "2026-10-20T04:31:27Z",
                 "updated_at": "2026-10-20T04:35:00Z", "status": "completed"}]
        self.assertEqual(self.columns(runs=runs)[0]["trigger"], "continuation")

    def test_the_limit_keeps_the_newest_columns(self):
        self.assertEqual([c["id"] for c in self.columns(limit=2)], ["m", "b"])


@unittest.skipUnless(NODE, "node is required for dashboard library tests")
class TestChangeFieldCells(unittest.TestCase):
    COLUMN = {"id": "b", "at": 1792531863000, "prevAt": 1792507852000, "sources": {"s1": 1792531863000}}  # 2026-10-20 21:31:03Z, prev 14:50:52Z

    def state(self, page: dict) -> str:
        page = dict({"sourceId": "s1"}, **page)
        return run_js("console.log(JSON.stringify(L.columnCellState(" + json.dumps(self.COLUMN) + ", " + json.dumps(page) + ")));")

    def test_a_scan_that_found_a_change_lights_up(self):
        self.assertEqual(self.state({"events": [{"event_type": "CHANGED_URL", "timestamp": "2026-10-20T21:31:03Z"}]}), "on")

    def test_problems_outrank_changes(self):
        page = {"events": [{"event_type": "CHANGED_URL", "timestamp": "2026-10-20T21:31:03Z"},
                           {"event_type": "FETCH_ERROR", "timestamp": "2026-10-20T21:31:03Z"}]}
        self.assertEqual(self.state(page), "bad")

    def test_a_current_error_marks_the_scan_of_the_last_check(self):
        self.assertEqual(self.state({"problemAt": "2026-10-20T21:31:03Z"}), "bad")

    def test_a_new_page_has_its_own_state(self):
        self.assertEqual(self.state({"events": [{"event_type": "NEW_URL", "timestamp": "2026-10-20T21:31:03Z"}]}), "new")
    def test_first_successful_scan_of_new_source_shows_new_even_without_baseline_event(self):
        self.assertEqual(self.state({"firstSeen": "2026-10-20T20:00:00Z", "events": []}), "new")

    def test_new_monitor_does_not_paint_scans_from_before_it_was_added(self):
        self.assertEqual(self.state({"firstSeen": "2026-10-20T21:32:00Z", "events": []}), "none")

    def test_a_scan_with_no_change_is_dim(self):
        self.assertEqual(self.state({}), "dim")

    def test_an_earlier_scans_event_belongs_to_that_scan(self):
        self.assertEqual(self.state({"events": [{"event_type": "CHANGED_URL", "timestamp": "2026-10-20T14:50:52Z"}]}), "dim")

    def test_other_sources_and_pages_not_yet_monitored_are_not_drawn(self):
        self.assertEqual(self.state({"sourceId": "s2"}), "none")
        self.assertEqual(self.state({"firstSeen": "2026-10-21T00:00:00Z"}), "none")


@unittest.skipUnless(NODE, "node is required for dashboard library tests")
class TestSlotColumns(unittest.TestCase):
    # Real shape of 21-22 Oct: a one-source continuation and manual runs between
    # scheduled scans used to add a column each, leaving holes in every other row.
    HISTORY = [
        {"run_id": "c", "generated_at": "2026-10-21T03:47:47Z", "source_id": "bcar", "trigger": "continuation"},
        {"run_id": "s1", "generated_at": "2026-10-21T04:25:45Z", "source_id": "fed", "trigger": "scheduled"},
        {"run_id": "s1", "generated_at": "2026-10-21T04:29:41Z", "source_id": "bcar", "trigger": "scheduled"},
        {"run_id": "m", "generated_at": "2026-10-21T06:54:49Z", "source_id": "fed", "trigger": "manual"},
        {"run_id": "s2", "generated_at": "2026-10-21T11:37:32Z", "source_id": "fed", "trigger": "scheduled"},
        {"run_id": "s2", "generated_at": "2026-10-21T11:41:03Z", "source_id": "bcar", "trigger": "scheduled"},
    ]

    def columns(self, limit=0):
        return run_js("console.log(JSON.stringify(L.slotColumns(" + json.dumps(self.HISTORY) + ", {limit: " + str(limit) + "})));")

    def cell(self, column, page):
        return run_js("console.log(JSON.stringify(L.columnCellState(" + json.dumps(column) + ", " + json.dumps(page) + ")));")

    def test_off_schedule_runs_merge_into_their_six_hour_slot(self):
        cols = self.columns()
        self.assertEqual(len(cols), 2)
        self.assertEqual(sorted(cols[0]["sources"]), ["bcar", "fed"])
        self.assertEqual(sorted(cols[1]["sources"]), ["bcar", "fed"])
        self.assertEqual([c["trigger"] for c in cols], ["scheduled", "scheduled"])
        self.assertEqual([c["manual"] for c in cols], [0, 1])

    def test_a_page_added_mid_window_lights_green_once(self):
        cols = self.columns()
        page = {"sourceId": "bcar", "firstSeen": "2026-10-21T03:47:47Z", "events": []}
        self.assertEqual([self.cell(c, page) for c in cols], ["new", "dim"])

    def test_a_page_whose_first_crawl_is_still_running_draws_nothing(self):
        cols = self.columns()
        page = {"sourceId": "bcar", "pending": True, "events": []}
        self.assertEqual([self.cell(c, page) for c in cols], ["none", "none"])

    def test_markup_only_changes_do_not_light_as_content_changes(self):
        cols = self.columns()
        page = {"sourceId": "bcar", "firstSeen": "2026-10-20T00:00:00Z",
                "events": [{"event_type": "HTML_SOURCE_CHANGED", "timestamp": "2026-10-21T11:41:03Z"}]}
        self.assertEqual([self.cell(c, page) for c in cols], ["dim", "markup"])

    def test_a_content_change_in_a_manual_run_lights_its_slot(self):
        cols = self.columns()
        page = {"sourceId": "fed", "firstSeen": "2026-10-18T00:00:00Z",
                "events": [{"event_type": "CHANGED_URL", "timestamp": "2026-10-21T06:54:49Z"}]}
        self.assertEqual([self.cell(c, page) for c in cols], ["dim", "on"])


@unittest.skipUnless(NODE, "node is required for dashboard library tests")
class TestPendingScan(unittest.TestCase):
    def pending(self, now, columns, runs=None):
        return run_js("const p = L.pendingScan(new Date(" + json.dumps(now) + "), " + json.dumps(columns) + ", " + json.dumps(runs or []) +
                      "); console.log(JSON.stringify({due: p.due.toISOString(), phase: p.phase, late: p.lateMinutes, trigger: p.trigger}));")

    SCHEDULED_0433 = {"id": "a", "at": 1792470793000, "trigger": "scheduled"}  # 2026-10-20T04:33:13Z
    MANUAL_0300 = {"id": "m", "at": 1792465200000, "trigger": "manual"}       # 2026-10-20T03:00:00Z

    def test_after_the_slots_scheduled_scan_the_clock_moves_to_the_next_slot(self):
        out = self.pending("2026-10-20T05:00:00Z", [self.SCHEDULED_0433])
        self.assertEqual((out["due"], out["phase"]), ("2026-10-20T06:00:00.000Z", "waiting"))

    def test_a_manual_run_does_not_stand_in_for_the_scheduled_scan(self):
        out = self.pending("2026-10-20T03:30:00Z", [self.MANUAL_0300])
        self.assertEqual((out["due"], out["phase"], out["late"]), ("2026-10-20T00:00:00.000Z", "late", 210))

    def test_just_after_a_slot_starts_the_scan_is_due(self):
        self.assertEqual(self.pending("2026-10-20T06:05:00Z", [self.SCHEDULED_0433])["phase"], "due")

    def test_a_run_in_progress_on_github_shows_as_running(self):
        runs = [{"id": 9, "event": "workflow_dispatch", "actor": "GauravDarwesh", "status": "in_progress",
                 "created_at": "2026-10-20T05:10:00Z", "run_started_at": "2026-10-20T05:10:05Z", "updated_at": "2026-10-20T05:11:00Z"}]
        out = self.pending("2026-10-20T05:12:00Z", [self.SCHEDULED_0433], runs)
        self.assertEqual((out["phase"], out["trigger"]), ("running", "manual"))

    def test_a_finished_run_whose_scan_is_saved_is_not_still_publishing(self):
        # The 8:06 am manual run: results saved at 02:36:32Z, GitHub marked the run done at 02:40:37Z.
        saved = {"id": "x", "at": 1792550192000, "trigger": "manual", "workflowRunId": "39533795876"}
        runs = [{"id": 39533795876, "event": "workflow_dispatch", "status": "completed", "conclusion": "success",
                 "created_at": "2026-10-21T02:34:40Z", "run_started_at": "2026-10-21T02:34:45Z", "updated_at": "2026-10-21T02:40:37Z"}]
        self.assertEqual(self.pending("2026-10-21T02:45:00Z", [saved], runs)["phase"], "late")
        saved.pop("workflowRunId")
        self.assertEqual(self.pending("2026-10-21T02:45:00Z", [saved], runs)["phase"], "late")

    def test_a_finished_run_not_yet_saved_is_publishing(self):
        runs = [{"id": 9, "event": "schedule", "status": "completed", "conclusion": "success",
                 "created_at": "2026-10-20T06:40:00Z", "updated_at": "2026-10-20T06:45:00Z"}]
        self.assertEqual(self.pending("2026-10-20T06:48:00Z", [self.SCHEDULED_0433], runs)["phase"], "publishing")


@unittest.skipUnless(NODE, "node is required for dashboard library tests")
class TestEventWording(unittest.TestCase):
    def test_every_engine_event_type_has_a_plain_label_and_tone(self):
        types = ["CHANGED_URL", "NEW_URL", "REMOVED_URL", "HTML_SOURCE_CHANGED", "FETCH_ERROR", "EXTRACTION_ERROR", "BASELINE_MIGRATION"]
        out = run_js("console.log(JSON.stringify(" + json.dumps(types) + ".map(t => [L.eventLabel(t), L.eventTone(t)])));")
        labels = [label for label, _tone in out]
        self.assertNotIn("Monitoring event", labels)
        self.assertEqual(dict(zip(types, [tone for _label, tone in out])), {
            "CHANGED_URL": "change", "NEW_URL": "change", "REMOVED_URL": "change", "HTML_SOURCE_CHANGED": "markup",
            "FETCH_ERROR": "problem", "EXTRACTION_ERROR": "problem", "BASELINE_MIGRATION": "info",
        })

    def test_engine_event_types_are_all_covered(self):
        engine = (ROOT / "regmon/engine.py").read_text(encoding="utf-8")
        lib = (ROOT / "dashboard-lib.js").read_text(encoding="utf-8")
        for event_type in ("CHANGED_URL", "NEW_URL", "REMOVED_URL", "HTML_SOURCE_CHANGED", "FETCH_ERROR", "EXTRACTION_ERROR", "BASELINE_MIGRATION"):
            with self.subTest(event_type=event_type):
                self.assertIn(event_type, engine)
                self.assertIn(event_type + ":", lib)

    def test_saved_federal_reserve_evidence_includes_both_deleted_rows(self):
        from pathlib import Path
        evidence = ROOT / "data/evidence/c89216e3b00e847563d351357b43772920c29a2266cccb68a8a651736f928219.json"
        if not evidence.exists():
            self.skipTest("current Federal Reserve evidence not present in this checkout")
        record = json.loads(evidence.read_text(encoding="utf-8"))
        self.assertEqual(record["url"], "https://www.federalreserve.gov/apps/reportingforms/review")
        out = run_js("console.log(JSON.stringify(L.renderRedline(" + json.dumps(record["diff"]) + ")));")
        self.assertEqual(out["stats"], {"added": 0, "removed": 2, "edited": 0})
        self.assertIn("FFIEC 031", out["html"])
        self.assertIn("FR Y-9C R-1867", out["html"])

    def test_table_row_matching_requires_the_full_cell_sequence(self):
        out = run_js(r"""
          const expected = L.splitMarkdownTableRow("| FFIEC 031 | link | FFIEC 031 | FFIEC 031 | N/A | link | link |");
          const sameLabelsWrongRow = ["FFIEC 031", "link", "FFIEC 031", "FFIEC 031", "N/A", "link", "Not Available"];
          const differentCollection = ["FR 3033p FR 3033s", "link", "FR 3033p/s", "N/A", "N/A", "link", "link"];
          console.log(JSON.stringify([
            L.tableCellsMatch(expected, expected),
            L.tableCellsMatch(expected, sameLabelsWrongRow),
            L.tableCellsMatch(expected, differentCollection),
            L.tableCellsMatch(expected, expected.slice(0, -1))
          ]));
        """)
        self.assertEqual(out, [True, False, False, False])

    def test_markup_only_events_are_visually_distinct_from_content_changes(self):
        out = run_js("""
          console.log(JSON.stringify([
            L.eventTone("CHANGED_URL"), L.eventTone("NEW_URL"),
            L.eventTone("REMOVED_URL"), L.eventTone("HTML_SOURCE_CHANGED"),
            L.eventTone("FETCH_ERROR")
          ]));
        """)
        self.assertEqual(out, ["change", "change", "change", "markup", "problem"])

    def test_federal_reserve_recently_approved_table_redline_keeps_both_deleted_rows(self):
        # Regression for a repeated form label near the top of the table:
        # it must not disappear when another distinct row is deleted near the bottom.
        diff = "\n".join([
            "--- before", "+++ after",
            "@@ -68,8 +68,6 @@",
            "| Recently Approved and Implemented | | | | | | |",
            "| --- | --- | --- | --- | --- | --- | --- |",
            "| Collection Identifier | OMB Supporting Statement | Reporting Form | Reporting Instructions | Other Attachments | Initial Federal Register Notice | Final Federal Register Notice |",
            "-| FFIEC 031 | link | FFIEC 031 | FFIEC 031 | N/A | link | link |",
            "| FR 3033p FR 3033s | link | FR 3033p/s | N/A | N/A | link | link |",
            "| FR 4031 | link | N/A | N/A | N/A | link | link |",
            "-| FR Y-9C R-1867 | link | N/A | FR Y-9C | Press Release 1 Press Release 2 | link | link |",
            "Last Update: Oct 16, 2026"
        ])
        out = run_js("console.log(JSON.stringify(L.renderRedline(" + json.dumps(diff) + ")));")
        self.assertEqual(out["stats"], {"added": 0, "removed": 2, "edited": 0})
        self.assertIn("FFIEC 031", out["html"])
        self.assertIn("FR Y-9C R-1867", out["html"])


if __name__ == "__main__":
    unittest.main()


@unittest.skipUnless(NODE, "node is required for dashboard library tests")
class TestCompareDiff(unittest.TestCase):
    def test_diff_runs_reconstruct_both_sequences_and_are_minimal(self):
        out = run_js(r"""
          let seed = 7; const rnd = () => (seed = (seed * 1103515245 + 12345) % 2147483648) / 2147483648;
          let ok = true, optimal = true;
          for (let t = 0; t < 400; t++) {
            const a = Array.from({length: Math.floor(rnd() * 30)}, () => "abcd"[Math.floor(rnd() * 4)]);
            const b = a.filter(() => rnd() > 0.25).map(x => rnd() < 0.1 ? "z" : x);
            const runs = L.diffSequences(a, b);
            let ia = 0, ib = 0, same = 0;
            for (const r of runs) {
              if (r.a[0] !== ia || r.b[0] !== ib) ok = false;
              if (r.type === "same") { same += r.a[1] - r.a[0]; for (let k = 0; k < r.a[1] - r.a[0]; k++) if (a[r.a[0] + k] !== b[r.b[0] + k]) ok = false; }
              ia = r.a[1]; ib = r.b[1];
            }
            if (ia !== a.length || ib !== b.length) ok = false;
            const dp = Array.from({length: a.length + 1}, () => new Array(b.length + 1).fill(0));
            for (let i = a.length - 1; i >= 0; i--) for (let j = b.length - 1; j >= 0; j--) dp[i][j] = a[i] === b[j] ? dp[i + 1][j + 1] + 1 : Math.max(dp[i + 1][j], dp[i][j + 1]);
            if (same !== dp[0][0]) optimal = false;
          }
          console.log(JSON.stringify({ok, optimal}));
        """)
        self.assertEqual(out, {"ok": True, "optimal": True})

    def test_large_documents_diff_quickly_through_patience_anchors(self):
        out = run_js(r"""
          const a = Array.from({length: 150000}, (_, i) => "line " + i);
          const b = a.slice(); b.splice(700, 1, "edited"); b.splice(120000, 0, "inserted");
          const t0 = Date.now(); const runs = L.diffSequences(a, b);
          console.log(JSON.stringify({ms: Date.now() - t0, changed: runs.filter(r => r.type !== "same").length}));
        """)
        self.assertLess(out["ms"], 5000)
        self.assertEqual(out["changed"], 3)

    def test_aligned_rows_pair_edits_word_by_word_and_keep_insertions_apart(self):
        out = run_js(r"""
          const r = L.alignLines("Title\nReports are due in 30 days.\nUnchanged\nOld only line here",
                                 "Title\nReports are due in 15 days.\nUnchanged\nA completely new sentence");
          console.log(JSON.stringify({types: r.rows.map(x => x.type), stats: r.stats,
            del: r.rows[1].segments.filter(s => s.type === "del").map(s => s.text.trim()),
            add: r.rows[1].segments.filter(s => s.type === "add").map(s => s.text.trim())}));
        """)
        self.assertEqual(out["types"], ["same", "edit", "same", "del", "add"])
        self.assertEqual(out["del"], ["30"])
        self.assertEqual(out["add"], ["15"])

    def test_changed_words_are_reported_per_side(self):
        out = run_js(r"""
          const w = s => s.split(" ").map(text => ({text}));
          console.log(JSON.stringify(L.changedWords(w("rate of 8.500 percent"), w("rate of 8.750 percent"))));
        """)
        self.assertEqual(out, {"removed": [2], "added": [2]})


@unittest.skipUnless(NODE, "node is required for dashboard library tests")
class TestScoutHelpers(unittest.TestCase):
    CANDIDATES = [
        {"key": "a:1", "origin": "new", "url": "https://r.org/a", "first_seen": "2026-10-19T04:20:00Z", "assessment": {"outcome": "relevant", "rank": 36}},
        {"key": "a:2", "origin": "new", "url": "https://r.org/b", "first_seen": "2026-10-19T04:20:00Z", "assessment": {"outcome": "uncertain", "rank": 21}},
        {"key": "a:3", "origin": "new", "url": "https://r.org/c", "first_seen": "2026-10-12T04:20:00Z", "assessment": {"outcome": "irrelevant", "rank": 12}},
        {"key": "a:4", "origin": "baseline", "url": "https://r.org/d", "first_seen": "2026-10-12T04:20:00Z", "assessment": {"outcome": "relevant"}},
        {"key": "a:5", "origin": "baseline", "url": "https://r.org/e", "first_seen": "2026-10-12T04:20:00Z", "assessment": {"outcome": "uncertain"}},
        {"key": "a:6", "origin": "new", "url": "https://r.org/f", "first_seen": "2026-10-19T04:20:00Z", "assessment": {"outcome": "relevant"},
         "decision": {"decision": "dismissed", "decided_at": "2026-10-20T10:00:00Z"}},
        {"key": "a:7", "origin": "new", "url": "http://r.org/g/", "first_seen": "2026-10-19T04:20:00Z", "assessment": {"outcome": "relevant"}},
    ]

    def test_findings_land_in_one_list_each(self):
        out = run_js("const b = L.scoutBuckets(" + json.dumps(self.CANDIDATES) + ", new Set([L.urlKey('https://r.org/g')]));" +
                     "console.log(JSON.stringify(Object.fromEntries(Object.entries(b).map(([k, v]) => [k, v.map(x => x.key)]))));")
        self.assertEqual(out, {"REVIEW": ["a:1", "a:2"], "BACKLOG": ["a:4"], "EXCLUDED": ["a:3"], "DONE": ["a:6", "a:7"]})

    def test_sentences_never_overstate_the_evidence(self):
        out = run_js("""console.log(JSON.stringify([
          L.scoutSentence({outcome: "relevant", publication_type: {id: "guidance"}, topics: ["Liquidity"]}),
          L.scoutSentence({outcome: "uncertain", confidence: "weak"}),
          L.scoutSentence({outcome: "irrelevant", reasons: [{text: "Address looks like a site utility page"}]}),
          L.scoutSentence({outcome: "uncertain", confidence: "conflicting", publication_type: {id: "research"}})
        ]));""")
        self.assertEqual(out[0], "Looks like guidance or a guideline about liquidity.")
        self.assertIn("couldn't tell", out[1])
        self.assertIn("utility page", out[2])
        self.assertIn("both ways", out[3])

    def test_next_sweep_is_the_coming_monday_0417_utc(self):
        out = run_js("console.log(JSON.stringify([L.nextScoutSweep(new Date('2026-10-21T05:00:00Z')).toISOString(), L.nextScoutSweep(new Date('2026-10-26T04:00:00Z')).toISOString()]));")
        self.assertEqual(out, ["2026-10-26T04:17:00.000Z", "2026-10-26T04:17:00.000Z"])


    def test_backlog_is_grouped_by_regulator_then_strongest_first(self):
        candidates = [
            {"key": "x:1", "regulator_id": "pra", "origin": "baseline", "url": "https://p.org/1", "title": "B", "assessment": {"outcome": "relevant", "rank": 10}},
            {"key": "y:1", "regulator_id": "eba", "origin": "baseline", "url": "https://e.org/1", "title": "A", "assessment": {"outcome": "relevant", "rank": 5}},
            {"key": "x:2", "regulator_id": "pra", "origin": "baseline", "url": "https://p.org/2", "title": "C", "assessment": {"outcome": "relevant", "rank": 30}},
            {"key": "y:2", "regulator_id": "eba", "origin": "baseline", "url": "https://e.org/2", "title": "D", "assessment": {"outcome": "relevant", "rank": 5},
             "evidence": {"published": "2026-10-01"}},
        ]
        out = run_js("console.log(JSON.stringify(L.scoutBuckets(" + json.dumps(candidates) + ", new Set()).BACKLOG.map(x => x.key)));")
        self.assertEqual(out, ["y:2", "y:1", "x:2", "x:1"])


@unittest.skipUnless(NODE, "node is required for dashboard library tests")
class TestWordDiffAndUrls(unittest.TestCase):
    def test_a_replaced_phrase_reads_as_one_edit(self):
        out = run_js('console.log(JSON.stringify(L.wordDiff("the quick brown fox", "a slow red fox")));')
        self.assertEqual(out, [{"type": "del", "text": "the quick brown"}, {"type": "add", "text": "a slow red"},
                               {"type": "same", "text": " fox"}])

    def test_segments_always_rebuild_both_texts(self):
        out = run_js("""
          const words = "alpha beta gamma 10% 2026 $5 billion Tier 1 2".split(" ");
          let seed = 7; const rnd = () => (seed = (seed * 1103515245 + 12345) % 2147483648) / 2147483648;
          const make = () => Array.from({length: Math.floor(rnd() * 12)}, () => words[Math.floor(rnd() * words.length)]).join(rnd() < .3 ? "  " : " ");
          let bad = 0;
          for (let i = 0; i < 2000; i++) {
            const a = make(), b = make(), s = L.wordDiff(a, b);
            const before = s.filter(x => x.type !== "add").map(x => x.text).join("");
            const after = s.filter(x => x.type !== "del").map(x => x.text).join("");
            if (before !== a || after !== b) bad++;
          }
          console.log(JSON.stringify(bad));""")
        self.assertEqual(out, 0)

    def test_a_long_line_does_not_build_a_quadratic_table(self):
        out = run_js("""
          const a = Array.from({length: 1500}, (_, i) => "w" + i).join(" ");
          const b = a.replace("w700", "x700");
          const t = Date.now(); const s = L.wordDiff(a, b);
          console.log(JSON.stringify({ms: Date.now() - t, dels: s.filter(x => x.type === "del").map(x => x.text)}));""")
        self.assertEqual(out["dels"], ["w700"])
        self.assertLess(out["ms"], 2000)

    def test_one_canonical_form_mirrors_the_engine(self):
        out = run_js("""console.log(JSON.stringify([
          L.canonicalUrl("https://Example.org/a/b/?utm_source=x&b=2&a=1&gclid=q#top"),
          L.urlKey("http://example.org/a/b?a=1&b=2&fbclid=z"),
          L.urlKey("https://example.org/a/b/?b=2&a=1"),
          L.canonicalUrl("not a url")
        ]));""")
        self.assertEqual(out[0], "https://example.org/a/b?a=1&b=2")
        self.assertEqual(out[1], out[2])
        self.assertEqual(out[3], "not a url")

    def test_python_and_javascript_canonical_forms_agree(self):
        from regmon.discovery import canonical
        urls = ["https://www.federalreserve.gov/apps/reportingforms/Download/DownloadAttachment?guid=ab-12&type=pdf",
                "https://x.gov/a/?z=1&utm_medium=e&a=2#f", "https://x.gov/", "https://x.gov/path?q=a b"]
        out = run_js("console.log(JSON.stringify(" + json.dumps(urls) + ".map(L.canonicalUrl)));")
        self.assertEqual(out, [canonical(url) for url in urls])


@unittest.skipUnless(NODE, "node is required for dashboard library tests")
class TestInlineMarkdown(unittest.TestCase):
    def render(self, text):
        return run_js("console.log(JSON.stringify(L.inlineMarkdown(" + json.dumps(text) + ")));")

    def test_an_asterisk_inside_a_link_never_becomes_emphasis(self):
        html = self.render("[form](https://x.gov/a*b*c) and **bold**")
        self.assertIn('href="https://x.gov/a*b*c"', html)
        self.assertNotIn("<em>", html)
        self.assertIn("<strong>bold</strong>", html)

    def test_sentence_punctuation_stays_outside_a_bare_link(self):
        html = self.render("See https://x.gov/page. Then (https://y.gov/p).")
        self.assertIn('href="https://x.gov/page"', html)
        self.assertIn('href="https://y.gov/p"', html)
        self.assertTrue(html.endswith(").") )

    def test_markup_and_script_urls_are_inert(self):
        html = self.render('<img src=x onerror=alert(1)> [x](javascript:alert(1)) "quoted"')
        self.assertNotIn("<img", html)
        self.assertNotIn('href="javascript', html)
        self.assertIn("&quot;quoted&quot;", html)


@unittest.skipUnless(NODE, "node is required for dashboard library tests")
class TestLocalXlsxReader(unittest.TestCase):
    def test_first_sheet_cells_are_read_as_text(self):
        import tempfile
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["URL", "Name", "Include Selector", "Ignore Selectors"])
        sheet.append(["https://www.example.gov/a?x=1&y=2", "Policy <A> & B", "main .content", ".cookie;.nav"])
        sheet["A4"] = "https://www.example.gov/b"
        sheet["B4"] = 2026
        workbook.create_sheet("Second")["A1"] = "ignored"
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "links.xlsx"
            workbook.save(path)
            out = run_js("const fs = require('fs'); L.xlsxFirstSheetRows(fs.readFileSync(" + json.dumps(str(path)) +
                         ")).then(rows => console.log(JSON.stringify(rows)));")
        self.assertEqual(out[0], ["URL", "Name", "Include Selector", "Ignore Selectors"])
        self.assertEqual(out[1], ["https://www.example.gov/a?x=1&y=2", "Policy <A> & B", "main .content", ".cookie;.nav"])
        self.assertEqual(out[2], [])
        self.assertEqual(out[3], ["https://www.example.gov/b", "2026"])
        self.assertEqual(len(out), 4)

    def test_a_non_workbook_is_rejected_with_a_readable_error(self):
        out = run_js("L.xlsxFirstSheetRows(new Uint8Array([1,2,3,4])).then(() => console.log(JSON.stringify('ok')), e => console.log(JSON.stringify(e.message)));")
        self.assertIn("not a valid .xlsx", out)


@unittest.skipUnless(NODE, "node is required for dashboard library tests")
class TestEditOffsets(unittest.TestCase):
    def test_a_repeated_token_is_located_where_it_changed(self):
        out = run_js("""
          const before = "Tier 1 capital for Tier 1 banks", after = "Tier 1 capital for Tier 2 banks";
          const parts = L.changedPartOffsets(L.wordDiff(before, after), "add");
          const line = L.normalizeForMatch(after);
          console.log(JSON.stringify(parts.map(p => [p.text, p.offset, line.substr(p.offset, p.text.length)])));""")
        self.assertEqual(out, [["2", 24, "2"]])

    def test_offsets_survive_whitespace_and_table_pipes(self):
        out = run_js("""
          const before = "| Item  | 30 days |  BHCK 2170 |", after = "| Item  | 15 days |  BHCK 2170 |";
          const parts = L.changedPartOffsets(L.wordDiff(before, after), "add");
          const line = L.normalizeForMatch(after.replace(/\\|/g, " "));
          console.log(JSON.stringify(parts.map(p => line.substr(p.offset, p.text.length))));""")
        self.assertEqual(out, ["15"])


@unittest.skipUnless(NODE, "node is required for dashboard library tests")
class TestReviewLayer(unittest.TestCase):
    """The dashboard's half of item codes, lineage, dates and review requests
    must agree with the Python half that writes and parses them."""

    def test_item_code_wording_matches_the_alert_line(self):
        from regmon import item_codes
        before = "\n".join(["| RCON1000 | Cash | 1 |", "| RCON2000 | Loans | 2 |", "| RCON3000 | Total | 3 |", "| RCON4000 | Old item | 4 |"])
        after = "\n".join(["| RCON1000 | Cash and balances | 1 |", "| RCON2000 | Loans | 2 |", "| RCON3000 | Total | 3 |", "| RCON5000 | New item | 5 |", "| RCON6000 | Newer | 6 |"])
        result = item_codes.diff(before, after)
        summary = item_codes.summary(result)
        out = run_js("console.log(JSON.stringify([L.describeItemCodes(" + json.dumps(result) + "), L.describeItemCodes(" + json.dumps(summary) + "), L.describeItemCodes(null)]));")
        self.assertEqual(out, [item_codes.describe(result), item_codes.describe(summary), ""])
        self.assertTrue(out[0].startswith("2 added, 1 removed, 1 amended"))

    def test_lineage_lookup_matches_the_python_lookup(self):
        from regmon import lineage
        payload = {"schema_version": 1, "mappings": [
            {"series": ["BHCK"], "assets": ["dbt: fr_y9c.all_items", "warehouse.fr_y9c.total_assets"], "owner": "Finance"},
            {"codes": ["BHCK2170"], "assets": ["warehouse.fr_y9c.total_assets"], "owner": "Regulatory Reporting"},
            {"codes": ["RCON1234"], "assets": ["edit check 12"]},
        ]}
        mappings = lineage.validate(payload)
        for code in ("BHCK2170", "BHCK9999", "RCON1234", "QTFL0001"):
            with self.subTest(code=code):
                out = run_js("console.log(JSON.stringify(L.lineageFor(" + json.dumps(code) + ", " + json.dumps(payload) + ")));")
                self.assertEqual(out, lineage.assets_for(code, mappings))

    def test_calendar_days_are_never_shifted_by_the_viewers_time_zone(self):
        script = ("const L = require(" + json.dumps(str(ROOT / "dashboard-lib.js")) + ");"
                  "const d = L.isoDayToDate('2027-03-31');"
                  "console.log(JSON.stringify([d.getDate(), d.getMonth(), L.dueLabel('2027-03-31', '2027-03-30'), L.daysBetween('2026-12-31', '2027-03-31')]));")
        for zone in ("America/Los_Angeles", "Pacific/Kiritimati", "UTC"):
            with self.subTest(zone=zone):
                result = subprocess.run([NODE, "-e", script], capture_output=True, text=True, timeout=30, check=True, env={"TZ": zone, "PATH": "/usr/bin:/bin"})
                self.assertEqual(json.loads(result.stdout), [31, 2, "Tomorrow", 90])

    def test_upcoming_dates_are_soonest_first_and_skip_the_past(self):
        calendar = {"schema_version": 1, "entries": [
            {"date": "2026-10-20", "title": "Past", "event_id": "a"},
            {"date": "2027-01-15", "title": "Later", "event_id": "b"},
            {"date": "2026-11-02", "title": "Sooner", "event_id": "a"},
            {"date": "not a date", "title": "Broken", "event_id": "a"},
        ]}
        out = run_js("const c = " + json.dumps(calendar) + "; console.log(JSON.stringify([L.upcomingDates(c, '2026-10-21', 5).map(e => e.title), "
                     "L.datesForChange(c, 'a').map(e => e.title), L.dueLabel('2026-10-20', '2026-10-21'), L.dueLabel('2026-11-02', '2026-10-21'), L.upcomingDates(null, '2026-10-21')]));")
        self.assertEqual(out, [["Sooner", "Later"], ["Past", "Sooner"], "Yesterday", "In 12 days", []])

    def test_review_request_is_what_the_review_workflow_parses(self):
        from regmon import reviews
        event_id = "a" * 64
        out = run_js("console.log(JSON.stringify(L.changeReviewRequest({decision: 'action_required', rationale: '  New due\\n date  applies ', "
                     "owner: 'Reg Reporting', ticket: 'https://tracker.example.com/REG-1', changes: [{event_id: " + json.dumps(event_id) + ", url: 'https://www.fed.gov/a?b=1'}]})));")
        self.assertEqual(out["errors"], [])
        self.assertTrue(out["title"].startswith(reviews.TITLE_PREFIX))
        self.assertIn(reviews.MARKER, out["body"])
        parsed = reviews.parse_request(out["body"])
        self.assertEqual(parsed["decision"], "action_required")
        self.assertEqual(parsed["rationale"], "New due date applies")
        self.assertEqual(parsed["owner"], "Reg Reporting")
        self.assertEqual(parsed["ticket"], "https://tracker.example.com/REG-1")
        self.assertEqual(parsed["events"], [{"event_id": event_id, "url": "https://www.fed.gov/a?b=1"}])

        for decision in ("material", "not_material", "acknowledged"):
            with self.subTest(decision=decision):
                body = run_js("console.log(JSON.stringify(L.changeReviewRequest({decision: " + json.dumps(decision) + ", rationale: 'Only affects a form we do not file.', "
                              "changes: [{event_id: " + json.dumps(event_id) + "}]}).body));")
                parsed = reviews.parse_request(body)
                self.assertEqual(parsed["decision"], decision)
                self.assertEqual((parsed["owner"], parsed["ticket"]), ("", ""))

    def test_review_request_refuses_what_the_workflow_would_reject(self):
        out = run_js("""
          const id = "b".repeat(64);
          console.log(JSON.stringify([
            L.changeReviewRequest({decision: "not_material", rationale: "short", changes: [{event_id: id}]}).errors.length,
            L.changeReviewRequest({decision: "material", ticket: "javascript:alert(1)", changes: [{event_id: id}]}).errors.length,
            L.changeReviewRequest({decision: "maybe", changes: [{event_id: id}]}).errors.length,
            L.changeReviewRequest({decision: "material", changes: [{event_id: "not-an-id"}]}).errors.length,
            L.changeReviewRequest({decision: "material", changes: [{event_id: id}]}).errors.length
          ]));""")
        self.assertEqual(out, [1, 1, 1, 1, 0])

    def test_csv_cells_cannot_run_as_formulas(self):
        out = run_js("console.log(JSON.stringify([L.csvCell('=HYPERLINK(\"x\")'), L.csvCell('+1'), L.csvCell('@SUM(A1)'), L.csvCell('-2'), "
                     "L.csvCell('a, b'), L.csvCell('say \"hi\"'), L.csvCell(null), L.csvCell('plain'), L.toCsv([['a', 'b'], ['c', 'd']])]));")
        self.assertEqual(out, ["\"'=HYPERLINK(\"\"x\"\")\"", "'+1", "'@SUM(A1)", "'-2", '"a, b"', '"say ""hi"""', "", "plain", "a,b\r\nc,d\r\n"])


@unittest.skipUnless(NODE, "node is required for dashboard library tests")
class TestAddLinkHelpers(unittest.TestCase):
    """The Add a link checklist: every link stays listed and selectable; requests are bounded and honest."""

    ITEMS = [
        {"key": "a", "url": "https://docs.example.gov/BC.pdf", "title": "Form BC", "format": "PDF", "section": "Current forms",
         "parentUrl": "https://home.example.gov/b-forms", "parentTitle": "TIC B Forms"},
        {"key": "b", "url": "https://docs.example.gov/BC-2019.pdf", "title": "Form BC, 2019", "format": "PDF", "section": "Historical forms",
         "parentUrl": "https://home.example.gov/b-forms", "parentTitle": "TIC B Forms"},
        {"key": "c", "url": "https://home.example.gov/auctions", "title": "Auction allotments", "format": "Web page", "section": "",
         "parentUrl": "https://home.example.gov/forms", "parentTitle": "TIC Forms"},
    ]

    def test_groups_keep_historical_links_and_filters_only_narrow_the_view(self):
        out = run_js("const items = " + json.dumps(self.ITEMS) + ";\nconsole.log(JSON.stringify({" +
                     "all: L.pickGroupBy(items, {}).map(g => [g.parentTitle, g.section, g.items.length])," +
                     "pdf: L.pickGroupBy(items, {format: 'PDF'}).reduce((n, g) => n + g.items.length, 0)," +
                     "search: L.pickGroupBy(items, {query: '2019'}).reduce((n, g) => n + g.items.length, 0)," +
                     "ticked: L.pickGroupBy(items, {show: 'SELECTED', selectedKeys: new Set(['a'])}).reduce((n, g) => n + g.items.length, 0)}));")
        self.assertEqual(out["all"], [["TIC B Forms", "Current forms", 1], ["TIC B Forms", "Historical forms", 1], ["TIC Forms", "", 1]])
        self.assertEqual(out["pdf"], 2)
        self.assertEqual(out["search"], 1)
        self.assertEqual(out["ticked"], 1)

    def test_following_links_reads_only_the_chosen_section(self):
        out = run_js("""
          const start = 'https://home.example.gov/data/tic/forms';
          console.log(JSON.stringify([
            L.pickInScope('https://home.example.gov/data/tic/forms/b-forms', start, 'https://home.example.gov/data/tic/'),
            L.pickInScope('https://home.example.gov/data/tic', start, 'https://home.example.gov/data/tic/'),
            L.pickInScope('https://home.example.gov/data/tic-archive', start, 'https://home.example.gov/data/tic/'),
            L.pickInScope('https://home.example.gov/data/auctions', start, 'https://home.example.gov/data/tic/'),
            L.pickInScope('https://home.example.gov/anything', start, ''),
            L.pickInScope('https://www.example.gov/data/tic/x', start, '')
          ]));
        """)
        self.assertEqual(out, [True, True, False, False, True, False])

    def test_summary_counts_monitored_and_duplicate_links_honestly(self):
        out = run_js("""
          const rows = [{url: 'https://a.gov/x.pdf'}, {url: 'https://a.gov/x.pdf/'}, {url: 'https://a.gov/y.pdf'}, {url: 'https://a.gov/z'}];
          const s = L.pickSummary(rows, new Set([L.urlKey('https://a.gov/y.pdf')]));
          console.log(JSON.stringify({selected: s.selected, monitored: s.monitored, duplicates: s.duplicates, toRequest: s.toRequest}));
        """)
        self.assertEqual(out, {"selected": 3, "monitored": 1, "duplicates": 1, "toRequest": 2})

    def test_requests_are_bounded_and_nothing_is_silently_dropped(self):
        out = run_js("""
          const rows = Array.from({length: 1214}, (_, i) => ({url: 'https://a.gov/d/' + i + '.pdf', name: 'Document ' + i}));
          rows.push({url: 'https://a.gov/' + 'x'.repeat(8000), name: 'too long'});
          const build = batch => 'https://github.com/o/r/issues/new?body=' + encodeURIComponent(JSON.stringify(batch));
          const r = L.pickBatches(rows, build, 25, 7000);
          console.log(JSON.stringify({batches: r.batches.length, max: Math.max(...r.batches.map(b => b.length)),
            total: r.batches.reduce((n, b) => n + b.length, 0), longest: Math.max(...r.batches.map(b => build(b).length)), tooLong: r.tooLong.length}));
        """)
        self.assertEqual(out["total"], 1214, "every ticked link is in exactly one request")
        self.assertLessEqual(out["max"], 25)
        self.assertLessEqual(out["longest"], 7000)
        self.assertEqual(out["tooLong"], 1, "an address too long for any request is reported, not truncated")

    def test_scope_choices_and_document_hosts(self):
        out = run_js("""
          console.log(JSON.stringify({
            scope: L.pickScopeOptions('https://home.example.gov/data/tic/forms').map(o => [o.id, o.prefix]),
            top: L.pickScopeOptions('https://home.example.gov/').map(o => o.id),
            hosts: L.pickHosts('www.example.gov, https://docs.example.gov/x, home.example.gov, localhost, 10.0.0.1, bad_host', 'https://home.example.gov/a')
          }));
        """)
        self.assertEqual(out["scope"], [["page", "https://home.example.gov/data/tic/forms/"],
                                        ["section", "https://home.example.gov/data/tic/"], ["site", ""]])
        self.assertEqual(out["top"], ["site"])
        self.assertEqual(out["hosts"]["hosts"], ["www.example.gov", "docs.example.gov"])
        self.assertEqual(out["hosts"]["invalid"], ["localhost", "10.0.0.1", "bad_host"])

    def test_the_weekly_check_request_is_what_registration_parses(self):
        out = run_js("""
          console.log(JSON.stringify(L.pickDiscoveryRequest({url: 'https://home.example.gov/data/tic/forms', name: 'TIC\\nforms', depth: 2,
            explore: 'https://home.example.gov/data/tic/', documentHosts: ['www.example.gov']})));
        """)
        import tempfile
        from regmon.register_monitor import register_request
        root = Path(__file__).resolve().parents[1]
        with tempfile.TemporaryDirectory() as tmp:
            sources = Path(tmp) / "sources.json"
            scout = Path(tmp) / "scout.json"
            sources.write_text(json.dumps({"version": 4, "defaults": {}, "sources": []}), encoding="utf-8")
            scout.write_text((root / "config" / "scout.json").read_text(encoding="utf-8"), encoding="utf-8")
            event = {"issue": {"title": out["title"], "body": out["body"], "user": {"login": "GauravDarwesh"}}}
            result = register_request(event, sources, "GauravDarwesh/ChangeSentinel", scout)
            saved_sources = json.loads(sources.read_text(encoding="utf-8"))["sources"]
            channel = next(c for r in json.loads(scout.read_text(encoding="utf-8"))["regulators"] for c in r["channels"]
                           if c["url"] == "https://home.example.gov/data/tic/forms")
        self.assertTrue(out["title"].startswith("[ChangeSentinel] Add monitor"))
        self.assertTrue(result["discovery_only"])
        self.assertEqual(saved_sources, [], "the weekly check never adds a monitor")
        self.assertEqual(channel["crawl_depth"], 2)
        self.assertEqual(channel["explore_prefixes"], ["https://home.example.gov/data/tic/"])
        self.assertEqual(channel["document_hosts"], ["www.example.gov"])

    def test_thin_link_text_is_named_with_its_section(self):
        out = run_js("console.log(JSON.stringify([L.pickMonitorName({title: 'Instructions', section: 'Form BC'}),"
                     " L.pickMonitorName({title: 'Treasury International Capital Form BC', section: 'Current'})]))")
        self.assertEqual(out, ["Form BC: Instructions", "Treasury International Capital Form BC"])


@unittest.skipUnless(NODE, "node is required for dashboard library tests")
class TestChangeFieldLenses(unittest.TestCase):
    """The scan history keeps a row for each page that changed or failed and folds the rest."""

    PAGES = """
      const mk = (id, title, states, sourceId) => ({ id, title, url: "https://x.gov/" + id, sourceId: sourceId || "s", sourceLabel: (sourceId || "s").toUpperCase(), states });
      const pages = [
        mk("a", "Alpha", ["dim", "on", "dim"]),
        mk("b", "Bravo", ["new", "dim", "dim"]),
        mk("c", "Charlie", ["dim", "dim", "markup"]),
        mk("d", "Delta", ["dim", "bad", "dim"]),
        mk("e", "Echo", ["none", "none", "none"]),
        mk("f", "Foxtrot", ["skip", "dim", "dim"], "t"),
        mk("g", "Golf", ["dim", "dim", "dim"], "t")
      ];
      const describe = rows => rows.map(r => r.type === "page" ? (r.member ? "  " : "") + r.page.title : r.type === "group" ? "[" + r.label + " " + r.count + (r.open ? " open" : "") + "] " + r.states.join(",") : "(+" + r.hidden + ")");
    """

    def test_activity_keeps_changes_and_problems_and_folds_the_rest(self):
        out = run_js(self.PAGES + "console.log(JSON.stringify(describe(L.fieldRows(pages, {}).rows)));")
        self.assertEqual(out, [
            "Delta", "Alpha",
            "[New pages 1] new,dim,dim",
            "[No change 3] dim,dim,markup",
            "[Not scanned yet 1] none,none,none",
        ])

    def test_a_summary_row_opens_in_place_and_long_lists_page(self):
        out = run_js(self.PAGES + "console.log(JSON.stringify(describe(L.fieldRows(pages, { open: new Set(['quiet']), limit: 4 }).rows)));")
        self.assertEqual(out, ["Delta", "Alpha", "[New pages 1] new,dim,dim", "[No change 3 open] dim,dim,markup", "  Charlie", "  Foxtrot", "[Not scanned yet 1] none,none,none", "(+1)"])

    def test_many_single_page_sources_past_the_limit_get_one_show_more(self):
        # Regression: every single-page source after the limit used to add its own "Show 1 more" row.
        out = run_js("""
          const pages = Array.from({ length: 30 }, (_, i) => ({ id: "p" + i, title: "Page " + String(i).padStart(2, "0"), sourceId: "s" + i, sourceLabel: "S" + i, states: ["dim"] }));
          const rows = L.fieldRows(pages, { lens: "all", limit: 10 }).rows;
          console.log(JSON.stringify({ more: rows.filter(r => r.type === "more").map(r => r.hidden), pages: rows.filter(r => r.type === "page").length }));""")
        self.assertEqual(out, {"more": [20], "pages": 10})

    def test_pages_are_sectioned_by_regulator_with_a_summary_strip(self):
        out = run_js(self.PAGES + """
          const tagged = pages.map(p => Object.assign(p, p.sourceId === "t" ? { groupKey: "osfi", groupLabel: "OSFI", groupNote: "Canada" } : { groupKey: "fed", groupLabel: "Federal Reserve", groupNote: "United States" }));
          const rows = L.fieldRows(tagged, { groupBy: "regulator" }).rows;
          const closedRows = L.fieldRows(tagged, { groupBy: "regulator", closed: new Set(["sec:fed"]) }).rows;
          console.log(JSON.stringify({
            rows: rows.map(r => r.type === "section" ? "## " + r.label + " " + r.count + " " + r.states.join(",") : describe([r])[0]),
            closed: closedRows.map(r => r.type === "section" ? "## " + r.label + (r.open ? "" : " (closed)") : describe([r])[0]),
            counts: rows.filter(r => r.type === "section").map(r => r.counts.activity)
          }));""")
        self.assertEqual(out["rows"], [
            "## Federal Reserve 5 new,bad,markup", "Delta", "Alpha", "[New pages 1] new,dim,dim", "[No change 1] dim,dim,markup", "[Not scanned yet 1] none,none,none",
            "## OSFI 2 dim,dim,dim", "[No change 2] dim,dim,dim",
        ])
        self.assertEqual(out["closed"], ["## Federal Reserve (closed)", "## OSFI", "[No change 2] dim,dim,dim"])
        self.assertEqual(out["counts"], [2, 0])

    def test_regulators_are_found_by_host_then_by_domain(self):
        out = run_js("""
          const profiles = [{ id: "fed", short_name: "Federal Reserve", jurisdiction: "United States", hosts: ["federalreserve.gov", "www.frbservices.org"] }];
          const pick = (url, name) => { const r = L.regulatorFor(url, profiles, name); return [r.id, r.short, r.jurisdiction]; };
          console.log(JSON.stringify([
            pick("https://www.frbservices.org/a.xlsx", "www.frbservices.org"),
            pick("https://www.federalreserve.gov/apps/x", ""),
            pick("https://www.centralbankbahamas.com/x", "www.centralbankbahamas.com"),
            pick("https://www.osfi-bsif.gc.ca/en/x", "www.osfi-bsif.gc.ca"),
            pick("https://www.eba.europa.eu/x", "European Banking Authority"),
            [L.pageKind("FR Y-9C instructions", ""), L.pageKind("Item 2b standard template", "https://x/item-2b.xlsx"), L.pageKind("TFC-1 Sample File (XML)", ""), L.pageKind("BCAR FAQ technical specification", "")]
          ]));""")
        self.assertEqual(out[0], ["fed", "Federal Reserve", "United States"])
        self.assertEqual(out[1], ["fed", "Federal Reserve", "United States"])
        self.assertEqual(out[2], ["centralbankbahamas.com", "centralbankbahamas.com", ""])
        self.assertEqual(out[3], ["osfi-bsif.gc.ca", "osfi-bsif.gc.ca", "Canada"])
        self.assertEqual(out[4], ["eba.europa.eu", "European Banking Authority", "European Union"])
        self.assertEqual(out[5], ["Instructions", "Forms and templates", "Sample files", "Technical specifications"])

    def test_problems_outrank_changes_when_lights_merge(self):
        out = run_js("console.log(JSON.stringify([L.mergeCellStates(['dim', 'on', 'bad']), L.mergeCellStates(['skip', 'markup']), L.mergeCellStates([])]));")
        self.assertEqual(out, ["bad", "markup", "none"])

    def test_lenses_search_and_source_groups(self):
        out = run_js(self.PAGES + """
          console.log(JSON.stringify({
            changed: describe(L.fieldRows(pages, { lens: "changed" }).rows),
            problems: describe(L.fieldRows(pages, { lens: "problems" }).rows),
            fresh: describe(L.fieldRows(pages, { lens: "new" }).rows),
            all: describe(L.fieldRows(pages, { lens: "all" }).rows),
            closed: describe(L.fieldRows(pages, { lens: "all", closed: new Set(["src:t"]) }).rows),
            search: describe(L.fieldRows(pages, { query: "golf" }).rows),
            counts: L.fieldRows(pages, {}).counts
          }));""")
        self.assertEqual(out["changed"], ["Alpha"])
        self.assertEqual(out["problems"], ["Delta"])
        self.assertEqual(out["fresh"], ["Bravo"])
        self.assertEqual(out["all"][0], "[S 5 open] new,bad,markup")
        self.assertIn("[T 2 open] dim,dim,dim", out["all"])
        self.assertIn("[T 2] dim,dim,dim", out["closed"])
        self.assertNotIn("Golf", out["closed"])
        self.assertEqual(out["search"], ["Golf"])
        self.assertEqual(out["counts"], {"activity": 2, "changed": 1, "problems": 1, "new": 1, "quiet": 3, "waiting": 1, "all": 7})

    def test_the_range_keeps_only_recent_windows(self):
        out = run_js("""
          const day = 86400000, now = Date.UTC(2026, 9, 23, 12);
          const cols = [10, 5, 1.5, 0.2].map(d => ({ id: String(d), start: now - d * day, end: now - d * day + 6 * 3600000 }));
          console.log(JSON.stringify(["2d", "7d", "30d", "all"].map(r => L.fieldWindow(cols, r, now).map(c => c.id))));""")
        self.assertEqual(out, [["1.5", "0.2"], ["5", "1.5", "0.2"], ["10", "5", "1.5", "0.2"], ["10", "5", "1.5", "0.2"]])


@unittest.skipUnless(NODE, "node is required for dashboard library tests")
class TestScoutDigest(unittest.TestCase):
    def test_findings_group_by_kind_regulator_and_day(self):
        out = run_js("""
          const f = (id, reg, kind, seen, confidence) => ({ key: id, regulator_id: reg, first_seen: seen, assessment: { confidence, publication_type: kind ? { id: kind, label: kind.toUpperCase() } : null } });
          const items = [f(1, "fed", "rule", "2026-10-21T09:00:00Z", "strong"), f(2, "eba", "rule", "2026-10-20T09:00:00Z", "weak"),
                         f(3, "fed", "guidance", "2026-10-21T10:00:00Z", "supported"), f(4, "fed", null, "2026-10-19T09:00:00Z", "conflicting")];
          const keys = groups => groups.map(g => g.label + ":" + g.items.map(i => i.key).join(","));
          console.log(JSON.stringify({
            kind: keys(L.scoutGroups(items, "kind")),
            regulator: keys(L.scoutGroups(items, "regulator", id => id.toUpperCase())),
            day: L.scoutGroups(items, "date").map(g => g.key),
            strength: items.map(i => L.scoutStrength(i.assessment).level + L.scoutStrength(i.assessment).tone)
          }));""")
        self.assertEqual(out["kind"], ["RULE:1,2", "GUIDANCE:3", "Kind not recognised:4"])
        self.assertEqual(out["regulator"], ["FED:1,3,4", "EBA:2"])
        self.assertEqual(out["day"], ["day:2026-10-21", "day:2026-10-20", "day:2026-10-19"])
        self.assertEqual(out["strength"], ["3strong", "1weak", "2supported", "1conflicting"])


@unittest.skipUnless(NODE, "node is required for dashboard library tests")
class TestSmartPicking(unittest.TestCase):
    LINKS = """[
      { url: "https://x.gov/forms/2026/y9c.pdf", text: "Current FR Y-9C form", section: "Current", format: "PDF" },
      { url: "https://x.gov/forms/2026/y9c-instr.pdf", text: "Current FR Y-9C instructions", section: "Current", format: "PDF" },
      { url: "https://x.gov/forms/2026/y9c.xlsx", text: "Current FR Y-9C template", section: "Current", format: "Spreadsheet" },
      { url: "https://x.gov/forms/archive/y9c-2024.pdf", text: "Annual report 2024", section: "Archive", format: "PDF" },
      { url: "https://x.gov/forms/archive/y9c-2025.pdf", text: "Annual report 2025", section: "Archive", format: "PDF" },
      { url: "https://x.gov/news", text: "News", section: "", format: "Web page", chrome: true }
    ]"""

    def test_scores_reasons_and_older_versions(self):
        out = run_js("const links = " + self.LINKS + """;
          console.log(JSON.stringify({ scores: links.map(l => L.pickScore(l).score), reasons: L.pickScore(links[1]).reasons,
            older: Array.from(L.pickVersions(links)) }));""")
        self.assertGreater(out["scores"][2], out["scores"][3])
        self.assertLess(out["scores"][5], 0)
        self.assertIn("Instructions", out["reasons"])
        self.assertEqual(out["older"], [3])

    def test_similar_links_and_the_inferred_rule(self):
        out = run_js("const links = " + self.LINKS + """;
          const similar = L.pickSimilar([links[0]], links, l => l.url).map(r => [r.label, r.extra]);
          console.log(JSON.stringify({ similar, recipe: L.pickRecipe("https://x.gov/forms", [links[0], links[2]], links) }));""")
        self.assertEqual(out["similar"], [["All PDFs starting “Current FR”", 1], ["Every PDF on this page", 3]])
        self.assertEqual(out["recipe"]["textPattern"], "^Current")
        self.assertEqual(out["recipe"]["prefix"], "https://x.gov/forms/2026/")
        self.assertEqual(out["recipe"]["count"], 3)

    def test_rule_matches_folders_and_request(self):
        out = run_js("const links = " + self.LINKS + """;
          const req = L.pickRuleRequest({ url: "https://x.gov/forms", name: "Forms\\nPage", folder: "https://x.gov/forms/2026/y9c.pdf",
            pattern: "^Current", include: "main", ignore: ["div.banner"] });
          console.log(JSON.stringify({
            folders: ["https://x.gov/forms", "https://x.gov/forms/a.pdf", "https://x.gov/forms/2026/", "ftp://x"].map(L.pickRuleFolder),
            match: L.pickRuleMatches(links, { folder: "https://x.gov/forms/2026/", pattern: "^current" }).count,
            anywhere: L.pickRuleMatches(links, { folder: "", pattern: "" }).count,
            bad: L.pickRuleMatches(links, { pattern: "(" }).error !== "",
            title: req.title, body: req.body }));""")
        self.assertEqual(out["folders"], ["https://x.gov/forms/", "https://x.gov/forms/", "https://x.gov/forms/2026/", ""])
        self.assertEqual(out["match"], 3)
        self.assertEqual(out["anywhere"], 5)
        self.assertTrue(out["bad"])
        self.assertTrue(out["title"].startswith("[ChangeSentinel] Add monitor request"))
        lines = out["body"].splitlines()
        self.assertIn("Name: Forms Page", lines)
        self.assertLess(lines.index("Documents under: https://x.gov/forms/2026/"), lines.index("Include selector: main"))
        self.assertIn("- div.banner", lines)


@unittest.skipUnless(NODE, "node is required for dashboard library tests")
class TestCadencePlanning(unittest.TestCase):
    def test_quarterly_plan_rolls_weekends_and_places_steps_in_business_days(self):
        out = run_js("""console.log(JSON.stringify(L.cadencePlan({ frequency: "quarterly", yearEndMonth: 11, lag: 40, lagLast: 45, business: true,
          milestones: [{ label: "Data cut-off", offset: 10 }, { label: "Sign-off", offset: 2 }, { label: "", offset: 3 }], cycles: 4, from: "2026-10-23" })));""")
        self.assertEqual([c["due"] for c in out], ["2026-11-09", "2027-02-15", "2027-05-10", "2027-08-09"])
        self.assertEqual([c["periodEnd"] for c in out], ["2026-09-30", "2026-12-31", "2027-03-31", "2027-06-30"])
        self.assertTrue(out[1]["rolled"])  # 14 Feb 2027 is a Sunday
        self.assertEqual(out[0]["period"].split(" ")[0], "Q3")
        self.assertEqual([(m["label"], m["date"]) for m in out[0]["milestones"]], [("Data cut-off", "2026-10-26"), ("Sign-off", "2026-11-05")])

    def test_fiscal_quarters_monthly_and_yearly_plans(self):
        out = run_js("""console.log(JSON.stringify({
          fiscal: L.cadencePlan({ frequency: "quarterly", yearEndMonth: 2, lag: 30, cycles: 2, from: "2026-10-23" }).map(c => [c.period.split(" ")[0], c.due]),
          monthly: L.cadencePlan({ frequency: "monthly", lag: 15, cycles: 3, from: "2026-10-23" }).map(c => c.due),
          yearly: L.cadencePlan({ frequency: "yearly", yearEndMonth: 2, lag: 90, cycles: 1, from: "2026-10-23" }).map(c => c.due),
          shift: [L.businessShift("2026-10-26", -1), L.businessShift("2026-10-23", 1), L.nextBusinessDay("2026-10-24")] }));""")
        self.assertEqual(out["fiscal"], [["Q2", "2026-10-30"], ["Q3", "2027-01-30"]])
        self.assertEqual(out["monthly"], ["2026-11-15", "2026-12-15", "2027-01-15"])
        self.assertEqual(out["yearly"], ["2027-06-29"])
        self.assertEqual(out["shift"], ["2026-10-23", "2026-10-26", "2026-10-26"])

    def test_weekly_load_and_calendar_links(self):
        out = run_js("""
          const events = [{ key: "a", date: "2026-10-27", title: "A", deadline: true }, { key: "b", date: "2026-10-28", title: "B", deadline: true },
            { key: "c", date: "2026-10-29", title: "C", deadline: true }, { key: "c", date: "2026-10-29", title: "C again", deadline: true },
            { key: "d", date: "2026-11-04", title: "D" }, { key: "e", date: "2026-11-05", title: "E", completed: true }];
          const weeks = L.cadenceWeeks(events, "2026-10-23", 3);
          console.log(JSON.stringify({ weeks: weeks.map(w => [w.start, w.total, w.deadlines, w.busy, w.level]),
            link: L.googleCalendarLink({ date: "2026-11-09", title: "FR Y-9C due", url: "https://x.gov" }), bad: L.googleCalendarLink({ date: "soon" }) }));""")
        self.assertEqual(out["weeks"], [["2026-10-19", 0, 0, False, 0], ["2026-10-26", 3, 3, True, 4], ["2026-11-02", 1, 0, False, 1]])
        self.assertIn("dates=20261109%2F20261110", out["link"])
        self.assertEqual(out["bad"], "")


@unittest.skipUnless(NODE, "node is required for dashboard library tests")
class TestTimelineModel(unittest.TestCase):
    """The timeline places events on lanes and tells them as a story."""

    SETUP = """
const H = 3600e3, s = Date.UTC(2026, 9, 17), e = s + 7 * 24 * H;
const ev = (id, h, kind, lane, page) => ({ id, at: s + h * H, kind, laneKey: lane, laneLabel: lane.toUpperCase(), pageId: page || id, title: id });
"""

    def model(self, body, lens="activity", gap_hours=6):
        return run_js(self.SETUP + "const E = [" + body + "];\n"
                      "const m = L.timelineModel(E, { start: s, end: e, gapMs: " + str(gap_hours) + " * H, lens: " + json.dumps(lens) + " });\n"
                      "console.log(JSON.stringify(m));")

    def test_changes_far_apart_keep_their_own_beads(self):
        m = self.model("ev('a', 30, 'on', 'fr'), ev('b', 78, 'on', 'fr')", gap_hours=24)
        self.assertEqual([b["count"] for b in m["lanes"][0]["beads"]], [1, 1])

    def test_events_close_together_share_one_bead_led_by_the_worst(self):
        m = self.model("ev('a', 30, 'new', 'fr'), ev('b', 31, 'bad', 'fr')")
        bead = m["lanes"][0]["beads"][0]
        self.assertEqual((bead["count"], bead["top"]), (2, "bad"))

    def test_lanes_with_real_changes_lead_a_more_recent_burst_of_new_pages(self):
        m = self.model("ev('a', 10, 'on', 'fr'), ev('n', 150, 'new', 'osfi')")
        self.assertEqual([lane["key"] for lane in m["lanes"]], ["fr", "osfi"])

    def test_a_burst_of_new_pages_is_one_story_row(self):
        m = self.model(",".join("ev('n%d', 100, 'new', 'fr')" % i for i in range(10)))
        rows = [row for day in m["days"] for row in day["rows"]]
        self.assertEqual(len(rows), 1)
        self.assertEqual((rows[0]["type"], rows[0]["count"]), ("burst", 10))

    def test_repeated_markup_edits_on_one_page_are_one_line_per_day(self):
        m = self.model("ev('m1', 100, 'markup', 'fr', 'p'), ev('m2', 103, 'markup', 'fr', 'p'), ev('c', 104, 'on', 'fr', 'q')")
        rows = [row for day in m["days"] for row in day["rows"]]
        self.assertEqual(sorted(row.get("repeat", 1) for row in rows), [1, 2])

    def test_changed_lens_leaves_out_markup_only_edits(self):
        m = self.model("ev('m', 100, 'markup', 'fr'), ev('c', 104, 'on', 'fr')", lens="changed")
        self.assertEqual(m["counts"]["total"], 1)

    def test_every_story_row_points_at_a_bead_that_exists(self):
        m = self.model("ev('a', 30, 'on', 'fr'), ev('b', 31, 'new', 'fr'), ev('c', 120, 'bad', 'osfi')")
        beads = {b["key"] for lane in m["lanes"] for b in lane["beads"]}
        for day in m["days"]:
            for row in day["rows"]:
                self.assertTrue(set(row["beadKeys"]) <= beads)
