"""Structured spec diffs: rule workbooks aligned by rule ID, field-spec tables
keyed by code, XSD/XBRL schemas, declared-vs-detected reconciliation, the
engine/alert wiring and the offline analyst workbook.

Real regulator files saved in data/snapshots are used where the repository has
them (OSFI BCAR technical specification, BCAR XSD); those tests skip if the
snapshots are ever pruned.
"""
from __future__ import annotations

import contextlib
import io
import json
import re
import tempfile
import unittest
from pathlib import Path

from engine_harness import Sandbox, ok
from regmon import spec_diff
from regmon.config import SourceConfig
from regmon.workbook import extract_workbook

ROOT = Path(__file__).resolve().parent.parent
SNAPSHOTS = ROOT / "data" / "snapshots"
OSFI_SPEC = "db3db09fda63b39f8dc40b97b712a5c23bcefd3c94615e1e9c990d45533cc00c"
BCAR_XSD = "2673ea372676d2552ed00099f87085c12b153d5447df3e81fbf08d4a1f7ed9a9"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

EDIT_HEADER = ["Series", "Database Table", "Edit No", "Version No.", "Start Date", "End Date", "Status", "Edit Type",
               "Technical Field Name", "Report Form Line Item", "Error Message", "Edit Logic", "Reporting Period Modified Date", "Comments"]


def edit_rows() -> list[list[str]]:
    rows = []
    for number in range(100, 130):
        rows.append(["FRY14Q", "CIL", str(number), "1", "19000101", "99991231", "No change", "Validity" if number % 2 else "Quality",
                     f"FIELD{number}", f"Line {number}", f"Field {number} must be valid", f"FIELD{number} IN (1,2,3)", "20250331", ""])
    rows.append(["FRY14Q", "CIL", "105", "1", "19000101", "99991231", "No change", "Validity", "TaxStatus", "Line 45",
                 "Tax status must be 1 or 2", "TaxStatus IN (1,2)", "20250331", ""])
    rows = [row for row in rows if not (row[2] == "105" and row[8] != "TaxStatus")]
    rows.append(["FRY14Q", "CIL", "329", "1", "19000101", "99991231", "No change", "Intraseries", "TaxStatus", "Line 45",
                 "Tax status changed from prior quarter", "TaxStatus = TaxStatus[P-1Q]", "20250331", ""])
    return rows


def edit_workbook(rows: list[list[str]]) -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "CIL Edits"
    sheet.append(EDIT_HEADER)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def current_edit_rows() -> list[list[str]]:
    rows = [list(row) for row in edit_rows()]
    for row in rows:
        if row[2] in {"105", "329"}:
            row[5], row[6], row[12] = "20260630", "Archived", "20260930"
        if row[2] == "110":
            row[11], row[6] = "FIELD110 IN (1,2,3,4)", "Modified"
        if row[2] == "112":
            row[7] = "Validity"          # Quality → Validity: now blocking, unannounced
        if row[2] == "114":
            row[10] = "Field 114 must be a valid code"
        if row[2] == "116":
            row[6] = "Modified"          # declared, nothing changed
        if row[2] == "120":
            row[5], row[6] = "20260630", "Archived"
    rows.append(["FRY14Q", "CIL", "531", "1", "19000101", "99991231", "New", "Syntax", "LMBEFD", "Line 60",
                 "LMBEFD must be a number", "ISNUMBER(LMBEFD)", "20260930", ""])
    rows.append(["FRY14Q", "CIL", "532", "1", "20260930", "99991231", "New", "Validity", "LMBEFD", "Line 60",
                 "LMBEFD must be >= 0", "LMBEFD >= 0", "20260930", ""])
    rows.append(["FRY14Q", "CIL", "120", "2", "20260930", "99991231", "New", "Validity", "FIELD120", "Line 120",
                 "Field 120 must be valid", "FIELD120 IN (1,2)", "20260930", ""])
    return rows


def workbook_text(body: bytes) -> str:
    text, kind, error = extract_workbook(body)
    assert error is None, error
    return text


def by_label(result: dict) -> dict[str, dict]:
    return {change["label"]: change for table in result["tables"] for change in table["changes"]}


class TestRuleWorkbooks(unittest.TestCase):
    """FR Y-14Q style edit checks, as the team's compare tool sees them."""

    @classmethod
    def setUpClass(cls):
        cls.result = spec_diff.diff(workbook_text(edit_workbook(edit_rows())), workbook_text(edit_workbook(current_edit_rows())))
        cls.changes = by_label(cls.result)

    def test_rows_are_aligned_by_the_smallest_unique_rule_key(self):
        table = self.result["tables"][0]
        self.assertEqual(table["kind"], "rules")
        self.assertEqual(table["key"], ["Edit No", "Version No."])
        self.assertEqual((table["rows_before"], table["rows_after"]), (31, 34))

    def test_new_rules_and_new_versions(self):
        self.assertEqual(self.changes["531 v1"]["category"], "added")
        self.assertEqual(self.changes["531 v1"]["effective"], "applies to all periods")
        self.assertEqual(self.changes["532 v1"]["effective"], "applies from 2026-09-30")
        self.assertEqual(self.changes["532 v1"]["quarter"], "Q3 2026")
        self.assertEqual(self.changes["120 v2"]["category"], "added_version")
        self.assertIn("120 v1", self.changes["120 v2"]["note"])

    def test_retired_rules_are_checked_for_a_replacement(self):
        retired = self.changes["105 v1"]
        self.assertEqual(retired["category"], "retired")
        self.assertEqual(retired["effective"], "last applies to 2026-06-30")
        self.assertEqual(retired["quarter"], "Q2 2026")
        self.assertIn("No live replacement", retired["note"])
        self.assertEqual(self.changes["120 v1"]["note"], "Superseded by 120 v2")
        self.assertEqual(self.changes["329 v1"]["field"], "TaxStatus")

    def test_edit_type_escalation_is_a_severity_change(self):
        change = self.changes["112 v1"]
        self.assertEqual(change["category"], "severity")
        self.assertEqual(change["severity"], {"before": "Quality", "after": "Validity", "direction": "escalated"})
        self.assertEqual(self.result["escalated"], 1)

    def test_logic_and_message_changes(self):
        self.assertEqual(self.changes["110 v1"]["category"], "logic")
        self.assertEqual(self.changes["114 v1"]["category"], "message")

    def test_declared_and_detected_changes_are_reconciled(self):
        self.assertEqual(self.changes["112 v1"]["mismatch"], "Changed but marked “No change”")
        self.assertIn("Changed but marked", self.changes["114 v1"]["mismatch"])
        self.assertIn("Marked “Modified”", self.changes["116 v1"]["mismatch"])
        self.assertNotIn("mismatch", self.changes["110 v1"])  # changed and declared Modified
        self.assertEqual(self.result["mismatches"], 3)

    def test_headline_summary_and_alert_line(self):
        self.assertIn("new rule 2 (531 v1, 532 v1)", self.result["headline"])
        self.assertIn("1 escalated", self.result["headline"])
        summary = spec_diff.summary(self.result)
        self.assertEqual(summary["counts"]["retired"], 3)
        self.assertIn("TaxStatus", summary["fields"])
        line = spec_diff.describe(summary)
        self.assertTrue(line.startswith("2 new, 1 new version, 1 severity changed, 1 logic changed, 3 retired, 1 message changed"), line)
        self.assertIn("(1 escalated)", line)

    def test_the_older_row_text_compares_equal_to_the_structured_text(self):
        from regmon.content import _extract_xlsx

        body = edit_workbook(edit_rows())
        older, _kind, error = _extract_xlsx(body)
        self.assertIsNone(error)
        self.assertIsNone(spec_diff.diff(older, workbook_text(body)))

    def test_no_structure_means_no_structured_diff(self):
        self.assertIsNone(spec_diff.diff("Plain prose.\nMore prose.", "Plain prose, edited."))
        self.assertIsNone(spec_diff.diff(workbook_text(edit_workbook(edit_rows())), "an extraction that lost the table"))


class TestFieldSpecTables(unittest.TestCase):
    """Data dictionary tables as the PDF extractors write them (page headers repeat)."""

    HEADER = ("| Series | Database Table | StartDate | EndDate | Technical Field Name / MDRM | Report Form Line Item | "
              "FieldType (DateTime / Varchar / Numeric) | Item Descriptions | Status |")

    def doc(self, name: str, rows: list[str]) -> str:
        pages = [rows[:3], rows[3:]]
        body = []
        for page in pages:
            body += [self.HEADER, "| --- | --- | --- | --- | --- | --- | --- | --- | --- |", *page, "", "Page footer text", ""]
        return f"[FILE] fry-14q-ti/{name} (pdf)\n" + "\n".join(body)

    def test_fields_are_keyed_by_code_across_renamed_period_documents(self):
        before = [
            "| FRY14Q | CIL | 20110930 | 99991231 | CustomerID | Customer ID | varchar(50) | Customer ID | No Change |",
            "| FRY14Q | CIL | 20110930 | 99991231 | TaxStatus | Interest Income Tax Status | Numeric (1,0) | Interest Income Tax Status | No Change |",
            "| FRY14Q | CIL | 20110930 | 99991231 | City | City | varchar(73) | City | No Change |",
            "| FRY14Q | CIL | 20110930 | 99991231 | Country | Country | varchar(2) | Country | No Change |",
        ]
        after = [
            "| FRY14Q | CIL | 20110930 | 99991231 | CustomerID | Customer ID | varchar(100) | Customer ID | Modified |",
            "| FRY14Q | CIL | 20110930 | 20260630 | TaxStatus | Interest Income Tax Status | Numeric (1,0) | Interest Income Tax Status | Archived |",
            "| FRY14Q | CIL | 20110930 | 99991231 | City | City | varchar(73) | City | No Change |",
            "| FRY14Q | CIL | 20110930 | 99991231 | Country | Country | varchar(2) | Country | No Change |",
            "| FRY14Q | CIL | 20260930 | 99991231 | LMBEFD | Leveraged Buyout Flag | Numeric (1,0) | LBO flag | New |",
        ]
        result = spec_diff.diff(self.doc("FRY14Q_2026Q2_CIL_DataDictionary_V1.pdf", before),
                                self.doc("FRY14Q_2026Q3_CIL_DataDictionary_V2.pdf", after))
        self.assertEqual(len(result["tables"]), 1, [table["name"] for table in result["tables"]])
        table = result["tables"][0]
        self.assertEqual(table["kind"], "fields")
        self.assertEqual(table["key"], ["Technical Field Name / MDRM"])
        changes = by_label(result)
        self.assertEqual(changes["TaxStatus"]["category"], "retired")
        self.assertEqual(changes["TaxStatus"]["effective"], "last applies to 2026-06-30")
        self.assertEqual(changes["CustomerID"]["category"], "definition")
        self.assertEqual(changes["LMBEFD"]["category"], "added")
        self.assertEqual(changes["LMBEFD"]["effective"], "applies from 2026-09-30")
        self.assertNotIn("City", changes)
        self.assertIn("LMBEFD", spec_diff.summary(result)["fields"])

    def test_inventory_lists_what_a_new_file_declares(self):
        rows = [
            "| FRY14Q | CIL | 20110930 | 20260630 | TaxStatus | Tax | Numeric (1,0) | Tax | Archived |",
            "| FRY14Q | CIL | 20260930 | 99991231 | LMBEFD | LBO | Numeric (1,0) | LBO | New |",
            "| FRY14Q | CIL | 20110930 | 99991231 | City | City | varchar(73) | City | No Change |",
            "| FRY14Q | CIL | 20110930 | 99991231 | Zip | Zip | varchar(9) | Zip | Modified |",
        ]
        inventory = spec_diff.inventory(self.doc("FRY14Q_2026Q3_CIL_DataDictionary_V2.pdf", rows))
        self.assertEqual(inventory["declared"], {"retired": 1, "new": 1, "updated": 1})
        summary = spec_diff.inventory_summary(inventory)
        self.assertEqual(summary["headline"], "Declared in this file: 1 marked new, 1 marked updated, 1 marked retired")
        self.assertEqual(spec_diff.describe(summary), summary["headline"])


class TestSystematicEdits(unittest.TestCase):
    def test_a_token_swap_across_many_comment_cells_is_one_pattern(self):
        def text(year: str, logic_change: bool = False) -> str:
            rows = ["[SHEET] Rules", "Rules · A: Rule ID · B: Rule Description · C: Comments"]
            for number in range(40):
                logic = f"[{number}] >= 0" + (" (+/- 5)" if logic_change and number == 7 else "")
                rows.append(f"Rules · A: CR-{number:04d} · B: {logic} · C: Existing Rule - Remains Effective Q1 {year}")
            return "\n".join(rows)

        result = spec_diff.diff(text("2026"), text("2027", logic_change=True))
        table = result["tables"][0]
        self.assertEqual(table["patterns"], [{"column": "Comments", "from": "2026", "to": "2027", "rows": 40}])
        self.assertEqual([change["label"] for change in table["changes"]], ["CR-0007"])
        self.assertEqual(table["changes"][0]["category"], "logic")
        self.assertIn("“2026” → “2027” in 40 rows", result["headline"])


SCHEMA = """<xs:schema xmlns:xs="http://www.w3.org/2001/XMLSchema">
  <xs:complexType name="Dollar_x0020_Amount"><xs:sequence><xs:element name="value"><xs:simpleType>
    <xs:restriction base="xs:decimal"><xs:totalDigits value="{digits}"/></xs:restriction></xs:simpleType></xs:element></xs:sequence></xs:complexType>
  <xs:complexType name="PD"><xs:sequence><xs:element name="value"><xs:simpleType><xs:restriction base="xs:string">
    <xs:enumeration value="1401"><xs:annotation><xs:documentation>Key '1401' = Label '{band1}'</xs:documentation></xs:annotation></xs:enumeration>
    <xs:enumeration value="{last}"><xs:annotation><xs:documentation>Key '{last}' = Label 'Not Applicable'</xs:documentation></xs:annotation></xs:enumeration>
  </xs:restriction></xs:simpleType></xs:element></xs:sequence></xs:complexType>
  <xs:element name="BA"><xs:complexType><xs:all>
    <xs:element name="_x0031_001" minOccurs="0"><xs:annotation><xs:documentation>Tier 1 capital ratio</xs:documentation></xs:annotation>
      <xs:complexType><xs:complexContent><xs:extension base="{ratio_type}"/></xs:complexContent></xs:complexType></xs:element>
    {extra}
  </xs:all></xs:complexType></xs:element>
</xs:schema>"""


class TestSchemas(unittest.TestCase):
    def schema(self, **overrides) -> str:
        values = {"digits": "14", "band1": "PD Band 01", "last": "1499", "ratio_type": "Two_x0020_Decimal_x0020_Places", "extra": ""}
        values.update(overrides)
        return SCHEMA.format(**values)

    def test_elements_types_and_code_lists(self):
        after = self.schema(digits="18", band1="PD Band 01 (0.00% to &lt;0.15%)", last="1498", ratio_type="Whole_x0020_Number",
                            extra='<xs:element name="_x0031_002" minOccurs="0"><xs:annotation><xs:documentation>CET1 ratio'
                                  "</xs:documentation></xs:annotation></xs:element>")
        result = spec_diff.diff(self.schema(), after)
        changes = by_label(result)
        self.assertEqual(changes["BA/1001"]["category"], "definition")          # escaped names are decoded
        self.assertEqual(changes["BA/1002"]["category"], "added")
        self.assertEqual(changes["Dollar Amount"]["category"], "definition")
        self.assertEqual(changes["PD=1401"]["category"], "description")
        self.assertEqual(changes["PD=1498"]["note"], "Likely re-keyed from PD=1499")
        self.assertEqual(changes["PD=1499"]["category"], "removed")
        self.assertNotIn("BA", changes)  # a container's type is not its children's type
        kinds = {table["name"]: table["kind"] for table in result["tables"]}
        self.assertEqual(kinds, {"Elements": "concepts", "Types": "types", "Code lists": "codes"})

    def test_the_real_bcar_schema_parses(self):
        paths = sorted(SNAPSHOTS.glob(BCAR_XSD + "-*.txt"))
        if not paths:
            self.skipTest("BCAR XSD snapshot not present")
        text = paths[0].read_text(encoding="utf-8")
        tables = {table.name: table for table in spec_diff.tables_from_text(text)}
        self.assertGreater(len(tables["Elements"].rows), 5000)
        self.assertGreater(len(tables["Code lists"].rows), 400)
        changed = text.replace("Label 'PD Band 24'", "Label 'PD Band 24 (revised)'", 1)
        result = spec_diff.diff(text, changed)
        self.assertEqual([change["label"] for table in result["tables"] for change in table["changes"]], ["PD=1424"])


LINKBASE = """<link:linkbase xmlns:link="http://www.xbrl.org/2003/linkbase" xmlns:xlink="http://www.w3.org/1999/xlink"
  xmlns:generic="http://xbrl.org/2008/generic" xmlns:va="http://xbrl.org/2008/assertion/value"
  xmlns:msg="http://xbrl.org/2010/message" xmlns:sev="http://xbrl.org/2016/assertion-severity">
  <generic:link xlink:type="extended" xlink:role="http://www.xbrl.org/2003/role/link">
    <va:valueAssertion xlink:type="resource" xlink:label="FFIEC_9510_formula" id="FFIEC_9510_formula" test="{test}"/>
    <generic:label xlink:type="resource" xlink:label="lbl_9510" xlink:role="http://www.ffiec.gov/2003/role/ruleType1">Validity</generic:label>
    <msg:message xlink:type="resource" xlink:label="msg_9510" xlink:role="http://www.ffiec.gov/2003/role/ruleMessage1">{message}</msg:message>
    <sev:{severity} xlink:type="resource" xlink:label="sev_9510"/>
    <generic:arc xlink:type="arc" xlink:arcrole="http://xbrl.org/arcrole/2008/element-label" xlink:from="FFIEC_9510_formula" xlink:to="lbl_9510"/>
    <generic:arc xlink:type="arc" xlink:arcrole="http://xbrl.org/arcrole/2010/assertion-unsatisfied-message" xlink:from="FFIEC_9510_formula" xlink:to="msg_9510"/>
    <generic:arc xlink:type="arc" xlink:arcrole="http://xbrl.org/arcrole/2016/assertion-unsatisfied-severity" xlink:from="FFIEC_9510_formula" xlink:to="sev_9510"/>
    <link:loc xlink:type="locator" xlink:href="concepts.xsd#RCON2170" xlink:label="loc_RCON2170"/>
    <link:label xlink:type="resource" xlink:label="lbl_RCON2170" xlink:role="http://www.xbrl.org/2003/role/label" xml:lang="en">{label}</link:label>
    <link:labelArc xlink:type="arc" xlink:arcrole="http://www.xbrl.org/2003/arcrole/concept-label" xlink:from="loc_RCON2170" xlink:to="lbl_RCON2170"/>
  </generic:link>
</link:linkbase>"""


class TestXbrlLinkbases(unittest.TestCase):
    def linkbase(self, **overrides) -> str:
        values = {"test": "$RCON2170 ge 0", "message": "Total assets must not be negative", "severity": "warning", "label": "Total assets"}
        values.update(overrides)
        return LINKBASE.format(**values)

    def test_assertions_messages_severity_and_concept_labels(self):
        result = spec_diff.diff(self.linkbase(), self.linkbase(test="$RCON2170 gt 0", severity="error", label="Total assets (consolidated)"))
        changes = by_label(result)
        rule = changes["9510"]
        self.assertEqual(rule["category"], "severity")
        self.assertEqual(rule["severity"], {"before": "Warning", "after": "Error", "direction": "escalated"})
        self.assertEqual(rule["also"], ["logic"])
        self.assertEqual(changes["RCON2170|label|en"]["category"], "description")
        message_only = spec_diff.diff(self.linkbase(), self.linkbase(message="Total assets must be positive"))
        self.assertEqual(by_label(message_only)["9510"]["category"], "message")


class TestRealOsfiSpecification(unittest.TestCase):
    """OSFI's BCAR technical specification, saved twice in two extraction formats."""

    @classmethod
    def setUpClass(cls):
        paths = sorted(SNAPSHOTS.glob(OSFI_SPEC + "-*.txt"))
        if len(paths) < 2:
            raise unittest.SkipTest("OSFI BCAR snapshots not present")
        texts = [path.read_text(encoding="utf-8") for path in paths]
        cls.structured = next(text for text in texts if text.startswith("[WORKBOOK]"))
        cls.older = next(text for text in texts if not text.startswith("[WORKBOOK]"))

    def test_the_four_spec_tables_are_recognised(self):
        tables = {table.name: (table.kind, len(table.rows)) for table in spec_diff.tables_from_text(self.structured)}
        self.assertEqual(tables["Conventional Rules"][0], "rules")
        self.assertGreater(tables["Conventional Rules"][1], 7000)
        self.assertEqual(tables["Dim Rules"][0], "rules")
        self.assertEqual(tables["Data Measure Definitions"][0], "fields")
        self.assertEqual(tables["Conventional DPAs"][0], "fields")

    def test_the_same_file_in_both_extraction_formats_has_no_structured_change(self):
        self.assertIsNone(spec_diff.diff(self.older, self.structured))

    def previous_version(self) -> str:
        """The current file with this release's annotated changes undone, as the prior release would read."""
        lines = []
        for line in self.structured.splitlines():
            if line.startswith(("Conventional Rules · ", "Dim Rules · ", "Conventional DPAs · ")):
                if re.search(r"C: New( Rule|$)", line):
                    continue
                if "C: Updated Rule" in line:
                    line = re.sub(r" · C: Updated Rule.*", " · C: Existing Rule - Remains Effective Q1 2026", line.replace(" · B: ", " · B: OLD "))
                line = re.sub(r"C: Deprecated( - Q1 2027)?$", "C: Existing", line)
            lines.append(line)
        return "\n".join(lines)

    def test_a_release_reads_as_the_regulator_declares_it(self):
        result = spec_diff.diff(self.previous_version(), self.structured)
        tables = {table["name"]: table for table in result["tables"]}
        self.assertEqual(tables["Conventional Rules"]["counts"], {"added": 58, "logic": 34})
        self.assertEqual(tables["Conventional Rules"]["declared"], {"new": 58, "updated": 34})
        self.assertEqual(tables["Dim Rules"]["counts"], {"added": 4})
        self.assertEqual(tables["Conventional DPAs"]["counts"], {"added": 138, "retired": 56})
        self.assertEqual(result["mismatches"], 0)

    def test_an_unannounced_rule_change_is_flagged(self):
        changed = self.structured.replace("CR-1003-10.01-001 · B: (BCAR Schedule 1) 1003 = (LR) D1502 (+/- 10)",
                                          "CR-1003-10.01-001 · B: (BCAR Schedule 1) 1003 = (LR) D1502 (+/- 20)")
        self.assertNotEqual(changed, self.structured)
        result = spec_diff.diff(self.structured, changed)
        change = by_label(result)["CR-1003-10.01-001"]
        self.assertEqual(change["category"], "logic")
        self.assertEqual(change["mismatch"], "Changed but marked “Existing Rule - Remains Effective Q1 2026”")


WORKBOOK_URL = "https://example.test/edit-checks.xlsx"
WORKBOOK_SOURCE = SourceConfig(
    id="edits", name="Edit checks", regulator="Synthetic Regulator", seed_urls=(WORKBOOK_URL,),
    allowed_prefixes=("https://example.test/",), allowed_domains=("example.test",), allowed_urls=(WORKBOOK_URL,),
)


class TestEngineAndAlerts(unittest.TestCase):
    def test_a_changed_edit_check_workbook_records_its_structured_changes(self):
        with Sandbox(source=WORKBOOK_SOURCE) as sandbox:
            sandbox.scan(lambda *a, **k: ok(edit_workbook(edit_rows()), XLSX))
            changed = sandbox.scan(lambda *a, **k: ok(edit_workbook(current_edit_rows()), XLSX))
            event = next(e for e in changed["report"]["events"] if e["event_type"] == "CHANGED_URL")
            self.assertEqual(event["spec"]["counts"]["added"], 2)
            self.assertEqual(event["spec"]["escalated"], 1)
            evidence = json.loads((sandbox.root / event["evidence_location"]).read_text(encoding="utf-8"))
            self.assertEqual(evidence["spec_diff"]["tables"][0]["key"], ["Edit No", "Version No."])
            entry = next(iter(sandbox.inventory().values()))["change_history"][0]
            self.assertEqual(entry["spec"]["counts"], event["spec"]["counts"])

    def test_a_failing_structured_diff_never_breaks_the_scan(self):
        from unittest.mock import patch

        with Sandbox(source=WORKBOOK_SOURCE) as sandbox:
            sandbox.scan(lambda *a, **k: ok(edit_workbook(edit_rows()), XLSX))
            with patch("regmon.engine.spec_diff.diff", side_effect=RuntimeError("boom")):
                changed = sandbox.scan(lambda *a, **k: ok(edit_workbook(current_edit_rows()), XLSX))
            event = next(e for e in changed["report"]["events"] if e["event_type"] == "CHANGED_URL")
            self.assertNotIn("spec", event)

    def test_alerts_carry_a_spec_changes_line(self):
        from regmon import channels, notify

        result = spec_diff.diff(workbook_text(edit_workbook(edit_rows())), workbook_text(edit_workbook(current_edit_rows())))
        report = {"events": [{"event_id": "e1", "event_type": "CHANGED_URL", "url": WORKBOOK_URL, "spec": spec_diff.summary(result)}],
                  "ai_results": []}
        notifications = notify.build_notifications(report)
        alert = next(item for item in notifications if item["event_id"] == "e1")
        self.assertIn("3 retired", alert["spec_changes"])
        self.assertIn("spec_changes", channels.webhook_payload(alert, "https://example.test")["alert"])


class TestAnalystWorkbook(unittest.TestCase):
    def test_compare_command_writes_a_reviewable_workbook(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            before, after = folder / "FRY14Q_2026Q2_CIL_Edits_v1.xlsx", folder / "FRY14Q_2026Q3_CIL_Edits_v2.xlsx"
            rows = current_edit_rows()
            rows[0][11] = "=FIELD100 > 0"  # rule text that looks like a formula must stay text
            before.write_bytes(edit_workbook(edit_rows()))
            after.write_bytes(edit_workbook(rows))
            output = folder / "compare.xlsx"
            with contextlib.redirect_stdout(io.StringIO()):
                code = spec_diff.main([str(before), str(after), "-o", str(output), "--json", str(folder / "out.json"), "--exit-code"])
            self.assertEqual(code, 2)
            book = load_workbook(output)
            self.assertEqual(book.sheetnames, ["Summary", "Spec Changes", "Cell Changes", "Patterns"])
            sheet = book["Spec Changes"]
            header = [cell.value for cell in sheet[1]]
            self.assertIn("Impact", header)
            self.assertIn("Mismatch", header)
            values = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=2)]
            self.assertTrue(any(row[2] == "112 v1" and row[0] == "Severity changed" for row in values))
            formula_like = [cell for row in sheet.iter_rows(min_row=2) for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")]
            self.assertTrue(formula_like)
            self.assertTrue(all(cell.data_type == "s" for cell in formula_like))
            self.assertEqual(sheet.data_validations.dataValidation[0].formula1, '"Impacted,Not impacted,Needs analysis,Out of scope"')
            outcome = json.loads((folder / "out.json").read_text(encoding="utf-8"))
            self.assertEqual(len(outcome["files"][0]["sha256"]), 64)


if __name__ == "__main__":
    unittest.main()


import shutil
import subprocess

NODE = shutil.which("node")


def run_js(body: str):
    script = "const L = require(" + json.dumps(str(ROOT / "dashboard-lib.js")) + ");\n" + body
    result = subprocess.run([NODE, "-e", script], capture_output=True, text=True, timeout=30, check=True)
    return json.loads(result.stdout)


@unittest.skipUnless(NODE, "node is required for dashboard library tests")
class TestDashboardHalf(unittest.TestCase):
    """The dashboard words, groups and exports spec changes as the Python half does."""

    @classmethod
    def setUpClass(cls):
        cls.result = spec_diff.diff(workbook_text(edit_workbook(edit_rows())), workbook_text(edit_workbook(current_edit_rows())))

    def test_the_alert_line_matches(self):
        summary = spec_diff.summary(self.result)
        inventory = spec_diff.inventory_summary({"declared": {"new": 2, "retired": 1}, "tables": []})
        out = run_js("console.log(JSON.stringify([" + ",".join(f"L.describeSpec({json.dumps(value)})" for value in (summary, inventory, None)) + "]));")
        self.assertEqual(out, [spec_diff.describe(summary), spec_diff.describe(inventory), ""])

    def test_groups_follow_importance_and_filters(self):
        table = self.result["tables"][0]
        out = run_js("const t = " + json.dumps(table) + ";\nconsole.log(JSON.stringify(['all','attention','mismatch'].map(f => "
                     "L.specGroups(t, f).map(g => [g.label, g.changes.length]))));")
        self.assertEqual(out[1], [["New rule", 2], ["New version of existing rule", 1], ["Severity changed", 1], ["Logic changed", 1],
                                  ["Retired", 3], ["Message changed", 1]])
        self.assertEqual(out[0][-1], ["Metadata only", 1])
        self.assertEqual(sum(count for _label, count in out[2]), 3)
        for kind, labels in spec_diff.CATEGORY_LABELS.items():
            js = run_js("console.log(JSON.stringify(" + json.dumps(list(labels)) + ".map(c => L.specCategoryLabel(" + json.dumps(kind) + ", c))));")
            self.assertEqual(js, list(labels.values()), kind)

    def test_csv_has_one_row_per_changed_item(self):
        rows = run_js("console.log(JSON.stringify(L.specCsvRows(" + json.dumps(self.result) + ")));")
        self.assertEqual(len(rows) - 1, sum(len(table["changes"]) for table in self.result["tables"]))
        severity = next(row for row in rows if row[2] == "112 v1")
        self.assertEqual(severity[5], "Quality → Validity (escalated)")
        self.assertEqual(severity[9], "Changed but marked “No change”")
