"""Review fixes and additions for structured spec diffs: column renames, weak
headers, severity scope, in-place version bumps, quiet-row budget, punctuation
in expressions, the release-notes audit, spec dates in the calendar and the
implementation conformance check."""

from __future__ import annotations

import contextlib
import io
import json
import re
import tempfile
import unittest
from datetime import date
from pathlib import Path

from regmon import deadlines, spec_conformance, spec_diff
from regmon.spec_report import write_workbook

ROOT = Path(__file__).resolve().parents[1]
SNAPSHOTS = ROOT / "data" / "snapshots"
OSFI_SPEC = "db3db09fda63b39f8dc40b97b712a5c23bcefd3c94615e1e9c990d45533cc00c"
LETTERS = "ABCDEFGHIJ"


def sheet(header: list[str], rows: list[list[str]], name: str = "Edits") -> str:
    lines = [f"[SHEET] {name}", f"{name} · " + " · ".join(f"{LETTERS[i]}: {h}" for i, h in enumerate(header))]
    for row in rows:
        lines.append(f"{name} · " + " · ".join(f"{LETTERS[i]}: {v}" for i, v in enumerate(row) if v))
    return "\n".join(lines)


def changes(result: dict) -> dict[str, dict]:
    return {change["label"]: change for table in result["tables"] for change in table["changes"]}


class TestColumnsAndRoles(unittest.TestCase):
    def test_a_renamed_logic_column_is_still_compared(self):
        before = sheet(["Edit No", "Version No.", "Edit Logic", "Error Message"], [["100", "1", "X > 0", "bad"], ["101", "1", "Y > 0", "bad"]])
        after = sheet(["Edit No", "Version No.", "Edit Logic (Excel format)", "Error Message"], [["100", "1", "X >= 0", "bad"], ["101", "1", "Y > 0", "bad"]])
        result = spec_diff.diff(before, after)
        table = result["tables"][0]
        self.assertEqual(table["columns_renamed"], [{"before": "Edit Logic", "after": "Edit Logic (Excel format)"}])
        self.assertEqual(table["columns_added"], [])
        self.assertEqual(table["columns_removed"], [])
        self.assertEqual(changes(result)["100 v1"]["category"], "logic")

    def test_a_renamed_unroled_column_in_the_same_place_is_paired(self):
        before = sheet(["Edit No", "Error Message", "Reviewer notes"], [["100", "bad", "a"], ["101", "bad", "b"]])
        after = sheet(["Edit No", "Error Message", "Reviewer notes 2026"], [["100", "bad", "a2"], ["101", "bad", "b"]])
        table = spec_diff.diff(before, after)["tables"][0]
        self.assertEqual(len(table["columns_renamed"]), 1)
        self.assertEqual(table["counts"], {"metadata": 1})

    def test_a_bare_rule_column_does_not_take_the_id_from_rule_id(self):
        table = spec_diff.tables_from_text(sheet(["Rule", "Rule ID", "Message"], [["X > 0", "R1", "bad"]]))[0]
        self.assertEqual(table.roles["id"], 1)

    def test_owner_is_not_a_table_column(self):
        roles = spec_diff._roles(["Rule ID", "Logic", "Owner"])
        self.assertNotIn("table", roles)

    def test_a_level_change_in_a_data_dictionary_is_a_definition_change(self):
        before = sheet(["Field Name", "Level", "Description"], [["F1", "1", "Total"], ["F2", "2", "Sub"]])
        after = sheet(["Field Name", "Level", "Description"], [["F1", "2", "Total"], ["F2", "2", "Sub"]])
        result = spec_diff.diff(before, after)
        self.assertEqual(changes(result)["F1"]["category"], "definition")
        self.assertNotIn("severity", changes(result)["F1"])


class TestExpressionsAndVersions(unittest.TestCase):
    def test_an_operator_change_is_logic_not_cosmetic(self):
        before = sheet(["Edit No", "Edit Logic", "Error Message"], [["100", "X > 0", "Value is bad."], ["101", "Y - 1", "bad"]])
        after = sheet(["Edit No", "Edit Logic", "Error Message"], [["100", "X >= 0", "value is bad"], ["101", "y  -  1", "bad"]])
        found = changes(spec_diff.diff(before, after))
        self.assertEqual(found["100"]["category"], "logic")
        self.assertEqual(found["101"]["category"], "cosmetic")  # case and spacing only

    def test_a_sign_change_in_a_data_type_is_a_definition_change(self):
        before = sheet(["Field Name", "Data Type", "Description"], [["F1", "Number(10,-2)", "Total"]])
        after = sheet(["Field Name", "Data Type", "Description"], [["F1", "Number(10,2)", "Total"]])
        self.assertEqual(changes(spec_diff.diff(before, after))["F1"]["category"], "definition")

    def test_a_version_bump_in_place_is_flagged(self):
        before = sheet(["Edit No", "Version No.", "Edit Logic"], [["100", "1", "X > 0"], ["101", "1", "Z"]])
        after = sheet(["Edit No", "Version No.", "Edit Logic"], [["100", "2", "X > 0"], ["101", "2", "Z > 1"]])
        found = changes(spec_diff.diff(before, after))
        self.assertEqual(found["100 v2"]["category"], "added_version")
        self.assertEqual(found["100 v2"]["version"], {"before": "1", "after": "2"})
        self.assertEqual(found["101 v2"]["category"], "logic")
        self.assertIn("Version 1 → 2", found["101 v2"]["note"])


class TestQuietBudget(unittest.TestCase):
    def test_metadata_rows_cannot_crowd_out_a_logic_change(self):
        header = ["Rule ID", "Logic", "Message", "Last Modified"]
        before = sheet(header, [[f"R{i}", f"logic {i}", "m", "2026-01-01"] for i in range(300)])
        after = sheet(header, [[f"R{i}", f"logic {i}" + (" changed" if i == 299 else ""), "m", f"2026-02-01 x{i}"] for i in range(300)])
        table = spec_diff.diff(before, after)["tables"][0]
        self.assertEqual(table["counts"], {"logic": 1, "metadata": 299})
        self.assertEqual(table["changes"][0]["category"], "logic")
        self.assertEqual(len(table["changes"]), 1 + spec_diff.MAX_QUIET_PER_TABLE)
        self.assertEqual(table["quiet_hidden"], 299 - spec_diff.MAX_QUIET_PER_TABLE)

    def test_over_budget_tables_still_list_their_substantive_rows(self):
        header = ["Rule ID", "Logic", "Last Modified"]
        texts_before, texts_after = [], []
        for number in range(3):
            texts_before.append(sheet(header, [[f"T{number}R{i}", "x", "a"] for i in range(80)], name=f"Sheet{number}"))
            texts_after.append(sheet(header, [[f"T{number}R{i}", "x" if i else "y", f"b{i}"] for i in range(80)], name=f"Sheet{number}"))
        from unittest import mock

        with mock.patch.object(spec_diff, "MAX_TOTAL_CHANGES", 1):
            result = spec_diff.diff("\n".join(texts_before), "\n".join(texts_after))
        self.assertEqual(len(result["tables"]), 3)
        self.assertTrue(all(table["changes"] and table["changes"][0]["category"] == "logic" for table in result["tables"]))
        self.assertEqual([len(table["changes"]) for table in result["tables"]], [1 + spec_diff.MAX_QUIET_PER_TABLE, 1, 1])
        self.assertTrue(result["tables"][1]["truncated"])


def change_log(rows: list[list[str]]) -> str:
    return sheet(["Version", "Publication date", "Details of change", "Worksheet"], rows, name="Version")


RULES_HEADER = ["Rule ID", "Rule Logic", "Status"]


class TestReleaseNotesAudit(unittest.TestCase):
    def files(self, notes: list[list[str]]):
        before = sheet(RULES_HEADER, [["IR-1", "A > 0", "Existing"], ["IR-2", "B > 0", "Existing"], ["IR-3", "C > 0", "Existing"]], name="Rules") + "\n" + \
            change_log([["1", "2026-01-01", "Initial release", "Rules"]])
        after = sheet(RULES_HEADER, [["IR-1", "A >= 0", "Updated"], ["IR-2", "B > 1", "Updated"], ["IR-3", "C > 0", "Deprecated"],
                                     ["IR-4", "D > 0", "New"]], name="Rules") + "\n" + change_log(notes + [["1", "2026-01-01", "Initial release", "Rules"]])
        return before, after

    def test_claims_are_checked_against_the_changes(self):
        before, after = self.files([["2", "2026-09-10", "Added: 1 rule Deprecated: 2 rules Updated: 2 rules", "Rules"]])
        notes = spec_diff.diff(before, after)["release_notes"]
        claims = {claim["action"]: claim for claim in notes["entries"][0]["claims"]}
        self.assertEqual(claims["added"]["status"], "matches")
        self.assertEqual((claims["retired"]["count"], claims["retired"]["detected"], claims["retired"]["status"]), (2, 1, "differs"))
        self.assertEqual(claims["updated"]["status"], "matches")
        self.assertEqual(notes["version"], "2")
        self.assertEqual(notes["differs"], 1)

    def test_a_listed_rule_that_did_not_change_and_an_unlisted_change_are_both_flagged(self):
        before, after = self.files([["2", "2026-09-10", "Updated rules: IR-1 IR-3", "Rules"]])
        result = spec_diff.diff(before, after)
        notes = result["release_notes"]
        self.assertEqual(notes["entries"][0]["ids_listed"], 2)
        self.assertEqual(notes["entries"][0]["ids_not_changed"], ["IR-3"])  # IR-3 was retired, not updated
        self.assertEqual(notes["undocumented"][0]["labels"], ["IR-2"])
        self.assertIn("Release notes v2", result["headline"])

    def test_summed_schedule_counts_are_one_claim(self):
        claims = spec_diff.release_notes(
            change_log([["1", "", "x", ""]]),
            change_log([["2", "", "Added DPA's to the following schedules: -Schedule 10.020 : 16 DPA's -Schedule 10.030 : 12 DPA's", ""],
                        ["1", "", "x", ""]]),
            {"tables": []}, {})["entries"][0]["claims"]
        self.assertEqual(claims, [{"action": "added", "count": 28, "noun": "dpa's", "status": "unchecked"}])

    def test_notes_already_present_in_the_previous_file_are_not_audited_again(self):
        before, after = self.files([])
        self.assertNotIn("release_notes", spec_diff.diff(before, after))

    def test_the_audit_lands_in_the_analyst_workbook(self):
        before, after = self.files([["2", "2026-09-10", "Updated rules: IR-1 IR-3", "Rules"]])
        with tempfile.TemporaryDirectory() as folder:
            from openpyxl import load_workbook

            path = write_workbook({"files": [], "result": spec_diff.diff(before, after)}, Path(folder) / "out.xlsx")
            book = load_workbook(path)
            self.assertIn("Release Notes Audit", book.sheetnames)
            self.assertIn("Columns", book.sheetnames)
            values = [cell.value for row in book["Release Notes Audit"].iter_rows() for cell in row]
            self.assertTrue(any(isinstance(value, str) and value.startswith("updated but not listed: IR-2") for value in values))


class TestRealOsfiReleaseNotes(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        paths = sorted(SNAPSHOTS.glob(OSFI_SPEC + "-*.txt"))
        texts = [path.read_text(encoding="utf-8") for path in paths]
        cls.structured = next((text for text in texts if text.startswith("[WORKBOOK]")), None)
        if cls.structured is None:
            raise unittest.SkipTest("OSFI BCAR snapshot not present")

    def test_the_version_sheet_is_read_and_audited(self):
        lines = []
        for line in self.structured.splitlines():
            if line.startswith("Version · A: 4 · "):
                continue  # the previous release had no v4 notes
            if line.startswith(("Conventional Rules · ", "Dim Rules · ", "Conventional DPAs · ")):
                if re.search(r"C: New( Rule|$)", line):
                    continue
                if "C: Updated Rule" in line:
                    line = re.sub(r" · C: Updated Rule.*", " · C: Existing Rule - Remains Effective Q1 2026", line.replace(" · B: ", " · B: OLD "))
                line = re.sub(r"C: Deprecated( - Q1 2027)?$", "C: Existing", line)
            lines.append(line)
        notes = spec_diff.diff("\n".join(lines), self.structured)["release_notes"]
        self.assertEqual((notes["version"], notes["date"]), ("4", "2026-09-10"))
        by_table = {tuple(entry["tables"]): entry for entry in notes["entries"]}
        conventional = {claim["action"]: claim for claim in by_table[("Conventional Rules",)]["claims"]}
        self.assertEqual((conventional["added"]["count"], conventional["added"]["status"]), (58, "matches"))
        # OSFI's notes say 33 rules were updated; 34 rows of the file are marked "Updated Rule".
        self.assertEqual((conventional["updated"]["count"], conventional["updated"]["detected"]), (33, 34))
        self.assertEqual(by_table[("Dim Rules",)]["ids_listed"], 4)
        dpas = {claim["action"]: claim for claim in by_table[("Conventional DPAs",)]["claims"]}
        self.assertEqual((dpas["added"]["count"], dpas["added"]["status"]), (138, "matches"))


class TestSpecDatesInCalendar(unittest.TestCase):
    def test_changed_rows_with_dates_become_one_entry_per_table_and_date(self):
        evidence = {"spec_diff": {"tables": [{"name": "Edits", "kind": "rules", "changes": [
            {"category": "added", "label": "0101 v2", "effective": "applies from 2027-03-31", "quarter": "Q1 2027"},
            {"category": "logic", "label": "0102 v1", "effective": "applies from 2027-03-31", "quarter": "Q1 2027"},
            {"category": "retired", "label": "0103 v1", "effective": "last applies to 2026-12-31", "quarter": "Q4 2026"},
            {"category": "metadata", "label": "0104 v1", "effective": "applies from 2027-03-31"},
            {"category": "added", "label": "0105 v1", "effective": "effective Q2 2027", "quarter": "Q2 2027"},
            {"category": "added", "label": "0106 v1", "effective": "applies from 2015-03-31"},
        ]}]}}
        hits = {(hit["kind"], hit["date"]): hit for hit in deadlines.spec_dates(evidence, date(2026, 10, 26), "2026-07-28")}
        self.assertEqual(set(hits), {("spec_effective", "2027-03-31"), ("spec_last", "2026-12-31"), ("spec_quarter", "2027-06-30")})
        self.assertEqual(hits[("spec_effective", "2027-03-31")]["quote"], "Edits: 2 changed rows apply from 2027-03-31 — 1 new, 1 logic changed (0101 v2, 0102 v1)")
        self.assertEqual(hits[("spec_last", "2026-12-31")]["basis"], "spec")

    def test_a_first_seen_file_contributes_its_declared_dates(self):
        evidence = {"spec_inventory": {"tables": [{"name": "Rules", "kind": "rules", "items": [
            {"label": "IR-9", "action": "new", "effective": "effective Q1 2027"}]}]}}
        hits = deadlines.spec_dates(evidence, date(2026, 10, 26), "")
        self.assertEqual(hits, [])  # quarter-only text without a computed quarter is left out
        evidence["spec_inventory"]["tables"][0]["items"][0]["effective"] = "applies from 2027-03-31"
        self.assertEqual(deadlines.spec_dates(evidence, date(2026, 10, 26), "")[0]["date"], "2027-03-31")

    def test_the_feed_says_where_a_spec_date_came_from(self):
        calendar = {"entries": [{"id": "x", "date": "2027-03-31", "kind": "spec_effective", "label": "Spec rows apply", "title": "Edits",
                                 "url": "https://example.test/e.xlsx", "event_id": "e1", "detected_at": "2026-10-26T00:00:00Z", "basis": "spec",
                                 "quote": "Edits: 2 changed rows apply from 2027-03-31"}]}
        text = deadlines.to_ics(calendar, "2026-10-26T00:00:00Z", "https://example.test/")
        self.assertIn("row-by-row comparison of the specification file", text.replace("\r\n ", ""))


SPEC = sheet(["Edit Number", "Version", "Edit Type", "Edit Logic", "Effective Start", "Effective End"], [
    ["0100", "1", "Validity", "RCON2170 > 0", "2020-03-31", "9999-12-31"],
    ["0101", "2", "Quality", "RCON2948 >= RCON3210", "2027-03-31", "9999-12-31"],
    ["0101", "1", "Quality", "RCON2948 > RCON3210", "2020-03-31", "2026-12-31"],
    ["0102", "1", "Validity", "RCON3210 <> 0", "2020-03-31", "2026-12-31"],
    ["0103", "1", "Validity", "RCONA000 = 1", "2027-03-31", "9999-12-31"],
], name="Spec")
OURS = sheet(["Rule ID", "Severity", "Rule Logic", "Owner"], [
    ["0100", "Warning", "RCON2170 > 0", "Ana"],
    ["0101", "Warning", "rcon2948 > rcon3210", "Ana"],
    ["0102", "Error", "RCON3210 <> 0", "Bo"],
    ["9001", "Warning", "internal check", "Bo"],
], name="Library")


class TestConformance(unittest.TestCase):
    def test_each_kind_of_gap_is_found(self):
        result = spec_conformance.compare(SPEC, OURS)
        found = {item["label"]: item for item in result["findings"]}
        self.assertEqual(result["counts"], {"retired_in_use": 1, "missing": 1, "differs": 2, "extra": 1})
        self.assertEqual(found["0102 v1"]["finding"], "retired_in_use")
        self.assertEqual(found["0103 v1"]["finding"], "missing")
        self.assertEqual(found["9001"]["finding"], "extra")
        # The live version 2 is compared, not the retired version 1.
        self.assertEqual(found["0101 v2"]["differences"], [{"role": "logic", "column": "Edit Logic", "spec": "RCON2948 >= RCON3210",
                                                           "ours": "rcon2948 > rcon3210"}])
        # "Validity" in the spec ranks as "Error"; our "Warning" relaxes it.
        self.assertEqual(found["0100 v1"]["differences"][0]["role"], "severity")
        self.assertEqual(result["findings"][0]["finding"], "retired_in_use")

    def test_a_matching_implementation_conforms(self):
        ours = sheet(["Rule ID", "Severity", "Rule Logic"], [["0100", "Error", "RCON2170 > 0"], ["0101", "Warning", "RCON2948 >= RCON3210"],
                                                             ["0103", "Error", "RCONA000 = 1"]], name="Library")
        result = spec_conformance.compare(SPEC, ours)
        self.assertEqual(result["counts"], {})
        self.assertTrue(result["headline"].startswith("Conforms"))

    def test_the_cli_writes_a_gaps_workbook(self):
        from openpyxl import Workbook, load_workbook

        def book(path: Path, text: str) -> None:
            workbook = Workbook()
            ws = workbook.active
            for line in text.splitlines()[1:]:
                cells = dict(re.findall(r"([A-Z]): ([^·]*?)(?= · [A-Z]: |$)", line.split(" · ", 1)[1]))
                ws.append([cells.get(letter, "").strip() or None for letter in LETTERS[:6]])
            workbook.save(path)

        with tempfile.TemporaryDirectory() as folder:
            spec_path, our_path, out = Path(folder) / "spec.xlsx", Path(folder) / "ours.xlsx", Path(folder) / "gaps.xlsx"
            book(spec_path, SPEC)
            book(our_path, OURS)
            with contextlib.redirect_stdout(io.StringIO()):
                code = spec_conformance.main([str(spec_path), str(our_path), "-o", str(out), "--json", str(Path(folder) / "g.json"), "--exit-code"])
            self.assertEqual(code, 2)
            gaps = load_workbook(out)["Gaps"]
            self.assertEqual(gaps.cell(row=2, column=1).value, "Retired in spec, still implemented")
            self.assertEqual(json.loads((Path(folder) / "g.json").read_text())["result"]["counts"]["missing"], 1)


if __name__ == "__main__":
    unittest.main()
