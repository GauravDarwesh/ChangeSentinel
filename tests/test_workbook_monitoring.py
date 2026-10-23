"""Excel templates are monitored sheet by sheet: values, formulas, merged
headings, drop-down lists, hidden sheets and named ranges, with the original
file kept so the dashboard can compare two versions cell by cell."""
from __future__ import annotations

import datetime
import io
import json
import unittest
from pathlib import Path

from engine_harness import Sandbox, ok
from regmon import engine
from regmon.config import SourceConfig
from regmon.content import extract_content


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
URL = "https://example.test/template.xlsx"
WORKBOOK_SOURCE = SourceConfig(
    id="workbook", name="Workbook", regulator="Synthetic Regulator", seed_urls=(URL,),
    allowed_prefixes=("https://example.test/",), allowed_domains=("example.test",), allowed_urls=(URL,),
)


def workbook(*, inserted_row=False, formula="=SUM(C2:C4)", validation='"Yes,No"', hidden=True, rename=None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation

    book = Workbook()
    sheet = book.active
    sheet.title = rename or "Schedule HC"
    sheet.append(["Item", "Description", "Amount"])
    rows = [("1", "Cash", 100), ("2", "Securities", 250), ("3", "Loans", 900)]
    if inserted_row:
        rows.insert(1, ("1a", "Cash equivalents", 40))
    for row in rows:
        sheet.append(list(row))
    sheet.cell(row=sheet.max_row + 1, column=1, value="Total")
    sheet.cell(row=sheet.max_row, column=3, value=formula)
    sheet.merge_cells("A10:C10")
    sheet["A10"] = "Notes"
    sheet["E1"] = datetime.date(2026, 12, 31)
    check = DataValidation(type="list", formula1=validation)
    sheet.add_data_validation(check)
    check.add("F2:F5")
    book.create_sheet("Instructions")["A1"] = "Report quarterly"
    if hidden:
        codes = book.create_sheet("Codes")
        codes["A1"] = "HC-01"
        codes.sheet_state = "hidden"
    book.defined_names.add(DefinedName("TotalAmount", attr_text=f"'{sheet.title}'!$C$5"))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


class TestWorkbookText(unittest.TestCase):
    def text(self, body: bytes) -> str:
        text, kind, error = extract_content(URL, XLSX, body)
        self.assertIsNone(error)
        self.assertEqual(kind, "xlsx")
        return text

    def test_every_sheet_and_its_structure_is_kept(self):
        text = self.text(workbook())
        self.assertIn("[WORKBOOK] 3 sheets: Schedule HC, Instructions, Codes (hidden)", text)
        self.assertIn("[SHEET] Schedule HC", text)
        self.assertIn("Schedule HC · A: Item · B: Description · C: Amount · E: 2026-12-31", text)
        self.assertIn("Schedule HC · A: Total · C: =SUM(C2:C4)", text)
        self.assertIn("[MERGED] Schedule HC · A10:C10", text)
        self.assertIn('[VALIDATION] Schedule HC · F2:F5 · list · "Yes,No"', text)
        self.assertIn("[SHEET hidden] Codes", text)
        self.assertIn("Codes · A: HC-01", text)
        self.assertIn("[NAME] TotalAmount = 'Schedule HC'!$C$5", text)

    def test_an_inserted_row_is_one_inserted_line(self):
        before = self.text(workbook()).splitlines()
        after = self.text(workbook(inserted_row=True, formula="=SUM(C2:C5)")).splitlines()
        added = [line for line in after if line not in before]
        removed = [line for line in before if line not in after]
        self.assertEqual(added, ["Schedule HC · A: 1a · B: Cash equivalents · C: 40", "Schedule HC · A: Total · C: =SUM(C2:C5)"])
        self.assertEqual(removed, ["Schedule HC · A: Total · C: =SUM(C2:C4)"])

    def test_formulas_drop_down_lists_and_visibility_are_changes(self):
        base = self.text(workbook())
        self.assertNotEqual(base, self.text(workbook(formula="=SUM(C2:C3)")))
        self.assertNotEqual(base, self.text(workbook(validation='"Yes,No,N/A"')))
        self.assertNotEqual(base, self.text(workbook(hidden=False)))
        self.assertEqual(base, self.text(workbook()))


class TestWorkbookMonitoring(unittest.TestCase):
    def test_the_original_is_kept_and_a_change_is_detected(self):
        with Sandbox(source=WORKBOOK_SOURCE) as sandbox:
            sandbox.scan(lambda *a, **k: ok(workbook(), XLSX))
            record = next(iter(sandbox.inventory().values()))
            self.assertEqual(record["extraction_type"], "xlsx")
            self.assertEqual(record["extractor_version"], engine.WORKBOOK_EXTRACTOR_VERSION)
            self.assertTrue(record["file_snapshot_location"].endswith(".xlsx"))
            self.assertTrue((sandbox.root / record["file_snapshot_location"]).is_file())
            quiet = sandbox.scan(lambda *a, **k: ok(workbook(), XLSX))
            self.assertFalse([e for e in quiet["report"]["events"] if e["event_type"] == "CHANGED_URL"])
            changed = sandbox.scan(lambda *a, **k: ok(workbook(formula="=SUM(C2:C3)"), XLSX))
            events = [e for e in changed["report"]["events"] if e["event_type"] == "CHANGED_URL"]
            self.assertEqual(len(events), 1)
            evidence = json.loads((sandbox.root / events[0]["evidence_location"]).read_text(encoding="utf-8"))
            self.assertTrue(evidence["before"]["file_snapshot_location"].endswith(".xlsx"))
            self.assertTrue(evidence["after"]["file_snapshot_location"].endswith(".xlsx"))
            self.assertNotEqual(evidence["before"]["file_snapshot_location"], evidence["after"]["file_snapshot_location"])
            self.assertIn("=SUM(C2:C3)", evidence.get("diff", ""))

    def test_a_workbook_read_by_the_earlier_extractor_moves_over_quietly(self):
        body = workbook()
        with Sandbox(source=WORKBOOK_SOURCE) as sandbox:
            sandbox.scan(lambda *a, **k: ok(body, XLSX))
            inventory = sandbox.inventory()
            uid, record = next(iter(inventory.items()))
            # Shape of a record written before this change: row text, no version, no original.
            legacy_text = "[SHEET] Schedule HC\nItem | Description | Amount"
            snapshot = Path(sandbox.root / record["snapshot_location"])
            snapshot.write_text(legacy_text, encoding="utf-8")
            import hashlib
            record.update({"extractor_version": None, "file_snapshot_location": None,
                           "normalized_hash": hashlib.sha256(legacy_text.encode("utf-8")).hexdigest()})
            sandbox.write_inventory(inventory)
            migrated = sandbox.scan(lambda *a, **k: ok(body, XLSX))
            types = [e["event_type"] for e in migrated["report"]["events"]]
            self.assertNotIn("CHANGED_URL", types)
            self.assertIn("BASELINE_MIGRATION", types)
            record = sandbox.inventory()[uid]
            self.assertEqual(record["extractor_version"], engine.WORKBOOK_EXTRACTOR_VERSION)
            self.assertTrue(record["file_snapshot_location"])


if __name__ == "__main__":
    unittest.main()
