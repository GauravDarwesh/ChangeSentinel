"""Regression tests for detection-correctness defects (assessment H-01, H-02, H-05,
H-06, H-08, M-04, M-05, M-06, M-12, M-13, M-14, L-01)."""
from __future__ import annotations

import hashlib
import io
import json
import os
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import Mock, patch

from engine_harness import ANALYSIS, Sandbox, html, ok
from regmon.config import SourceConfig
from regmon.content import extract_content, normalize_html
from regmon.fetch import FetchResult
from regmon.notify import build_notifications
from regmon.relevance import triage

OFFICE_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}


def make_docx() -> bytes:
    from docx import Document

    document = Document()
    document.sections[0].header.paragraphs[0].text = "Form FR Y-9C instructions"
    document.add_paragraph("Reporting instructions")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Line item"
    table.cell(0, 1).text = "Threshold"
    table.cell(1, 0).text = "Total assets"
    table.cell(1, 1).text = "USD 10 billion"
    document.add_paragraph("After the table")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Threshold"
    sheet["B1"] = 10
    sheet["C1"] = "=B1*2"  # saved without a cached value
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Q3 changes"
    rows, cols = 2, 2
    table = slide.shapes.add_table(rows, cols, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Schedule"
    table.cell(0, 1).text = "Change"
    table.cell(1, 0).text = "HC-R"
    table.cell(1, 1).text = "8% to 10%"
    group = slide.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(4), Inches(2), Inches(1))
    box.text_frame.text = "Grouped note text"
    slide.notes_slide.notes_text_frame.text = "Speaker note about the revision"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


class TestOfficeExtraction(unittest.TestCase):
    """H-02: Office MIME types contain 'openxmlformats' and were parsed as XML."""

    def test_every_office_format_extracts_under_every_content_type(self):
        bodies = {"docx": make_docx(), "xlsx": make_xlsx(), "pptx": make_pptx()}
        for kind, body in bodies.items():
            for content_type in ("", OFFICE_TYPES[kind], "application/octet-stream"):
                with self.subTest(kind=kind, content_type=content_type):
                    text, extracted_kind, error = extract_content(f"https://example.test/file.{kind}", content_type, body)
                    self.assertIsNone(error)
                    self.assertEqual(extracted_kind, kind)
                    self.assertTrue(text)

    def test_docx_tables_and_headers_are_extracted_in_order(self):
        text, _kind, _error = extract_content("https://example.test/a.docx", OFFICE_TYPES["docx"], make_docx())
        self.assertIn("| Total assets | USD 10 billion |", text)
        self.assertLess(text.index("Reporting instructions"), text.index("Total assets"))
        self.assertLess(text.index("Total assets"), text.index("After the table"))
        self.assertIn("[HEADER] Form FR Y-9C instructions", text)

    def test_pptx_tables_groups_and_notes_are_extracted(self):
        text, _kind, _error = extract_content("https://example.test/a.pptx", OFFICE_TYPES["pptx"], make_pptx())
        self.assertIn("| HC-R | 8% to 10% |", text)
        self.assertIn("Grouped note text", text)
        self.assertIn("[NOTES] Speaker note about the revision", text)

    def test_xlsx_formula_without_cached_value_is_kept(self):
        text, _kind, _error = extract_content("https://example.test/a.xlsx", OFFICE_TYPES["xlsx"], make_xlsx())
        self.assertIn("Threshold | 10 | =B1*2", text)

    def test_zip_with_an_office_member_extracts(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("form.docx", make_docx())
            archive.writestr("readme.txt", "bundle notes")
        text, kind, error = extract_content("https://example.test/bundle.zip", "application/zip", buffer.getvalue())
        self.assertIsNone(error)
        self.assertEqual(kind, "zip")
        self.assertIn("USD 10 billion", text)


class TestHtmlNoiseAndMarkup(unittest.TestCase):
    def test_page_state_classes_on_structural_tags_never_empty_the_page(self):
        """H-05: <body class='has-cookie-banner'> used to remove the whole page."""
        text = normalize_html(
            '<html><body class="page has-cookie-banner"><main class="modal-open"><h1>Rule 12</h1>'
            "<p>Capital ratio raised to 8%.</p></main></body></html>"
        )
        self.assertEqual(text, "Rule 12\nCapital ratio raised to 8%.")

    def test_regulatory_blocks_that_share_words_with_chrome_are_kept(self):
        page = (
            "<html><body><main>"
            '<section class="shareholder-reporting"><p>Shareholder deadline moved to 1 March.</p></section>'
            '<div id="consent-orders"><p>Consent order issued against Bank X.</p></div>'
            '<div class="alert-banner"><p>Effective 1 Jan 2027 the form is withdrawn.</p></div>'
            '<div class="social-security-benefits"><p>Benefit table updated.</p></div>'
            '<div class="share-class-table"><p>Class A shares only.</p></div>'
            "</main></body></html>"
        )
        text = normalize_html(page)
        for phrase in ("Shareholder deadline", "Consent order issued", "form is withdrawn", "Benefit table", "Class A shares"):
            self.assertIn(phrase, text)

    def test_real_site_chrome_is_still_removed(self):
        page = (
            "<html><body><main><p>Guideline text.</p>"
            '<div class="cookie-banner">Accept all cookies</div>'
            '<div id="cookie-consent">We use cookies</div>'
            '<div class="consent-banner">Manage consent</div>'
            '<ul class="social-share"><li>Share on X</li></ul>'
            '<div class="share-buttons">Share</div>'
            '<nav class="breadcrumbs">Home / Rules</nav>'
            '<div class="pagination">Page 2 of 9</div>'
            "</main></body></html>"
        )
        self.assertEqual(normalize_html(page), "Guideline text.")

    def test_html_comments_are_not_page_text(self):
        """H-06: '<!--row-->' leaked into Fed snapshots; rotating comments caused changes."""
        a = normalize_html("<main><!--row--><p>Rule text.</p><!-- generated 2026-10-19T12:00:01Z --></main>")
        b = normalize_html("<main><!--row--><p>Rule text.</p><!-- generated 2026-10-19T18:00:07Z --></main>")
        self.assertEqual(a, "Rule text.")
        self.assertEqual(a, b)

    def test_nested_table_content_is_kept(self):
        """M-04: content inside a table nested in a cell was dropped."""
        text = normalize_html(
            "<main><table><tr><td>Outer label<table><tr><td>Inner threshold 42</td><td>units</td></tr>"
            "<tr><td>Second row</td></tr></table></td><td>B</td></tr></table></main>"
        )
        self.assertIn("Inner threshold 42 ; units / Second row", text)
        self.assertIn("| B |", text)

    def test_declared_charset_is_honoured(self):
        """M-13: windows-1252 pages were decoded as UTF-8 into replacement characters."""
        body = "<html><body><main><p>Caf\xe9 r\xe8gle — € 5</p></main></body></html>".encode("cp1252")
        header_text, _kind, _error = extract_content("https://example.test/p", "text/html; charset=windows-1252", body)
        self.assertEqual(header_text, "Caf\xe9 r\xe8gle — € 5")
        meta_body = b'<html><head><meta charset="windows-1252"></head>' + body[6:]
        meta_text, _kind, _error = extract_content("https://example.test/p", "text/html", meta_body)
        self.assertEqual(meta_text, "Caf\xe9 r\xe8gle — € 5")


class TestRelevanceGate(unittest.TestCase):
    def test_regulatory_subjects_are_not_treated_as_site_utilities(self):
        """M-12: /rules/privacy and /about/regulations were excluded from AI and alerts."""
        self.assertTrue(triage("https://www.regulator.gov/rules/privacy", "Privacy rule amendments")["candidate"])
        self.assertTrue(triage("https://www.regulator.gov/about/regulations", "Regulation text")["candidate"])
        self.assertTrue(triage("https://www.regulator.gov/accessibility/rules")["candidate"])
        self.assertFalse(triage("https://www.regulator.gov/contact")["candidate"])
        self.assertFalse(triage("https://www.regulator.gov/careers/openings")["candidate"])


class TestExtractionDriftGuard(unittest.TestCase):
    """H-01: identical source bytes with different extracted text raised CHANGED_URL."""

    def drifted(self, real):
        def extract(url, content_type, body, **kwargs):
            text, kind, error = real(url, content_type, body, **kwargs)
            return text.replace("Rule text.", "Rule text ."), kind, error
        return extract

    def test_identical_bytes_with_new_text_is_a_quiet_migration(self):
        body = html("<p>Rule text.</p>")
        with Sandbox() as sandbox:
            sandbox.scan(lambda *a, **k: ok(body))
            result = sandbox.scan(lambda *a, **k: ok(body), extract=self.drifted(extract_content))
            events = result["report"]["events"]
            self.assertEqual([event["event_type"] for event in events], ["BASELINE_MIGRATION"])
            self.assertEqual(events[0]["migration_reason"], "extraction_drift")
            self.assertEqual(result["report"]["counts"]["extraction_drift"], 1)
            self.assertEqual(result["analyze_mock"].call_count, 0)
            self.assertEqual(build_notifications(result["report"]), [])

    def test_cloudflare_only_byte_changes_with_new_text_are_also_drift(self):
        def page(ray):
            return html(f"<p>Rule text.</p><script>window.__CF$cv$params={{r:'{ray}',t:'MTc5'}};</script>")
        with Sandbox() as sandbox:
            sandbox.scan(lambda *a, **k: ok(page("aaa")))
            result = sandbox.scan(lambda *a, **k: ok(page("bbb")), extract=self.drifted(extract_content))
            self.assertEqual(result["report"]["events"][0]["migration_reason"], "extraction_drift")

    def test_real_source_change_is_never_hidden(self):
        with Sandbox() as sandbox:
            sandbox.scan(lambda *a, **k: ok(html("<p>Threshold 10 percent.</p>")))
            result = sandbox.scan(lambda *a, **k: ok(html("<p>Threshold 25 percent.</p>")))
            self.assertEqual([event["event_type"] for event in result["report"]["events"]], ["CHANGED_URL"])

    def test_change_after_an_extraction_error_is_still_a_change(self):
        """A baseline recorded during an error has the failed response's raw hash."""
        good, fixed = html("<p>Threshold 10 percent.</p>"), html("<p>Threshold 25 percent.</p>")
        responses = iter([ok(good), ok(fixed), ok(fixed)])
        calls = {"n": 0}

        def flaky_extract(url, content_type, body, **kwargs):
            calls["n"] += 1
            if calls["n"] == 2:
                return None, "html", "HTML normalization failed: synthetic"
            return extract_content(url, content_type, body, **kwargs)

        with Sandbox() as sandbox:
            sandbox.scan(lambda *a, **k: next(responses), extract=flaky_extract)
            errored = sandbox.scan(lambda *a, **k: next(responses), extract=flaky_extract)
            self.assertEqual(errored["report"]["events"][0]["event_type"], "EXTRACTION_ERROR")
            result = sandbox.scan(lambda *a, **k: next(responses), extract=flaky_extract)
            self.assertEqual([event["event_type"] for event in result["report"]["events"]], ["CHANGED_URL"])


class TestDiscoveryAndRemovalGuards(unittest.TestCase):
    def test_empty_complete_discovery_never_removes_the_baseline(self):
        """H-08: an empty COMPLETE discovery reported every URL removed and wiped the inventory."""
        with Sandbox() as sandbox:
            sandbox.scan(lambda *a, **k: ok(html("<p>x</p>")))
            sandbox.discovered = []
            result = sandbox.scan(lambda *a, **k: ok(html("<p>x</p>")))
            self.assertFalse(result["report"]["baseline_update_allowed"])
            self.assertEqual(result["report"]["events"], [])
            self.assertEqual(result["report"]["discovery"]["state"], "FAILED")
            self.assertEqual(len(sandbox.inventory()), 1)

    def mass_removal_source(self):
        urls = tuple(f"https://example.test/page-{index}" for index in range(5))
        return SourceConfig(
            id="mass", name="Mass", regulator="Synthetic", seed_urls=urls,
            allowed_prefixes=("https://example.test/",), allowed_domains=("example.test",), allowed_urls=urls,
        )

    def test_mass_removal_is_held_and_alerted_once(self):
        source = self.mass_removal_source()
        with Sandbox(source=source) as sandbox:
            sandbox.scan(lambda url, **k: ok(html(f"<p>{url}</p>")))
            sandbox.discovered = list(source.seed_urls[:1])
            result = sandbox.scan(lambda url, **k: ok(html(f"<p>{url}</p>")))
            report = result["report"]
            self.assertNotIn("REMOVED_URL", [event["event_type"] for event in report["events"]])
            self.assertEqual(report["removal_guard"]["held_removals"], 4)
            self.assertEqual(len(sandbox.inventory()), 5)
            self.assertTrue(all(record.get("missing_since") for url, record in sandbox.inventory().items()
                                if record["canonical_url"] != source.seed_urls[0]))
            alerts = build_notifications({"removal_guards": [report["removal_guard"]]})
            self.assertEqual([alert["kind"] for alert in alerts], ["REMOVAL_HELD"])
            again = sandbox.scan(lambda url, **k: ok(html(f"<p>{url}</p>")))
            self.assertEqual(again["report"]["removal_guard"]["notification_id"], report["removal_guard"]["notification_id"])

    def test_operator_can_confirm_a_mass_removal(self):
        source = self.mass_removal_source()
        with Sandbox(source=source) as sandbox:
            sandbox.scan(lambda url, **k: ok(html(f"<p>{url}</p>")))
            sandbox.discovered = list(source.seed_urls[:1])
            with patch.dict(os.environ, {"REGMON_ALLOW_MASS_REMOVAL": "mass"}):
                result = sandbox.scan(lambda url, **k: ok(html(f"<p>{url}</p>")))
            removed = [event for event in result["report"]["events"] if event["event_type"] == "REMOVED_URL"]
            self.assertEqual(len(removed), 4)
            self.assertIsNone(result["report"]["removal_guard"])

    def test_document_seed_is_its_own_inventory(self):
        """H-08b: a registered PDF monitor discovered 0 URLs with state COMPLETE forever."""
        from regmon.discovery import discover
        from regmon.register_monitor import source_for_url

        source = SourceConfig.from_dict(source_for_url("https://www.example.gov/files/guidance-2026.pdf", "Guidance"))
        with tempfile.TemporaryDirectory() as temporary:
            self.assertEqual(discover(source, Path(temporary)), ["https://www.example.gov/files/guidance-2026.pdf"])

    def test_disabled_http_discovery_ignores_stale_complete_metadata(self):
        """H-08a: stale COMPLETE metadata skipped the browser path and returned no URLs."""
        from regmon import discovery

        source = SourceConfig(
            id="browser-only", name="b", regulator="r", seed_urls=("https://www.example.com/a",),
            allowed_prefixes=("https://www.example.com/",), allowed_domains=("www.example.com",),
            use_http_discovery=False,
        )
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            (root / "discovery").mkdir()
            (root / "discovery" / "browser-only.json").write_text(json.dumps({"source_id": "browser-only", "state": "COMPLETE"}))
            with patch("regmon.discovery._stealth_discover", return_value=["https://www.example.com/a"]) as stealth:
                self.assertEqual(discovery.discover(source, root), ["https://www.example.com/a"])
            stealth.assert_called_once()
            fresh_metadata = json.loads((root / "discovery" / "browser-only.json").read_text())
            self.assertEqual(fresh_metadata["state"], "DEGRADED")
            self.assertEqual(fresh_metadata["method"], "stealth")

    def test_browser_only_success_can_complete_discovery(self):
        """A successful browser-only crawl must not be downgraded because HTTP is disabled."""
        from regmon.discovery import discover

        source = SourceConfig(
            id="browser-success", name="b", regulator="r",
            seed_urls=("https://www.example.com/a",),
            allowed_prefixes=("https://www.example.com/",),
            allowed_domains=("www.example.com",),
            use_http_discovery=False,
        )
        result = Mock(
            returncode=0,
            stdout="https://www.example.com/a\n",
            stderr="",
        )
        with tempfile.TemporaryDirectory() as temporary:
            with patch("regmon.discovery_browser.subprocess.run", return_value=result):
                urls = discover(source, Path(temporary))
            metadata = json.loads((Path(temporary) / "discovery" / "browser-success.json").read_text())
        self.assertEqual(urls, ["https://www.example.com/a"])
        self.assertEqual(metadata["state"], "COMPLETE")
        self.assertEqual(metadata["method"], "stealth")
        self.assertEqual(metadata["failed_pages"], 0)

    def test_browser_only_partial_crawl_cannot_be_marked_complete(self):
        """Useful URLs from a failing crawler remain visible but never bless the baseline."""
        from regmon.discovery import discover

        source = SourceConfig(
            id="browser-partial", name="b", regulator="r",
            seed_urls=("https://www.example.com/a",),
            allowed_prefixes=("https://www.example.com/",),
            allowed_domains=("www.example.com",),
            use_http_discovery=False,
        )
        result = Mock(
            returncode=2,
            stdout="https://www.example.com/a\n",
            stderr="crawler exited before completing the crawl",
        )
        with tempfile.TemporaryDirectory() as temporary:
            with patch("regmon.discovery_browser.subprocess.run", return_value=result):
                urls = discover(source, Path(temporary))
            metadata = json.loads((Path(temporary) / "discovery" / "browser-partial.json").read_text())
        self.assertEqual(urls, ["https://www.example.com/a"])
        self.assertEqual(metadata["state"], "DEGRADED")
        self.assertEqual(metadata["failed_pages"], 1)
        self.assertEqual(metadata["failed_seeds"], 1)


class TestConditionalFetchAndSnapshots(unittest.TestCase):
    def test_extraction_error_forces_full_refetch_and_survives_a_304(self):
        """M-06: a 304 after an extraction error cleared the error and reported UNCHANGED."""
        from regmon.engine import make_current_item
        from engine_harness import SOURCE

        old = {
            "url_id": "u", "canonical_url": SOURCE.seed_urls[0], "first_seen": "2026-10-19T00:00:00+00:00",
            "last_full_fetch": "2026-10-19T06:00:00+00:00", "content_type": "application/pdf",
            "extraction_type": "pdf", "raw_hash": "bytes", "normalized_hash": "older-good-text",
            "etag": '"e2"', "extraction_error": "PDF extraction failed: damaged xref",
            "error_signature": "EXTRACTION_ERROR|PDF extraction failed|", "error_since": "2026-10-19T06:00:00+00:00",
        }
        not_modified = FetchResult(304, "", None, '"e2"', None, "", b"", True)
        with patch("regmon.engine.DEFAULTS", {"conditional_refresh_interval_hours": 24}), \
             patch("regmon.engine.fetch", return_value=not_modified) as fetch:
            item, _text, _html = make_current_item(old["canonical_url"], old, SOURCE, "2026-10-19T12:00:00+00:00")
        self.assertIsNone(fetch.call_args.kwargs["etag"], "an errored record must be fully re-fetched")
        self.assertEqual(item["extraction_error"], old["extraction_error"])
        self.assertEqual(item["content_type"], "application/pdf")
        self.assertEqual(item["error_since"], old["error_since"])

    def test_cloudflare_token_rotation_reuses_the_saved_html_snapshot(self):
        """M-14/O-07: a new ~100 KB HTML file was committed per page per run."""
        def page(ray):
            return html(f"<p>Rule text.</p><script>window.__CF$cv$params={{r:'{ray}',t:'MTc5'}};</script>")
        with Sandbox() as sandbox:
            for ray in ("aaa", "bbb", "ccc"):
                result = sandbox.scan(lambda *a, ray=ray, **k: ok(page(ray)))
            files = list((sandbox.data / "snapshots" / "html").glob("*.html"))
            self.assertEqual(len(files), 1)
            record = next(iter(sandbox.inventory().values()))
            # html_hash identifies the saved snapshot; raw_hash the latest response.
            self.assertEqual(record["html_hash"], hashlib.sha256(files[0].read_bytes()).hexdigest())
            self.assertEqual(record["raw_hash"], hashlib.sha256(page("ccc")).hexdigest())
            self.assertNotEqual(record["html_hash"], record["raw_hash"])
            self.assertIn("saved HTML snapshot", result["report"]["change_detector"]["html_hash"])
            sandbox.scan(lambda *a, **k: ok(html("<p>Rule text changed.</p>")))
            self.assertEqual(len(list((sandbox.data / "snapshots" / "html").glob("*.html"))), 2)

    def test_evidence_uses_the_configured_diff_budget_and_links_original_files(self):
        """L-01: evidence diffs were cut at a hard-coded 200 lines."""
        before = "\n".join(f"<p>line {index}</p>" for index in range(300))
        after = "\n".join(f"<p>line {index} revised</p>" for index in range(300))
        with Sandbox(defaults={"max_diff_lines": 250}) as sandbox:
            sandbox.scan(lambda *a, **k: ok(html(before)))
            result = sandbox.scan(lambda *a, **k: ok(html(after)))
            evidence = json.loads((sandbox.root / result["report"]["events"][0]["evidence_location"]).read_text())
            self.assertEqual(len(evidence["diff"].splitlines()), 251)
            self.assertIn("file_snapshot_location", evidence["after"])


if __name__ == "__main__":
    unittest.main()


class TestGoldenCorpusCheck(unittest.TestCase):
    """E-01: saved originals must reproduce their extracted baseline in CI."""

    def run_check(self, sandbox):
        import importlib.util

        spec = importlib.util.spec_from_file_location(
            "golden_corpus_check", Path(__file__).resolve().parents[1] / "scripts" / "golden_corpus_check.py"
        )
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        config = sandbox.root / "sources.json"
        config.write_text(json.dumps({"sources": [{
            "id": sandbox.source.id, "name": "s", "regulator": "r", "seed_urls": list(sandbox.source.seed_urls),
            "allowed_prefixes": list(sandbox.source.allowed_prefixes), "allowed_domains": list(sandbox.source.allowed_domains),
            "allowed_urls": list(sandbox.source.allowed_urls),
        }]}), encoding="utf-8")
        return module.check(sandbox.data, config)

    def test_current_baselines_reproduce_and_drift_is_reported(self):
        with Sandbox() as sandbox:
            sandbox.scan(lambda *a, **k: ok(html("<p>Rule text.</p>")))
            self.assertEqual(self.run_check(sandbox)["matched"], 1)
            inventory = sandbox.inventory()
            next(iter(inventory.values()))["normalized_hash"] = "0" * 64
            sandbox.write_inventory(inventory)
            self.assertEqual(len(self.run_check(sandbox)["mismatched"]), 1)

    def test_docling_variance_warns_instead_of_failing(self):
        with Sandbox() as sandbox:
            sandbox.scan(lambda *a, **k: ok(html("<p>Rule text.</p>")))

            def docling_variant(url, content_type, body, _extraction_profile=None, **kwargs):
                if _extraction_profile is not None:
                    _extraction_profile["profile"] = "docling"
                return "slightly different layout output", "html", None

            with patch("regmon.content.extract_content", side_effect=docling_variant):
                import importlib.util

                spec = importlib.util.spec_from_file_location(
                    "golden_corpus_check", Path(__file__).resolve().parents[1] / "scripts" / "golden_corpus_check.py"
                )
                module = importlib.util.module_from_spec(spec)
                spec.loader.exec_module(module)
                config = sandbox.root / "sources.json"
                config.write_text(json.dumps({"sources": [{
                    "id": sandbox.source.id, "name": "s", "regulator": "r", "seed_urls": list(sandbox.source.seed_urls),
                    "allowed_prefixes": list(sandbox.source.allowed_prefixes),
                    "allowed_domains": list(sandbox.source.allowed_domains), "allowed_urls": list(sandbox.source.allowed_urls),
                }]}), encoding="utf-8")
                with patch.object(module, "extract_content", side_effect=docling_variant):
                    results = module.check(sandbox.data, config)
            self.assertEqual(results["mismatched"], [])
            self.assertEqual(len(results["model_variance"]), 1)
