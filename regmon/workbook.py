"""Workbook extraction that keeps an Excel template's structure.

Regulatory templates and taxonomies are workbooks of many sheets where a change
can be a value, a formula, a merged heading, a drop-down list or a named range,
and where the sheet a line belongs to matters as much as the line. This turns a
workbook into stable text for change detection and the redline:

    [WORKBOOK] 3 sheets: Schedule HC, Instructions, Codes (hidden)
    [SHEET] Schedule HC
    Schedule HC · A: Item · B: Description · C: Amount
    Schedule HC · A: Total · C: =SUM(C2:C4) → 1250
    [MERGED] Schedule HC · A10:C10
    [VALIDATION] Schedule HC · F2:F5 · list · "Yes,No"
    [NAME] TotalAmount = 'Schedule HC'!$C$5

Every row line names its sheet and its columns but not its row number, so a row
inserted near the top of a sheet is one inserted line rather than a renumbering
of everything below it. Formulas are kept with their saved result, so a changed
formula is a change even when the number it produces is not. The cell-by-cell
comparison itself is drawn by the dashboard from the saved original files.
"""
from __future__ import annotations

import datetime
import io
import re
import zipfile
from xml.sax.saxutils import unescape

MAX_CELLS = 2_000_000
MAX_PART_BYTES = 64 * 1024 * 1024
PROFILE = "xlsx2"


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime.datetime):
        if value.time() == datetime.time(0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, datetime.time):
        return value.isoformat(timespec="minutes")
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else format(value, ".15g")
    return re.sub(r"\s+", " ", str(value)).strip()


def _attr(tag: str, name: str) -> str:
    match = re.search(r"\b" + re.escape(name) + r'="([^"]*)"', tag)
    return unescape(match.group(1), {"&quot;": '"', "&apos;": "'"}) if match else ""


def _structure(body: bytes) -> tuple[dict[str, str], dict[str, list[str]], dict[str, list[str]]]:
    """Sheet visibility, merged ranges and data-validation rules, read from the package XML.

    openpyxl's streaming reader (needed for large templates) exposes none of these.
    """
    states: dict[str, str] = {}
    merges: dict[str, list[str]] = {}
    validations: dict[str, list[str]] = {}
    with zipfile.ZipFile(io.BytesIO(body)) as archive:
        def read(name: str) -> str:
            try:
                info = archive.getinfo(name)
            except KeyError:
                return ""
            if info.file_size > MAX_PART_BYTES:
                raise ValueError(f"Workbook part {name} is too large to read safely")
            return archive.read(info).decode("utf-8", errors="replace")

        workbook = read("xl/workbook.xml")
        rels = read("xl/_rels/workbook.xml.rels")
        targets = {}
        for tag in re.findall(r"<Relationship\b[^>]*>", rels):
            target = _attr(tag, "Target")
            targets[_attr(tag, "Id")] = target[1:] if target.startswith("/") else "xl/" + target.lstrip("./")
        for tag in re.findall(r"<sheet\b[^>]*>", workbook):
            name = _attr(tag, "name")
            states[name] = _attr(tag, "state") or "visible"
            path = targets.get(_attr(tag, "r:id"), "")
            if "worksheets/" not in path:
                continue
            xml = read(path)
            merges[name] = re.findall(r'<mergeCell\b[^>]*\bref="([^"]+)"', xml)
            rules = []
            for attrs, inner in re.findall(r"<dataValidation\b([^>]*)>([\s\S]*?)</dataValidation>", xml):
                first = re.search(r"<formula1>([\s\S]*?)</formula1>", inner)
                second = re.search(r"<formula2>([\s\S]*?)</formula2>", inner)
                rule = unescape(first.group(1)) if first else ""
                if second:
                    rule += " … " + unescape(second.group(1))
                rules.append(" · ".join(part for part in (_attr(attrs, "sqref"), _attr(attrs, "type") or "any", rule) if part))
            validations[name] = rules
    return states, merges, validations


def extract_workbook(body: bytes):
    """Return (text, "xlsx", error) for an .xlsx/.xlsm workbook."""
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter

        states, merges, validations = _structure(body)
        values = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
        formulas = load_workbook(io.BytesIO(body), read_only=True, data_only=False)
        try:
            names = []
            try:
                defined = formulas.defined_names
                items = defined.items() if hasattr(defined, "items") else [(entry.name, entry) for entry in defined.definedName]
                for name, entry in sorted(items, key=lambda pair: pair[0]):
                    if name.startswith("_xlnm._FilterDatabase"):
                        continue
                    names.append(f"[NAME] {name} = {entry.attr_text}")
            except Exception:
                names = []
            titles = [sheet.title for sheet in values.worksheets]
            labels = [title + ("" if states.get(title, "visible") == "visible" else f" ({states[title]})") for title in titles]
            lines = [f"[WORKBOOK] {len(titles)} sheet{'s' if len(titles) != 1 else ''}: " + ", ".join(labels)]
            cells = 0
            truncated = False
            for sheet, formula_sheet in zip(values.worksheets, formulas.worksheets):
                state = states.get(sheet.title, "visible")
                lines.append(f"[SHEET] {sheet.title}" if state == "visible" else f"[SHEET {state}] {sheet.title}")
                for value_row, formula_row in zip(sheet.iter_rows(values_only=True), formula_sheet.iter_rows(values_only=True)):
                    parts = []
                    for index, (value, formula) in enumerate(zip(value_row, formula_row)):
                        shown = _text(value)
                        if isinstance(formula, str) and formula.startswith("="):
                            shown = formula + (f" → {shown}" if shown else "")
                        elif formula is not None and not shown:
                            shown = _text(formula)
                        if not shown:
                            continue
                        parts.append(f"{get_column_letter(index + 1)}: {shown}")
                        cells += 1
                    if parts:
                        lines.append(f"{sheet.title} · " + " · ".join(parts))
                    if cells >= MAX_CELLS:
                        truncated = True
                        break
                for merged in merges.get(sheet.title, []):
                    lines.append(f"[MERGED] {sheet.title} · {merged}")
                for rule in validations.get(sheet.title, []):
                    lines.append(f"[VALIDATION] {sheet.title} · {rule}")
                if truncated:
                    lines.append(f"[TRUNCATED] Stopped after {MAX_CELLS:,} filled cells")
                    break
            lines.extend(names)
        finally:
            values.close()
            formulas.close()
        from regmon.content import normalize_text

        return normalize_text("\n".join(lines)), "xlsx", None
    except Exception as exc:
        return None, "xlsx", f"{type(exc).__name__}: {exc}"


# ---------------------------------------------------------------------------
# Legacy and look-alike ".xls" workbooks.
#
# Regulators still publish Excel 97-2003 files (BIFF inside an OLE2 container),
# and many ".xls" downloads are really Excel 2003 XML (SpreadsheetML) or an HTML
# table saved with an .xls name. All three become the same structured text as an
# .xlsx above, so the redline, the spec comparison and the dashboard read an .xls
# exactly like a modern workbook. BIFF stores formulas as compiled tokens, so an
# .xls keeps each formula's saved result (the value a reviewer sees), not its text.
# ---------------------------------------------------------------------------

LEGACY_PROFILE = "xls1"
OLE2_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_XLS_VISIBILITY = {0: "visible", 1: "hidden", 2: "veryHidden"}


def _column_letter(index: int) -> str:
    name = ""
    number = index + 1
    while number > 0:
        number, remainder = divmod(number - 1, 26)
        name = chr(65 + remainder) + name
    return name


def _render(sheets: list[dict], names: list[str]) -> str:
    """The line format of extract_workbook, for sheets read some other way.

    Each sheet is {"title", "state", "rows": [[str, ...], ...], "merges": [...]}.
    """
    from regmon.content import normalize_text

    labels = [sheet["title"] + ("" if sheet["state"] == "visible" else f" ({sheet['state']})") for sheet in sheets]
    lines = [f"[WORKBOOK] {len(sheets)} sheet{'s' if len(sheets) != 1 else ''}: " + ", ".join(labels)]
    cells = 0
    for sheet in sheets:
        title, state = sheet["title"], sheet["state"]
        lines.append(f"[SHEET] {title}" if state == "visible" else f"[SHEET {state}] {title}")
        truncated = False
        for row in sheet["rows"]:
            parts = [f"{_column_letter(index)}: {value}" for index, value in enumerate(row) if value]
            if parts:
                lines.append(f"{title} · " + " · ".join(parts))
                cells += len(parts)
            if cells >= MAX_CELLS:
                truncated = True
                break
        for merged in sheet.get("merges", []):
            lines.append(f"[MERGED] {title} · {merged}")
        if truncated:
            lines.append(f"[TRUNCATED] Stopped after {MAX_CELLS:,} filled cells")
            break
    lines.extend(sorted(names))
    return normalize_text("\n".join(lines))


def _biff_sheets(body: bytes) -> tuple[list[dict], list[str]]:
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=body, formatting_info=True, on_demand=True)
    except NotImplementedError:
        book = xlrd.open_workbook(file_contents=body, on_demand=True)
    try:
        sheets = []
        for index in range(book.nsheets):
            sheet = book.sheet_by_index(index)
            rows = []
            for r in range(sheet.nrows):
                row = []
                for c in range(sheet.ncols):
                    cell = sheet.cell(r, c)
                    kind, value = cell.ctype, cell.value
                    if kind == xlrd.XL_CELL_DATE:
                        try:
                            value = xlrd.xldate_as_datetime(value, book.datemode)
                        except Exception:
                            pass
                    elif kind == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif kind == xlrd.XL_CELL_ERROR:
                        value = xlrd.error_text_from_code.get(value, "#ERR")
                    elif kind in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        value = None
                    row.append(_text(value))
                rows.append(row)
            merges = [
                f"{_column_letter(clo)}{rlo + 1}:{_column_letter(chi - 1)}{rhi}"
                for rlo, rhi, clo, chi in sorted(getattr(sheet, "merged_cells", []) or [])
            ]
            sheets.append({
                "title": sheet.name,
                "state": _XLS_VISIBILITY.get(getattr(sheet, "visibility", 0), "visible"),
                "rows": rows,
                "merges": merges,
            })
            book.unload_sheet(index)
        names = []
        for name in getattr(book, "name_obj_list", []) or []:
            label = ("_xlnm." if name.builtin else "") + str(name.name)
            if label.startswith("_xlnm._FilterDatabase") or label == "_FilterDatabase":
                continue
            formula = str(getattr(name, "formula_text", "") or "").strip()
            names.append(f"[NAME] {label} = {formula}" if formula else f"[NAME] {label}")
        return sheets, names
    finally:
        book.release_resources()


_SS_NS = "urn:schemas-microsoft-com:office:spreadsheet"


def _spreadsheetml_sheets(text: str) -> tuple[list[dict], list[str]]:
    """Excel 2003 XML: <Workbook><Worksheet ss:Name><Table><Row><Cell><Data>."""
    import xml.etree.ElementTree as ET

    root = ET.fromstring(text)
    ss = "{" + _SS_NS + "}"
    sheets, names = [], []
    for named in root.iter(ss + "NamedRange"):
        label = named.get(ss + "Name") or ""
        if label and not label.startswith("_FilterDatabase"):
            names.append(f"[NAME] {label} = {named.get(ss + 'RefersTo') or ''}".rstrip(" ="))
    for worksheet in root.iter(ss + "Worksheet"):
        title = worksheet.get(ss + "Name") or f"Sheet{len(sheets) + 1}"
        state = "visible"
        for option in worksheet.iter("{urn:schemas-microsoft-com:office:excel}Visible"):
            state = {"SheetHidden": "hidden", "SheetVeryHidden": "veryHidden"}.get((option.text or "").strip(), state)
        rows: list[list[str]] = []
        merges: list[str] = []
        table = worksheet.find(ss + "Table")
        r = 0
        for row_node in (table.findall(ss + "Row") if table is not None else []):
            index = row_node.get(ss + "Index")
            r = int(index) - 1 if index and index.isdigit() else r
            while len(rows) < r:
                rows.append([])
            row: list[str] = []
            c = 0
            for cell in row_node.findall(ss + "Cell"):
                at = cell.get(ss + "Index")
                c = int(at) - 1 if at and at.isdigit() else c
                data = cell.find(ss + "Data")
                value = "".join(data.itertext()) if data is not None else ""
                if data is not None and data.get(ss + "Type") == "Number":
                    try:
                        value = _text(float(value))
                    except ValueError:
                        value = _text(value)
                elif data is not None and data.get(ss + "Type") == "Boolean":
                    value = "TRUE" if value.strip() == "1" else "FALSE"
                else:
                    value = _text(value)
                if data is not None and data.get(ss + "Type") == "DateTime" and value.endswith("T00:00:00.000"):
                    value = value[:10]
                formula = cell.get(ss + "Formula")
                if formula:
                    value = formula + (f" → {value}" if value else "")
                while len(row) < c:
                    row.append("")
                row.append(value)
                across, down = int(cell.get(ss + "MergeAcross") or 0), int(cell.get(ss + "MergeDown") or 0)
                if across or down:
                    merges.append(f"{_column_letter(c)}{r + 1}:{_column_letter(c + across)}{r + 1 + down}")
                c += 1 + across
            rows.append(row)
            r += 1
        sheets.append({"title": title, "state": state, "rows": rows, "merges": merges})
    if not sheets:
        raise ValueError("Excel 2003 XML file contains no worksheets")
    return sheets, names


def _html_table_sheets(text: str) -> tuple[list[dict], list[str]]:
    """An HTML table saved with an .xls name: one sheet per <table>."""
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(text, "html.parser")
    sheets = []
    for table in soup.find_all("table"):
        if table.find_parent("table") is not None:
            continue
        rows = []
        for tr in table.find_all("tr"):
            if tr.find_parent("table") is not table:
                continue
            rows.append([_text(cell.get_text(" ", strip=True)) for cell in tr.find_all(["td", "th"], recursive=False)])
        sheets.append({"title": f"Table {len(sheets) + 1}", "state": "visible", "rows": rows, "merges": []})
    if not sheets:
        raise ValueError("The .xls file is an HTML page with no table")
    return sheets, []


def legacy_workbook_format(body: bytes) -> str | None:
    """'biff', 'spreadsheetml' or 'html' when the bytes are an .xls-style workbook."""
    if body.startswith(OLE2_SIGNATURE):
        return "biff"
    head = body[:4096].lstrip(b"\xef\xbb\xbf \t\r\n").lower()
    if b"urn:schemas-microsoft-com:office:spreadsheet" in body[:8192] and head.startswith(b"<?xml"):
        return "spreadsheetml"
    if head.startswith((b"<html", b"<!doctype html", b"<table", b"<meta")) and b"<table" in body[:2_000_000].lower():
        return "html"
    return None


def extract_legacy_workbook(body: bytes):
    """Return (text, "xls", error) for an Excel 97-2003, Excel 2003 XML or HTML-table .xls."""
    from regmon.content import decode_text

    try:
        kind = legacy_workbook_format(body)
        if kind == "biff":
            sheets, names = _biff_sheets(body)
        elif kind == "spreadsheetml":
            sheets, names = _spreadsheetml_sheets(decode_text(body))
        elif kind == "html":
            sheets, names = _html_table_sheets(decode_text(body, sniff_html=True))
        else:
            raise ValueError("The file is not an Excel 97-2003 workbook, an Excel 2003 XML file or an HTML table")
        return _render(sheets, names), "xls", None
    except Exception as exc:
        message = str(exc)
        if "workbook is encrypted" in message.lower():
            message = "The workbook is password-protected, so its cells cannot be read"
        elif "can't find workbook in ole2" in message.lower():
            message = "This Office file is a Word or PowerPoint 97-2003 document, not a workbook"
        return None, "xls", f"{type(exc).__name__}: {message}"
