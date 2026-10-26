"""Conformance: does our implementation match the regulator's specification?

``spec_diff`` compares two versions of the *same* file. This compares two
*different* files that describe the same items: the regulator's edit-check
workbook, data dictionary or taxonomy on one side, and our own implementation
(the rule library a reporting tool runs, an internal data dictionary, a
mapping sheet) on the other. Rows are paired by rule ID or item code, so the
two files may have different columns, orders and sheet names.

Findings, most urgent first:

* ``retired_in_use``  we still implement a rule or field the spec has retired
* ``missing``         a live spec item we don't implement
* ``differs``         implemented, but logic, severity, data type or dates disagree
* ``extra``           we implement something the spec doesn't mention

Run offline:

    python -m regmon.spec_conformance FR_Y-9C_edits_2027Q1.xlsx our_edit_library.xlsx -o gaps.xlsx
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from regmon import spec_diff as S

VERSION = 1
FINDINGS = ["retired_in_use", "missing", "differs", "extra"]
FINDING_LABELS = {"retired_in_use": "Retired in spec, still implemented", "missing": "In spec, not implemented",
                  "differs": "Implemented differently", "extra": "Implemented, not in spec"}
# Columns compared for items on both sides. Descriptions and messages are
# wording an implementation legitimately rewrites; these define behaviour.
COMPARED_ROLES = ["logic", "severity", "dtype", "start", "end", "line_item"]
MAX_FINDINGS = 3000


def _pool(tables: list[S.Table], kind: str) -> dict[str, tuple[S.Table, list[str]]]:
    """Item key → (table, row); for rules with several versions, the live (latest, not retired) one."""
    role = "id" if kind == "rules" else "field"
    pool: dict[str, tuple[S.Table, list[str]]] = {}
    for table in tables:
        if table.kind != kind:
            continue
        for row in table.rows:
            key = S._canon(S._cell(table, row, role)).upper()
            if not key:
                continue
            current = pool.get(key)
            if current is None:
                pool[key] = (table, row)
                continue
            old_table, old_row = current
            # Prefer a live row; among equals, the higher version.
            rank_new = (not S._is_retired(table, row), S._version_tuple(S._cell(table, row, "version")))
            rank_old = (not S._is_retired(old_table, old_row), S._version_tuple(S._cell(old_table, old_row, "version")))
            if rank_new > rank_old:
                pool[key] = (table, row)
    return pool


def _differences(spec: tuple[S.Table, list[str]], ours: tuple[S.Table, list[str]]) -> list[dict]:
    (spec_table, spec_row), (our_table, our_row) = spec, ours
    out = []
    # Severity is often worded differently on each side ("Validity" vs "Error"): compare by rank when either uses an edit type.
    spec_severity_role = "severity" if "severity" in spec_table.roles else ("etype" if "etype" in spec_table.roles else "")
    our_severity_role = "severity" if "severity" in our_table.roles else ("etype" if "etype" in our_table.roles else "")
    if spec_severity_role and our_severity_role and "etype" in (spec_severity_role, our_severity_role):
        expected, actual = S._cell(spec_table, spec_row, spec_severity_role), S._cell(our_table, our_row, our_severity_role)
        left, right = S._severity_rank(expected), S._severity_rank(actual)
        if left and right and left != right:
            out.append({"role": "severity", "column": spec_table.header[spec_table.roles[spec_severity_role]], "spec": S._short(expected),
                        "ours": S._short(actual)})
    for role in COMPARED_ROLES:
        if role not in spec_table.roles or role not in our_table.roles:
            continue
        expected, actual = S._cell(spec_table, spec_row, role), S._cell(our_table, our_row, role)
        if not expected and not actual:
            continue
        if role in ("start", "end"):
            left, right = S._as_date(expected), S._as_date(actual)
            if (left and right and left == right) or (S._is_open(expected) and S._is_open(actual)):
                continue
        if role == "severity" and S._severity_rank(expected) and S._severity_rank(expected) == S._severity_rank(actual):
            continue
        if S._same_loosely(expected, actual, role, spec_table):
            continue
        out.append({"role": role, "column": spec_table.header[spec_table.roles[role]], "spec": S._short(expected),
                    "ours": S._short(actual)})
    return out


def compare(spec_text: str | None, our_text: str | None) -> dict | None:
    """Findings between a specification text and an implementation text, or None when either holds no spec table."""
    spec_tables, our_tables = S.tables_from_text(spec_text), S.tables_from_text(our_text)
    if not spec_tables or not our_tables:
        return None
    findings: list[dict] = []
    compared: dict[str, dict] = {}
    for kind in ("rules", "fields", "concepts", "codes"):
        spec_pool, our_pool = _pool(spec_tables, kind), _pool(our_tables, kind)
        if not spec_pool or not our_pool:
            continue
        compared[kind] = {"spec": len(spec_pool), "ours": len(our_pool), "both": len(set(spec_pool) & set(our_pool))}
        for key, (table, row) in spec_pool.items():
            label = S._label(table, row)
            effective, quarter = S._effective(table, row)
            base = {"kind": kind, "key": key, "label": label, "spec_table": table.name}
            if effective:
                base["effective"] = effective
            if quarter:
                base["quarter"] = quarter
            ours = our_pool.get(key)
            retired = S._is_retired(table, row)
            if ours is None:
                if not retired:
                    findings.append({**base, "finding": "missing", "values": S._values(table, row)})
                continue
            base["our_table"] = ours[0].name
            if retired and not S._is_retired(*ours):
                findings.append({**base, "finding": "retired_in_use"})
                continue
            differences = _differences((table, row), ours)
            if differences:
                findings.append({**base, "finding": "differs", "differences": differences})
        for key, (table, row) in our_pool.items():
            if key not in spec_pool:
                findings.append({"kind": kind, "key": key, "label": S._label(table, row), "our_table": table.name, "finding": "extra",
                                 "values": S._values(table, row)})
    if not compared:
        return None
    findings.sort(key=lambda item: (FINDINGS.index(item["finding"]), item["kind"], item["label"]))
    counts = Counter(item["finding"] for item in findings)
    result = {"version": VERSION, "compared": compared, "counts": dict(counts), "findings": findings[:MAX_FINDINGS]}
    if len(findings) > MAX_FINDINGS:
        result["truncated"] = True
    result["headline"] = headline(result)
    return result


def headline(result: dict | None) -> str:
    if not result:
        return ""
    counts = result.get("counts") or {}
    if not counts:
        both = sum(item["both"] for item in result.get("compared", {}).values())
        return f"Conforms: all {both} items in both files agree on logic, severity, type and dates."
    parts = []
    for finding in FINDINGS:
        if counts.get(finding):
            examples = [item["label"] for item in result["findings"] if item["finding"] == finding][:3]
            parts.append(f"{FINDING_LABELS[finding].lower()} {counts[finding]} (" + ", ".join(examples) +
                         (", …" if counts[finding] > len(examples) else "") + ")")
    return "; ".join(parts)


def write_workbook(outcome: dict, path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    from regmon.spec_report import _header, _put

    result = outcome.get("result") or {"counts": {}, "findings": [], "compared": {}, "headline": "No rule or field table recognised in both files."}
    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    _put(summary, 1, 1, "Implementation conformance").font = Font(bold=True, size=14)
    _put(summary, 2, 1, "Generated " + datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC") + " by ChangeSentinel (regmon.spec_conformance)")
    _put(summary, 4, 1, "In short").font = Font(bold=True)
    cell = _put(summary, 5, 1, result.get("headline") or "")
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    summary.merge_cells("A5:F5")
    summary.row_dimensions[5].height = 60
    row = 7
    _header(summary, row, ["File", "Role", "Bytes", "SHA-256"])
    for item in outcome.get("files", []):
        row += 1
        for column, value in enumerate([item["name"], item["role"], item["bytes"], item["sha256"]], start=1):
            _put(summary, row, column, value)
    row += 2
    _header(summary, row, ["Kind", "In spec", "In ours", "In both"] + [FINDING_LABELS[name] for name in FINDINGS])
    for kind, numbers in result.get("compared", {}).items():
        row += 1
        kind_counts = Counter(item["finding"] for item in result.get("findings", []) if item["kind"] == kind)
        for column, value in enumerate([kind, numbers["spec"], numbers["ours"], numbers["both"]] + [kind_counts.get(name) for name in FINDINGS], start=1):
            _put(summary, row, column, value)
    summary.column_dimensions["A"].width = 40
    for letter in "BCDEFGH":
        summary.column_dimensions[letter].width = 18

    gaps = book.create_sheet("Gaps")
    columns = ["Finding", "Kind", "Item", "Spec table", "Our table", "Effective", "Column", "Spec says", "We have", "Owner", "Status", "Comment"]
    _header(gaps, 1, columns)
    fills = {"retired_in_use": "FADBD8", "missing": "FCE4D6", "differs": "FFF2CC", "extra": "F2F2F2"}
    row = 1
    for item in result.get("findings", []):
        differences = item.get("differences") or [None]
        for difference in differences:
            row += 1
            values = [FINDING_LABELS[item["finding"]], item["kind"], item["label"], item.get("spec_table"), item.get("our_table"),
                      item.get("effective"), difference["column"] if difference else None,
                      difference["spec"] if difference else (item.get("values") or {}).get("logic") if item["finding"] == "missing" else None,
                      difference["ours"] if difference else (item.get("values") or {}).get("logic") if item["finding"] == "extra" else None,
                      None, None, None]
            for column, value in enumerate(values, start=1):
                written = _put(gaps, row, column, value)
                if column == 1:
                    written.fill = PatternFill("solid", fgColor=fills[item["finding"]])
                if column in (8, 9):
                    written.alignment = Alignment(wrap_text=True, vertical="top")
    if row > 1:
        validation = DataValidation(type="list", formula1='"Open,In progress,Fixed,Accepted difference,Not applicable"', allow_blank=True)
        validation.add(f"K2:K{row}")
        gaps.add_data_validation(validation)
        gaps.auto_filter.ref = gaps.dimensions
    for index, width in enumerate([30, 10, 26, 24, 24, 24, 18, 60, 60, 16, 18, 40], start=1):
        gaps.column_dimensions[gaps.cell(row=1, column=index).column_letter].width = width
    gaps.freeze_panes = "A2"
    path = Path(path)
    try:
        book.save(path)
    except PermissionError:
        path = path.with_name(f"{path.stem}_{datetime.now().strftime('%Y%m%d-%H%M%S')}{path.suffix}")
        book.save(path)
    return path


def compare_files(spec: Path, ours: Path) -> dict:
    spec_text, our_text = S._read_file(spec), S._read_file(ours)
    files = []
    for role, path in (("specification", spec), ("implementation", ours)):
        body = path.read_bytes()
        files.append({"role": role, "name": path.name, "bytes": len(body), "sha256": hashlib.sha256(body).hexdigest()})
    return {"files": files, "result": compare(spec_text, our_text)}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="python -m regmon.spec_conformance",
        description="Check an implementation (our edit library or data dictionary) against the regulator's specification.")
    parser.add_argument("specification", type=Path, help="The regulator's file: edit checks, data dictionary, XSD or linkbase")
    parser.add_argument("implementation", type=Path, help="Our file describing the same rules or fields")
    parser.add_argument("-o", "--output", type=Path, help="Write a gaps workbook (.xlsx) with Owner, Status and Comment columns")
    parser.add_argument("--json", type=Path, help="Write the full result as JSON")
    parser.add_argument("--exit-code", action="store_true", help="Exit 2 when any gap is found")
    args = parser.parse_args(argv)
    outcome = compare_files(args.specification, args.implementation)
    result = outcome["result"]
    print(f"{args.specification.name} (spec) vs {args.implementation.name} (ours)")
    if not result:
        print("No rule or field table was recognised in both files.")
    else:
        for kind, numbers in result["compared"].items():
            print(f"  {kind}: {numbers['spec']} in spec, {numbers['ours']} in ours, {numbers['both']} in both")
        print(result["headline"])
    if args.json:
        args.json.write_text(json.dumps(outcome, indent=2, ensure_ascii=False), encoding="utf-8")
        print(f"JSON: {args.json}")
    if args.output:
        print(f"Workbook: {write_workbook(outcome, args.output)}")
    return 2 if args.exit_code and result and result.get("counts") else 0


if __name__ == "__main__":
    sys.exit(main())
