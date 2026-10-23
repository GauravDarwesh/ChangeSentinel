"""Content extraction and deterministic normalization."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import re
import stat
import time
import unicodedata
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path, PurePosixPath
from urllib.parse import urlparse

from bs4 import BeautifulSoup
from regmon.pdf_grid import rebuild_page
from regmon.docling_extractor import extract_pdf_with_docling

REMOVE_TAGS = {
    "script", "style", "noscript", "template", "svg", "canvas", "iframe",
    "object", "embed", "nav", "footer", "aside",
}
# Recurring site chrome is recognised from whole id/class *tokens*, split on
# separators ("cookie-banner" -> {"cookie", "banner"}). Substring matching used to
# delete regulatory content: "shareholder-reporting" matched "share",
# "consent-orders" matched "consent", "alert-banner" notices matched "banner", and
# a <body class="has-cookie-banner"> emptied the whole page.
NOISE_PARTS = {
    "cookie", "cookies", "gdpr", "breadcrumb", "breadcrumbs", "pagination",
    "pager", "toolbar", "modal", "popup", "popover", "sharethis", "addthis",
}
# Parts that are chrome only alongside a qualifier ("social-share", "share-buttons",
# "consent-banner"); alone they name real content (social security, share classes,
# consent orders, alert banners).
QUALIFIED_NOISE_PARTS = {
    "social": {"share", "sharing", "links", "media", "icons", "buttons", "follow", "networks", "bar", "nav"},
    "share": {"buttons", "button", "links", "bar", "tools", "toolbar", "icons", "widget", "social", "menu", "this"},
    "sharing": {"buttons", "button", "links", "bar", "tools", "icons", "widget", "social"},
    "consent": {"banner", "bar", "cookie", "cookies", "dialog", "manager", "modal", "popup", "notice", "gdpr", "prompt"},
    "banner": {"cookie", "cookies", "consent", "gdpr"},
}
# Structural containers are never removed by a class/id marker: their classes
# describe page state ("has-cookie-banner", "modal-open"), not their content.
PROTECTED_TAGS = {"html", "body", "main", "article"}
CONTENT_SELECTORS = (
    "main", "[role='main']", "article", "#main-content", "#main",
    ".main-content", ".region-content",
)


def _attrs(tag):
    return getattr(tag, "attrs", None) or {}


def _marker_tokens(attrs) -> list[str]:
    classes = attrs.get("class") or []
    if isinstance(classes, str):
        classes = classes.split()
    return [str(value).lower() for value in [attrs.get("id") or "", *classes] if str(value).strip()]


def _should_remove(tag) -> bool:
    if getattr(tag, "name", None) in PROTECTED_TAGS:
        return False
    attrs = _attrs(tag)
    if str(attrs.get("role") or "").lower() == "main":
        return False
    for token in _marker_tokens(attrs):
        parts = {part for part in re.split(r"[^a-z0-9]+", token) if part}
        if parts & NOISE_PARTS:
            return True
        for part, qualifiers in QUALIFIED_NOISE_PARTS.items():
            if part in parts and parts & qualifiers:
                return True
    return False


def normalize_text(text: str) -> str:
    lines = []
    for line in text.splitlines():
        line = unicodedata.normalize("NFKC", line).replace("\u00a0", " ")
        line = re.sub(r"\s+", " ", line).strip()
        if line:
            lines.append(line)
    return "\n".join(lines)


def _remove_default_noise(root) -> None:
    for tag in list(root.find_all(True)):
        attrs = _attrs(tag)
        if (
            tag.name in REMOVE_TAGS
            or "hidden" in attrs
            or str(attrs.get("aria-hidden") or "").lower() == "true"
            or _should_remove(tag)
        ):
            tag.decompose()


def _remove_ignored_nodes(root, ignore_selectors: tuple[str, ...]) -> None:
    for selector in ignore_selectors:
        for node in reversed(root.select(selector)):
            if node.parent is not None:
                node.decompose()


HTML_BLOCK_TAGS = {
    "address", "article", "blockquote", "dd", "details", "dialog", "div", "dl",
    "dt", "fieldset", "figcaption", "figure", "h1", "h2", "h3", "h4", "h5",
    "h6", "header", "hr", "li", "main", "ol", "p", "pre", "section", "summary",
    "ul",
}


def _html_inline_text(node) -> str:
    """Collect inline HTML text without inserting separators inside sentences."""
    from bs4 import NavigableString
    from bs4.element import PreformattedString

    # Comments, CDATA, processing instructions and doctypes are markup, not page
    # text. They used to leak in ("<!--row-->" became a line of every Fed page)
    # and a rotating "<!-- generated ... -->" comment caused false changes.
    if isinstance(node, PreformattedString):
        return ""
    if isinstance(node, NavigableString):
        value = str(node)
        # HTML source formatting newlines/spaces collapse in normal browser
        # layout. Preserve literal whitespace only inside preformatted blocks.
        if getattr(getattr(node, "parent", None), "name", None) == "pre":
            return value
        return re.sub(r"\s+", " ", value)
    name = getattr(node, "name", None)
    if not name:
        return ""
    if name == "br":
        return "\n"
    if name == "table":
        return "\n" + _html_table_markdown(node) + "\n"
    if name == "li":
        value = "".join(_html_inline_text(child) for child in node.children).strip()
        return "\n- " + value + "\n"
    value = "".join(_html_inline_text(child) for child in node.children)
    if name in HTML_BLOCK_TAGS:
        return "\n" + value.strip() + "\n"
    return value


def _nested_table_text(table) -> str:
    """Flatten a table nested inside a cell, keeping its row/cell boundaries."""
    rows = []
    for row in table.find_all("tr"):
        if row.find_parent("table") is not table:
            continue
        cells = [
            _html_cell_text(cell, escape=False)
            for cell in row.find_all(["th", "td"], recursive=False)
        ]
        cells = [cell for cell in cells if cell]
        if cells:
            rows.append(" ; ".join(cells))
    return " / ".join(rows)


def _html_cell_text(cell, escape: bool = True) -> str:
    # Nested tables (common in older government layouts) used to be dropped
    # entirely; their content is now rendered inline within the outer cell.
    parts = []
    for child in cell.children:
        if getattr(child, "name", None) == "table":
            nested = _nested_table_text(child)
            if nested:
                parts.append(" [" + nested + "] ")
        else:
            parts.append(_html_inline_text(child))
    value = re.sub(r"\s+", " ", "".join(parts)).strip()
    return value.replace("|", r"\|") if escape else value


def _html_table_markdown(table) -> str:
    """Render a DOM table as row/column-preserving Markdown."""
    rows = []
    for row in table.find_all("tr"):
        # Do not accidentally absorb rows from nested tables.
        if row.find_parent("table") is not table:
            continue
        cells = [
            _html_cell_text(cell)
            for cell in row.find_all(["th", "td"], recursive=False)
        ]
        if any(cells):
            rows.append((cells, bool(row.find("th", recursive=False))))
    if not rows:
        return ""
    width = max(len(cells) for cells, _ in rows)
    rendered = []
    for index, (cells, has_header) in enumerate(rows):
        cells = cells + [""] * (width - len(cells))
        rendered.append("| " + " | ".join(cells) + " |")
        if index == 0:
            rendered.append("| " + " | ".join(["---"] * width) + " |")
    return "\n".join(rendered)


def normalize_html(
    html: str,
    include_selector: str | None = None,
    ignore_selectors: tuple[str, ...] | list[str] | None = None,
) -> str:
    """Extract stable, block-aware text while preserving inline sentences and tables.

    A configured inclusion selector that matches nothing raises an explicit
    error instead of silently falling back to whole-page monitoring. Ignore
    selectors apply inside the included fragment, or to the page before
    default main-content detection when no inclusion selector is configured.
    """
    soup = BeautifulSoup(html, "html.parser")
    include_selector = str(include_selector or "").strip()
    ignored = tuple(
        str(selector).strip()
        for selector in (ignore_selectors or ())
        if str(selector).strip()
    )

    if include_selector:
        soup.select(include_selector)
    for selector in ignored:
        soup.select(selector)

    _remove_default_noise(soup)

    if include_selector:
        selected = soup.select(include_selector)
        if not selected:
            raise ValueError(
                f"Configured include selector matched no elements: {include_selector}"
            )
        selected_set = set(selected)
        selected_roots = [
            node for node in selected
            if not any(parent in selected_set for parent in node.parents)
        ]
        fragments = []
        for node in selected_roots:
            fragment = BeautifulSoup(str(node), "html.parser")
            _remove_default_noise(fragment)
            _remove_ignored_nodes(fragment, ignored)
            value = normalize_text(_html_inline_text(fragment))
            if value:
                fragments.append(value)
        return normalize_text("\n\n".join(fragments))

    _remove_ignored_nodes(soup, ignored)
    candidates = []
    for selector in CONTENT_SELECTORS:
        for node in soup.select(selector):
            length = len(node.get_text(" ", strip=True))
            if length:
                candidates.append((length, node))
    content = max(candidates, key=lambda item: item[0])[1] if candidates else (soup.body or soup)
    return normalize_text(_html_inline_text(content))

def _extension(url: str) -> str:
    return Path(urlparse(url).path.lower()).suffix


_CHARSET_PARAM = re.compile(r"charset\s*=\s*[\"']?([A-Za-z0-9._:-]+)", re.IGNORECASE)
_META_CHARSET = re.compile(rb"<meta[^>]+charset\s*=\s*[\"']?([A-Za-z0-9._:-]+)", re.IGNORECASE)


def decode_text(body: bytes, content_type: str = "", *, sniff_html: bool = False) -> str:
    """Decode text using the declared charset (header, then <meta>), else UTF-8.

    Decoding every response as UTF-8 turned windows-1252/ISO-8859-1 pages into
    replacement characters. Unknown or invalid charsets fall back to UTF-8 with
    replacement, so decoding itself never fails.
    """
    candidates = []
    header = _CHARSET_PARAM.search(content_type or "")
    if header:
        candidates.append(header.group(1))
    if sniff_html:
        meta = _META_CHARSET.search(body[:4096])
        if meta:
            candidates.append(meta.group(1).decode("ascii", "ignore"))
    for charset in candidates:
        try:
            return body.decode(charset)
        except (LookupError, UnicodeDecodeError):
            continue
    return body.decode("utf-8-sig", "replace")


OFFICE_MEDIA_TYPES = {
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xlsm": "application/vnd.ms-excel.sheet.macroenabled.12",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}


def _office_kind(ctype: str, suffix: str) -> str | None:
    """Recognise Office Open XML before the generic XML/JSON branches.

    Their MIME types contain "openxmlformats", which the XML branch used to claim,
    so every DOCX/XLSX/PPTX was parsed as XML and failed.
    """
    if "wordprocessingml" in ctype or suffix == ".docx":
        return "docx"
    if "spreadsheetml" in ctype or "ms-excel.sheet.macroenabled" in ctype or suffix in {".xlsx", ".xlsm"}:
        return "xlsx"
    if "presentationml" in ctype or suffix == ".pptx":
        return "pptx"
    return None


MAX_ZIP_MEMBERS = 256
MAX_ZIP_MEMBER_BYTES = 32 * 1024 * 1024
MAX_ZIP_TOTAL_BYTES = 128 * 1024 * 1024
MAX_ZIP_COMPRESSION_RATIO = 1000
ZIP_IGNORED_EXTENSIONS = {
    ".bmp", ".css", ".gif", ".ico", ".jpeg", ".jpg", ".js",
    ".png", ".svg", ".webp",
}
ZIP_MEMBER_MEDIA_TYPES = {
    ".csv": "text/csv",
    ".htm": "text/html",
    ".html": "text/html",
    ".json": "application/json",
    ".xml": "application/xml",
    ".xsd": "application/xml",
    ".txt": "text/plain",
}


def _validated_zip_infos(body: bytes):
    """Return validated archive metadata without extracting members to disk."""
    try:
        archive = zipfile.ZipFile(io.BytesIO(body))
    except (OSError, zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
        raise ValueError(f"Invalid ZIP archive: {type(exc).__name__}: {exc}") from exc

    try:
        infos = archive.infolist()
        if not infos:
            raise ValueError("ZIP archive contains no entries")
        if len(infos) > MAX_ZIP_MEMBERS:
            raise ValueError(f"ZIP archive contains more than {MAX_ZIP_MEMBERS} entries")

        seen_names = set()
        total_uncompressed = 0
        for info in infos:
            name = info.filename.replace(chr(92), "/")
            path = PurePosixPath(name)
            if (
                not name
                or path.is_absolute()
                or name.startswith("/")
                or re.match(r"^[A-Za-z]:", name)
                or ".." in path.parts
            ):
                raise ValueError(f"Unsafe ZIP member path: {info.filename!r}")
            if name in seen_names:
                raise ValueError(f"Duplicate ZIP member path: {name!r}")
            seen_names.add(name)

            mode = (info.external_attr >> 16) & 0o170000
            if mode == stat.S_IFLNK:
                raise ValueError(f"Symbolic links are not supported in ZIP archives: {name}")
            if info.flag_bits & 0x1:
                raise ValueError(f"Encrypted ZIP members are not supported: {name}")
            if info.is_dir():
                continue
            if info.file_size > MAX_ZIP_MEMBER_BYTES:
                raise ValueError(f"ZIP member exceeds {MAX_ZIP_MEMBER_BYTES} bytes: {name}")
            total_uncompressed += info.file_size
            if total_uncompressed > MAX_ZIP_TOTAL_BYTES:
                raise ValueError(f"ZIP contents exceed {MAX_ZIP_TOTAL_BYTES} uncompressed bytes")
            if info.file_size and info.file_size / max(1, info.compress_size) > MAX_ZIP_COMPRESSION_RATIO:
                raise ValueError(f"ZIP member has an unsafe compression ratio: {name}")
        return archive, infos
    except Exception:
        archive.close()
        raise


def _office_open_xml_suffix(names: set[str]) -> str | None:
    if "word/document.xml" in names and "[Content_Types].xml" in names:
        return ".docx"
    if "xl/workbook.xml" in names and "[Content_Types].xml" in names:
        return ".xlsx"
    if "ppt/presentation.xml" in names and "[Content_Types].xml" in names:
        return ".pptx"
    return None


def _extract_zip_bundle(body: bytes, extraction_profile: dict | None = None):
    if extraction_profile is not None:
        extraction_profile["profile"] = "native"
    archive, infos = _validated_zip_infos(body)
    try:
        extracted = []
        member_profiles = []
        skipped_assets = []
        for info in sorted(infos, key=lambda entry: (entry.filename.casefold(), entry.filename)):
            if info.is_dir():
                continue
            name = info.filename.replace(chr(92), "/")
            lowered_name = name.casefold()
            if lowered_name.startswith("__macosx/") or PurePosixPath(name).name.casefold() in {
                ".ds_store", "thumbs.db", "desktop.ini"
            }:
                continue
            suffix = PurePosixPath(name).suffix.lower()
            if suffix in ZIP_IGNORED_EXTENSIONS:
                skipped_assets.append(name)
                continue
            if suffix not in {
                ".csv", ".docx", ".htm", ".html", ".json", ".pdf",
                ".pptx", ".txt", ".xlsm", ".xlsx", ".xml", ".xsd",
            }:
                raise ValueError(f"Unsupported file type inside ZIP archive: {name}")

            member_bytes = archive.read(info)
            media_type = ZIP_MEMBER_MEDIA_TYPES.get(suffix, "")
            member_profile = {}
            text, kind, error = extract_content(
                name, media_type, member_bytes, _extraction_profile=member_profile, _workbook_detail=False
            )
            if error:
                raise ValueError(f"Could not extract ZIP member {name}: {error}")
            member_profiles.append((name, member_profile.get("profile", "native")))
            if text is not None and text.strip():
                extracted.append((name, kind, text))

        if not extracted:
            suffix_note = f" (ignored {len(skipped_assets)} media/style assets)" if skipped_assets else ""
            raise ValueError("ZIP archive contains no supported extractable documents" + suffix_note)

        # A ZIP may contain several PDFs and only some may use Docling.
        # Record a deterministic per-member profile fingerprint; a single
        # aggregate label like "fallback" could otherwise hide one member's
        # Docling<->fallback transition and cause a false change alert.
        if extraction_profile is not None and any(profile != "native" for _, profile in member_profiles):
            profile_material = json.dumps(
                member_profiles, ensure_ascii=False, separators=(",", ":")
            ).encode("utf-8")
            extraction_profile["profile"] = (
                "zipprofile-" + hashlib.sha256(profile_material).hexdigest()[:16]
            )

        parts = []
        for name, kind, text in extracted:
            parts.append(f"[FILE] {name} ({kind})\n{text}")
        return normalize_text("\n\n".join(parts)), "zip", None
    except Exception as exc:
        return None, "zip", f"ZIP extraction failed: {type(exc).__name__}: {exc}"
    finally:
        archive.close()


# Bump this when PDF/HTML normalized output changes. The engine re-fetches these
# records and records a baseline migration (not a regulatory change) if the
# original response bytes are identical.
CONTENT_EXTRACTOR_VERSION = "content-extractor-v7"
LOW_TEXT_PAGE_CHARS = 40
MAX_OCR_PAGES = 1000
OCR_PAGE_TIMEOUT_SECONDS = 15
MAX_OCR_TOTAL_SECONDS = 120


def normalize_pdf_prose(text: str) -> str:
    """Join ordinary PDF line wraps while preserving paragraph and heading boundaries."""
    raw_lines = []
    for source_line in str(text or "").replace("\r", "\n").splitlines():
        line = unicodedata.normalize("NFKC", source_line).replace("\u00a0", " ")
        line = re.sub(r"[ \t]+", " ", line).strip()
        if line:
            raw_lines.append(line)
        elif raw_lines and raw_lines[-1] != "":
            raw_lines.append("")

    paragraphs = []
    current = ""

    def flush():
        nonlocal current
        if current.strip():
            paragraphs.append(current.strip())
        current = ""

    for line in raw_lines:
        if not line:
            flush()
            continue
        if not current:
            current = line
            continue

        # Wrapped words split by a PDF line-end hyphen are usually hyphenation,
        # not a true word boundary. Keep the normalized word as one token.
        if current.endswith("-") and line[:1].islower():
            current = current[:-1] + line
            continue

        # Preserve numbered items, bullets, all-caps section headings and
        # explicit paragraph endings. Other hard line breaks are usually layout
        # wraps, so join them with a space.
        starts_new_item = bool(re.match(r"^(?:[-•▪◦]|\(?\d{1,3}[.)])\s+", line))
        current_is_heading = bool(
            re.fullmatch(r"[A-Z][A-Z0-9 /&()’'–—:.-]{3,}", current)
            or re.match(r"^(?:GENERAL INSTRUCTIONS|LINE ITEM INSTRUCTIONS|QUESTIONS AND ANSWERS)\b", current, re.I)
        )
        if current.endswith((".", "!", "?", ";")) or starts_new_item or current_is_heading:
            flush()
            current = line
        else:
            current += " " + line
    flush()
    return "\n".join(paragraphs)


def _looks_like_fragmented_pdf_page(text: str) -> bool:
    """Fast heuristic to route likely forms/tables to geometry-aware extraction."""
    lines = [re.sub(r"\s+", " ", line).strip() for line in str(text or "").splitlines()]
    lines = [line for line in lines if line]
    if len(lines) < 18:
        return False
    short_lines = sum(len(line) <= 32 for line in lines)
    fragment_lines = sum(
        len(line.split()) <= 3 and len(line) <= 45
        for line in lines
    )
    coded_fields = len(re.findall(r"\b[A-Z]{2,5}\s+[A-Z0-9]{3,5}\b", text))
    return (
        (short_lines / len(lines) >= 0.62 and fragment_lines / len(lines) >= 0.38)
        or coded_fields >= 4
    )


def _pdf_text_quality(text: str) -> float:
    """Rank text candidates by useful content versus excessive fragmentation."""
    value = str(text or "").strip()
    if not value:
        return 0.0
    lines = [line.strip() for line in value.splitlines() if line.strip()]
    if not lines:
        return 0.0
    fragments = sum(
        1 for line in lines
        if len(line) <= 18 and len(line.split()) <= 2
    )
    replacement_chars = value.count("\ufffd")
    return (
        min(len(value), 12000)
        + min(len(lines), 300) * 1.5
        - fragments * 7
        - replacement_chars * 20
    )


def _markdown_table_count(text: str) -> int:
    """Count Markdown table separator rows as evidence of reconstructed grids."""
    count = 0
    for line in str(text or "").splitlines():
        value = line.strip()
        if not (value.startswith("|") and value.endswith("|")):
            continue
        cells = [cell.strip() for cell in value[1:-1].split("|")]
        if len(cells) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            count += 1
    return count


_ROW_NUMBER = re.compile(r"^\(?\d{1,3}[.)]?$")
_SEPARATOR_CELL = re.compile(r":?-{3,}:?")


def _split_markdown_row(line: str) -> list[str] | None:
    value = line.strip()
    if not (value.startswith("|") and value.endswith("|") and len(value) > 1):
        return None
    return [cell.strip() for cell in re.split(r"(?<!\\)\|", value[1:-1])]


def _tidy_markdown_table(rows: list[list[str]]) -> list[str]:
    """Make a Docling form grid readable and diff-stable.

    Blank fill-in rows (empty apart from a row number) are dropped and counted;
    empty and duplicated columns are removed. The grid itself is kept as a table so
    column labels stay inside table cells (scripts/docling_integration_smoke.py).
    """
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    header, body = rows[0], rows[1:]
    kept, blank = [], 0
    for row in body:
        if all(not c or _ROW_NUMBER.match(c) for c in row):
            blank += 1
        else:
            kept.append(row)
    grid = [header] + kept
    columns = [list(col) for col in zip(*grid)]
    columns = [col for col in columns if any(col)]
    deduped: list[list[str]] = []
    for col in columns:
        if not deduped or col != deduped[-1]:
            deduped.append(col)
    columns = deduped
    if not columns:
        return [f"({blank} blank entry rows)"] if blank else []
    grid = [list(row) for row in zip(*columns)]
    out: list[str] = ["| " + " | ".join(grid[0]) + " |", "|" + "|".join(" --- " for _ in grid[0]) + "|"]
    out.extend("| " + " | ".join(row) + " |" for row in grid[1:])
    if blank:
        out.append(f"({blank} blank entry rows)")
    return out


def tidy_docling_markdown(markdown: str) -> str:
    """Clean Docling's Markdown tables; all non-table text is kept unchanged."""
    lines = str(markdown or "").splitlines()
    out: list[str] = []
    index = 0
    while index < len(lines):
        header = _split_markdown_row(lines[index])
        separator = _split_markdown_row(lines[index + 1]) if index + 1 < len(lines) else None
        if header is not None and separator and all(_SEPARATOR_CELL.fullmatch(c) for c in separator):
            rows = [header]
            index += 2
            while index < len(lines):
                row = _split_markdown_row(lines[index])
                if row is None:
                    break
                rows.append(row)
                index += 1
            out.extend(_tidy_markdown_table(rows))
            continue
        out.append(lines[index])
        index += 1
    return "\n".join(out)


GRID_MIN_TOKEN_COVERAGE = 0.9
_LEADER_LINE = re.compile(r"\.{6,}")
_CODE_TOKEN = re.compile(r"\b[A-Z]{4}\s?[A-Z0-9]{4}\b|\b[A-Z]{2}\d{2}\b")
_GLUED_CODES = re.compile(r"[A-Z]{2}\d{2}[A-Z]{2}\d{2}|[A-Z]{4}\d{4}[A-Z]{4}")


def _grid_candidate_page(text: str) -> bool:
    """Report-form pages the fragmentation test misses: dot-leader line items next to
    item codes, or codes the text layer glued together ("HQ78HQ76")."""
    value = str(text or "")
    if _GLUED_CODES.search(value):
        return True
    leaders = sum(1 for line in value.splitlines() if _LEADER_LINE.search(line))
    return leaders >= 3 and len(_CODE_TOKEN.findall(value)) >= 2
_COVERAGE_TOKEN = re.compile(r"[a-z0-9]+(?:[/%._-][a-z0-9]+)*", re.IGNORECASE)


def _token_coverage(reference: str, candidate: str) -> float:
    """Share of the reference's distinct words (3+ characters) present in the candidate."""
    ref = {t.casefold() for t in _COVERAGE_TOKEN.findall(reference or "") if len(t) >= 3}
    if not ref:
        return 1.0
    cand = {t.casefold() for t in _COVERAGE_TOKEN.findall(candidate or "") if len(t) >= 3}
    return len(ref & cand) / len(ref)


def _prefer_docling_candidate(existing: str, candidate: str) -> bool:
    """Use Docling when it yields more table structure without losing most text."""
    if not str(candidate or "").strip():
        return False
    old_tables = _markdown_table_count(existing)
    new_tables = _markdown_table_count(candidate)
    if new_tables == 0:
        return False
    old_quality = _pdf_text_quality(existing)
    new_quality = _pdf_text_quality(candidate)
    if new_tables > old_tables:
        return new_quality >= max(40, int(old_quality * 0.90))
    return new_tables >= old_tables and new_quality > max(40, int(old_quality * 1.12))


def _table_to_markdown(rows) -> str:
    """Render extracted PDF table cells as an explicit grid."""
    cleaned = []
    for row in rows or []:
        cells = [
            re.sub(r"\s+", " ", str(cell or "")).strip().replace("|", r"\|")
            for cell in (row or [])
        ]
        while cells and not cells[-1]:
            cells.pop()
        if any(cells):
            cleaned.append(cells)
    if not cleaned:
        return ""
    width = max(len(row) for row in cleaned)
    output = ["[TABLE]"]
    for row_index, row in enumerate(cleaned):
        row = row + [""] * (width - len(row))
        output.append("| " + " | ".join(row) + " |")
        if row_index == 0:
            output.append("| " + " | ".join(["---"] * width) + " |")
    return "\n".join(output)



def _has_split_word_boundaries(rows) -> bool:
    """Reject layout guesses that split ordinary words across table cells."""
    common_words = {
        "a", "an", "and", "are", "as", "at", "be", "been", "being", "by",
        "can", "could", "did", "do", "does", "for", "from", "had", "has",
        "have", "he", "her", "here", "his", "how", "if", "in", "into", "is",
        "it", "its", "may", "might", "more", "most", "much", "must", "no",
        "not", "of", "off", "on", "one", "or", "other", "our", "out", "over",
        "per", "same", "she", "so", "some", "such", "than", "that", "the",
        "their", "them", "then", "there", "these", "they", "this", "those",
        "through", "to", "too", "under", "up", "us", "was", "we", "were",
        "what", "when", "where", "which", "who", "will", "with", "would",
        "you", "your", "about", "above", "after", "before", "between",
        "below", "during", "within", "without", "because", "while", "each",
        "every", "all", "any", "both", "few", "many", "several", "therefore",
        "form", "forms", "report", "reports", "reported", "reporting", "date",
        "dates", "month", "months", "year", "years", "day", "days", "name",
        "address", "city", "state", "zip", "code", "number", "status", "page",
        "item", "items", "type", "value", "values", "amount", "total", "rate",
        "rates", "period", "periods", "table", "field", "fields", "current",
        "federal", "reserve", "system", "board", "general", "instructions",
        "instruction", "survey", "small", "business", "loan", "loans", "farm",
        "farmers", "lending", "bank", "banks", "please", "read", "completing",
        "effective", "purpose", "include", "exclude", "information", "person",
        "contact", "phone", "telephone", "monday", "tuesday", "wednesday",
        "thursday", "friday", "january", "february", "march", "april", "june",
        "july", "august", "september", "october", "november", "december",
        "approval", "expires",
    }
    continuation_fragments = {
        "tion", "tions", "sion", "sions", "ing", "ings", "ment", "ments",
        "ness", "less", "able", "ible", "ance", "ence", "ity", "ities",
        "ies", "ous", "ive", "ally", "edly", "ed", "ers", "ure", "ary",
        "ory", "ship", "ization", "izations", "ically", "ially", "ative",
    }

    compared_boundaries = 0
    suspicious_boundaries = 0
    for row in rows or []:
        cells = [
            re.sub(r"\s+", " ", str(cell or "")).strip()
            for cell in (row or [])
        ]
        cells = [cell for cell in cells if cell]
        for left, right in zip(cells, cells[1:]):
            left_words = re.findall(r"[A-Za-z]+", left)
            right_words = re.findall(r"[A-Za-z]+", right)
            if not left_words or not right_words:
                continue
            left_word = left_words[-1].casefold()
            right_word = right_words[0].casefold()
            compared_boundaries += 1

            right_is_suffix = right_word in continuation_fragments
            both_look_fragmentary = (
                2 <= len(left_word) <= 7
                and 2 <= len(right_word) <= 7
                and left_word not in common_words
                and right_word not in common_words
                and (len(left_word) <= 5 or len(right_word) <= 5)
            )
            if right_is_suffix or both_look_fragmentary:
                suspicious_boundaries += 1

    return (
        suspicious_boundaries >= 3
        and suspicious_boundaries / max(1, compared_boundaries) >= 0.12
    )


def _table_score(rows) -> int:
    cleaned = [
        [str(cell or "").strip() for cell in (row or [])]
        for row in (rows or [])
    ]
    cleaned = [row for row in cleaned if any(row)]
    if len(cleaned) < 2:
        return 0
    width = max((len(row) for row in cleaned), default=0)
    filled = sum(bool(cell) for row in cleaned for cell in row)
    if width < 2 or filled < 4:
        return 0
    # Reject text-layout guesses that cut ordinary words across cells. The
    # linear extraction path is safer than injecting a malformed table grid.
    if _has_split_word_boundaries(cleaned):
        return 0
    nonempty_ratio = filled / max(1, sum(len(row) for row in cleaned))
    return filled if nonempty_ratio >= 0.20 else 0

def _find_pdf_tables(page):
    """Try ruled-table detection first, then alignment-based detection."""
    candidates = []
    settings_list = (
        {
            "vertical_strategy": "lines",
            "horizontal_strategy": "lines",
            "snap_tolerance": 4,
            "join_tolerance": 4,
            "intersection_tolerance": 5,
        },
        {
            "vertical_strategy": "text",
            "horizontal_strategy": "text",
            "min_words_vertical": 3,
            "min_words_horizontal": 1,
            "intersection_tolerance": 5,
            "text_x_tolerance": 3,
            "text_y_tolerance": 3,
        },
    )
    for settings in settings_list:
        try:
            tables = page.find_tables(table_settings=settings)
        except Exception:
            continue
        chosen = []
        total_score = 0
        for table in tables:
            try:
                rows = table.extract()
                score = _table_score(rows)
                bbox = tuple(float(value) for value in table.bbox)
            except Exception:
                continue
            if not score:
                continue
            if any(
                all(abs(left - right) < 4 for left, right in zip(bbox, existing[0]))
                for existing in chosen
            ):
                continue
            chosen.append((bbox, rows, score))
            total_score += score
        candidates.append((total_score, chosen))
    return max(candidates, key=lambda item: item[0])[1] if candidates else []


def _extract_pdfplumber_page(page, fast_text: str) -> str:
    """Recover page text and table cells while retaining their vertical order."""
    tables = _find_pdf_tables(page)
    if tables:
        tables.sort(key=lambda entry: (entry[0][1], entry[0][0]))

        # The vertical-crop strategy below is safe for page-wide tables, but it
        # can discard prose to the left/right of a narrow table that shares the
        # same vertical band (for example, the FR 2028S prime-rate grid). Keep
        # the complete fast-path text in that case and add a structured table
        # view, accepting stable duplication rather than losing source content.
        narrow_tables = [
            entry for entry in tables
            if float(entry[0][2]) - float(entry[0][0]) < float(page.width) * 0.90
        ]
        if narrow_tables:
            parts = [normalize_pdf_prose(fast_text)]
            parts.extend(_table_to_markdown(rows) for _bbox, rows, _score in tables)
            return "\n\n".join(part for part in parts if part.strip())

        blocks = []
        cursor = 0.0
        for bbox, rows, _score in tables:
            top = max(cursor, max(0.0, min(float(page.height), bbox[1])))
            bottom = max(top, max(0.0, min(float(page.height), bbox[3])))
            if top - cursor >= 2:
                try:
                    region_text = page.crop((0, cursor, page.width, top)).extract_text()
                except Exception:
                    region_text = ""
                region_text = normalize_pdf_prose(region_text or "")
                if region_text:
                    blocks.append(region_text)
            table_text = _table_to_markdown(rows)
            if table_text:
                blocks.append(table_text)
            cursor = max(cursor, bottom)
        if cursor < float(page.height) - 2:
            try:
                region_text = page.crop((0, cursor, page.width, page.height)).extract_text()
            except Exception:
                region_text = ""
            region_text = normalize_pdf_prose(region_text or "")
            if region_text:
                blocks.append(region_text)
        if blocks:
            return "\n\n".join(blocks)

    candidates = [normalize_pdf_prose(fast_text)]
    for kwargs in (
        {"layout": False, "x_tolerance": 2.5, "y_tolerance": 3},
        {"layout": True, "x_tolerance": 2.5, "y_tolerance": 3},
    ):
        try:
            candidate = page.extract_text(**kwargs)
        except Exception:
            continue
        if candidate:
            candidates.append(normalize_pdf_prose(candidate))
    return max(candidates, key=_pdf_text_quality)


def _resources_have_image_xobject(resources, depth: int = 0) -> bool:
    """Find image objects, including images nested inside PDF Form XObjects."""
    if resources is None or depth > 5:
        return False
    resources = resources.get_object()
    xobjects = resources.get("/XObject")
    if xobjects is None:
        return False
    for reference in xobjects.get_object().values():
        obj = reference.get_object()
        subtype = str(obj.get("/Subtype"))
        if subtype == "/Image":
            return True
        if subtype == "/Form" and _resources_have_image_xobject(obj.get("/Resources"), depth + 1):
            return True
    return False


def _page_has_image_xobject(page) -> bool:
    """Recognize image scans, including nested images, on low-text pages."""
    try:
        if _resources_have_image_xobject(page.get("/Resources")):
            return True
        # Some outlined/vector PDFs have no Image XObjects. Because this helper
        # is called only for low-text pages, inspecting their content stream is
        # acceptable; do not decompress every page in an ordinary text PDF.
        contents = page.get_contents()
        if contents is not None:
            return len(contents.get_data()) > 2048
    except Exception:
        return False
    return False


def _ocr_pdf_pages(body: bytes, page_indexes: list[int]) -> dict[int, str]:
    """OCR only pages that appear image-based; native-text pages stay on the fast path."""
    import pytesseract
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(body)
    results = {}
    deadline = time.monotonic() + MAX_OCR_TOTAL_SECONDS
    try:
        for index in page_indexes:
            remaining = deadline - time.monotonic()
            if remaining <= 0.5:
                raise TimeoutError(
                    f"PDF OCR exceeded its total time budget of {MAX_OCR_TOTAL_SECONDS} seconds"
                )
            page = None
            bitmap = None
            try:
                page = document[index]
                bitmap = page.render(scale=1.8)
                image = bitmap.to_pil()
                try:
                    text = pytesseract.image_to_string(
                        image,
                        config="--psm 3",
                        timeout=min(OCR_PAGE_TIMEOUT_SECONDS, remaining),
                    )
                finally:
                    image.close()
                # OCR engines already return geometric line groups; joining them as if
                # they were paragraph-wrapped digital text can collapse form
                # fields and table rows into false sentences.
                results[index] = normalize_text(text or "")
            finally:
                if bitmap is not None:
                    try:
                        bitmap.close()
                    except Exception:
                        pass
                if page is not None:
                    try:
                        page.close()
                    except Exception:
                        pass
    finally:
        document.close()
    return results


def _extract_pdf(
    body: bytes, extraction_profile: dict | None = None
) -> tuple[str | None, str | None]:
    if extraction_profile is not None:
        extraction_profile["profile"] = "native"
    """Fast-first PDF extraction with layout/table and selective OCR fallbacks."""
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(body), strict=False)
        if reader.is_encrypted:
            decrypted = reader.decrypt("")
            if not decrypted:
                return None, "PDF is encrypted and requires a password; refusing to replace the trusted baseline."
        pages = list(reader.pages)
        if not pages:
            return None, "PDF contains no readable pages."

        page_texts = []
        raw_page_texts = []
        low_text_indexes = []
        for index, page in enumerate(pages):
            try:
                raw_text = page.extract_text() or ""
            except Exception:
                raw_text = ""
            raw_page_texts.append(raw_text)
            page_texts.append(normalize_pdf_prose(raw_text))
            is_low_text = len((raw_text or "").strip()) < LOW_TEXT_PAGE_CHARS
            # Keep the common digital-PDF path cheap: only inspect resources and
            # content streams on pages that actually need OCR triage.
            has_image = _page_has_image_xobject(page) if is_low_text else False
            if is_low_text and has_image:
                low_text_indexes.append(index)

        # Run geometry-aware parsing only on pages whose fast text is highly
        # fragmented (common for forms, multi-column PDFs and table-heavy pages).
        suspicious_indexes = [
            index for index, raw in enumerate(raw_page_texts)
            if _looks_like_fragmented_pdf_page(raw)
        ]
        grid_pages = 0
        if suspicious_indexes or any(_grid_candidate_page(raw) for raw in raw_page_texts):
            try:
                import pdfplumber

                with pdfplumber.open(io.BytesIO(body)) as document:
                    grid_only = [i for i, raw in enumerate(raw_page_texts) if i not in set(suspicious_indexes) and _grid_candidate_page(raw)]
                    for index in sorted(set(suspicious_indexes) | set(grid_only)):
                        if index >= len(document.pages):
                            continue
                        # Ruled tables are rebuilt from their drawn lines first (exact cells,
                        # rotated headers, report-form fields); kept only when they retain
                        # the page's words.
                        try:
                            grid_text = rebuild_page(document.pages[index], normalize_pdf_prose)
                        except Exception as exc:
                            grid_text = None
                            logging.getLogger(__name__).warning(
                                "Ruled-grid reconstruction failed on page %s: %s: %s", index + 1, type(exc).__name__, exc,
                            )
                        if grid_text and _token_coverage(raw_page_texts[index], grid_text) >= GRID_MIN_TOKEN_COVERAGE:
                            page_texts[index] = grid_text
                            grid_pages += 1
                            continue
                        if index not in suspicious_indexes:
                            continue  # line-item pages without a usable grid keep their text layer
                        try:
                            candidate = _extract_pdfplumber_page(document.pages[index], page_texts[index])
                            # Successful table reconstruction wins even if it is
                            # shorter than the linear text stream: the cell grid
                            # is the useful semantic structure, not text length.
                            if "[TABLE]" in candidate or _pdf_text_quality(candidate) > _pdf_text_quality(page_texts[index]) * 1.03:
                                page_texts[index] = candidate
                        except Exception as exc:
                            # Keep the fast extraction if a malformed page defeats
                            # table reconstruction. OCR/quality logic still applies.
                            logging.getLogger(__name__).warning(
                                "pdfplumber table reconstruction failed on page %s: %s: %s",
                                index + 1, type(exc).__name__, exc,
                            )
            except Exception as exc:
                logging.getLogger(__name__).warning(
                    "pdfplumber could not open the PDF for layout recovery: %s: %s", type(exc).__name__, exc
                )

        if low_text_indexes:
            if len(low_text_indexes) > MAX_OCR_PAGES:
                return None, (
                    f"PDF has {len(low_text_indexes)} image-only/low-text pages; "
                    f"OCR safety limit is {MAX_OCR_PAGES}. Increase MAX_OCR_PAGES for this workload."
                )
            try:
                ocr_results = _ocr_pdf_pages(body, low_text_indexes)
            except Exception as exc:
                return None, (
                    "PDF contains image-based pages with no usable text layer, but the OCR fallback failed: "
                    f"{type(exc).__name__}: {exc}"
                )
            for index, ocr_text in ocr_results.items():
                if ocr_text.strip():
                    # Preserve the PDFium-rendered/OCR result only when it yields
                    # real text; otherwise retain an existing sparse text stream.
                    if _pdf_text_quality(ocr_text) > _pdf_text_quality(page_texts[index]):
                        page_texts[index] = ocr_text
                # Individual image pages may be decorative covers or blank
                # scanned inserts. Ignore empty OCR for that page; fail safely
                # only if the document as a whole contains no usable text.

        # Keep the established page-by-page extraction as the safe default. For
        # difficult PDFs only, Docling can replace it when it demonstrably recovers
        # structured tables without discarding most of the extracted text.
        text = normalize_text("\n\n".join(page_texts))
        if grid_pages:
            # Ruled tables were rebuilt exactly; a layout model would only re-guess them.
            if extraction_profile is not None:
                extraction_profile["profile"] = "grid"
        elif suspicious_indexes:
            docling_text, docling_error = extract_pdf_with_docling(
                body, expected_pages=len(pages)
            )
            if docling_text and _prefer_docling_candidate(text, docling_text):
                text = normalize_text(tidy_docling_markdown(docling_text))
                if extraction_profile is not None:
                    extraction_profile["profile"] = "docling"
            else:
                if extraction_profile is not None:
                    extraction_profile["profile"] = "fallback"
                if docling_error and docling_error.startswith(
                    "Docling intentionally skipped for a previously time-budget-exhausted PDF"
                ):
                    logging.getLogger(__name__).info("%s", docling_error)
                elif docling_error:
                    logging.getLogger(__name__).warning(
                        "Docling fallback unavailable for SHA-256 %s: %s",
                        hashlib.sha256(body).hexdigest(),
                        docling_error,
                    )
                elif docling_text:
                    logging.getLogger(__name__).info(
                        "Docling candidate rejected by conservative content-quality gate "
                    "(baseline score=%s, Docling score=%s, baseline tables=%s, "
                    "Docling tables=%s)",
                    _pdf_text_quality(text),
                    _pdf_text_quality(docling_text),
                    _markdown_table_count(text),
                    _markdown_table_count(docling_text),
                )
        if not text.strip():
            return None, (
                "PDF produced no usable text. It may contain outlines, damaged text, or scans that OCR could not read."
            )
        return text, None
    except Exception as exc:
        return None, f"PDF extraction failed: {type(exc).__name__}: {exc}"


def extract_content(
    url: str,
    content_type: str,
    body: bytes,
    include_selector: str | None = None,
    ignore_selectors: tuple[str, ...] | list[str] | None = None,
    _extraction_profile: dict | None = None,
    _workbook_detail: bool = True,
):
    if _extraction_profile is not None:
        _extraction_profile["profile"] = "native"
    ctype, suffix = (content_type or "").lower(), _extension(url)
    zip_signature = body.startswith(b"PK")
    if "zip" in ctype or suffix == ".zip" or zip_signature:
        try:
            archive, infos = _validated_zip_infos(body)
            try:
                names = {info.filename.replace(chr(92), "/") for info in infos}
                office_suffix = _office_open_xml_suffix(names)
            finally:
                archive.close()
        except Exception as exc:
            return None, "zip", f"ZIP validation failed: {type(exc).__name__}: {exc}"
        if office_suffix is not None:
            suffix = office_suffix
            ctype = {
                ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            }[office_suffix]
        else:
            return _extract_zip_bundle(body, _extraction_profile)

    # Route binary formats before decoding bytes as UTF-8. Decoding a large
    # PDF into a replacement-filled string wastes CPU and can double peak memory.
    if "pdf" in ctype or suffix == ".pdf" or body.startswith(b"%PDF-"):
        if _extraction_profile is None:
            text, error = _extract_pdf(body)
        else:
            text, error = _extract_pdf(body, _extraction_profile)
        return text, "pdf", error

    office = _office_kind(ctype, suffix)
    if office == "docx":
        return _extract_docx(body)
    if office == "xlsx":
        # A workbook monitored on its own keeps its structure (sheets, cell columns,
        # formulas, merges, drop-down lists, names). Workbooks inside a ZIP bundle keep
        # the earlier row text, so saved bundle baselines still reproduce exactly.
        if not _workbook_detail:
            return _extract_xlsx(body)
        from regmon.workbook import PROFILE, extract_workbook

        if _extraction_profile is not None:
            _extraction_profile["profile"] = PROFILE
        return extract_workbook(body)
    if office == "pptx":
        return _extract_pptx(body)

    if "json" in ctype or suffix == ".json":
        try:
            payload = json.loads(decode_text(body, ctype))
            normalized_json = json.dumps(payload, ensure_ascii=False, sort_keys=True, indent=2)
            return normalized_json, "json", None
        except (json.JSONDecodeError, UnicodeError) as exc:
            return None, "json", f"{type(exc).__name__}: {exc}"

    if "xhtml" in ctype:
        ctype = ctype.replace("xhtml+xml", "html")
    if "xml" in ctype or suffix in {".xml", ".xsd"}:
        try:
            canonical_xml = ET.canonicalize(xml_data=decode_text(body, ctype))
            root = ET.fromstring(canonical_xml)
            for node in root.iter():
                if node.text is not None and not node.text.strip():
                    node.text = None
                if node.tail is not None and not node.tail.strip():
                    node.tail = None
            ET.indent(root, space="  ")
            return ET.tostring(root, encoding="unicode"), "xml", None
        except (ET.ParseError, ValueError) as exc:
            return None, "xml", f"{type(exc).__name__}: {exc}"

    if "html" in ctype or suffix in {".html", ".htm"} or (
        not ctype and suffix == "" and not body.startswith(b"%PDF-")
    ):
        try:
            return normalize_html(
                decode_text(body, ctype, sniff_html=True),
                include_selector=include_selector,
                ignore_selectors=ignore_selectors,
            ), "html", None
        except Exception as exc:
            return None, "html", f"HTML normalization failed: {type(exc).__name__}: {exc}"

    if "csv" in ctype or suffix == ".csv":
        try:
            rows = [
                " | ".join(cell.replace("|", r"\|") for cell in row)
                for row in csv.reader(io.StringIO(decode_text(body, ctype)))
            ]
            return normalize_text("\n".join(rows)), "csv", None
        except Exception as exc:
            return None, "csv", f"{type(exc).__name__}: {exc}"
    if ctype.startswith("text/") or suffix == ".txt":
        return normalize_text(decode_text(body, ctype)), "text", None
    return None, "binary", "Unsupported content type for text extraction"


def _docx_table_rows(table) -> list[list[str]]:
    rows = []
    for row in table.rows:
        cells, seen = [], set()
        for cell in row.cells:
            # Horizontally merged cells are returned once per grid column.
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            cells.append(cell.text)
        rows.append(cells)
    return rows


def _extract_docx(body: bytes):
    """Body paragraphs AND tables in document order, plus headers and footers."""
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = Document(io.BytesIO(body))
        parts = []
        for child in document.element.body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                parts.append(Paragraph(child, document).text)
            elif tag == "tbl":
                parts.append(_table_to_markdown(_docx_table_rows(Table(child, document))))
        seen_margins = set()
        for section in document.sections:
            for label, part in (("HEADER", section.header), ("FOOTER", section.footer)):
                text = "\n".join(paragraph.text for paragraph in part.paragraphs).strip()
                if text and (label, text) not in seen_margins:
                    seen_margins.add((label, text))
                    parts.append(f"[{label}] {text}")
        return normalize_text("\n".join(parts)), "docx", None
    except Exception as exc:
        return None, "docx", f"{type(exc).__name__}: {exc}"


def _extract_xlsx(body: bytes):
    """Cached cell values; a formula is shown when no cached value was saved."""
    try:
        from openpyxl import load_workbook

        values = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
        formulas = load_workbook(io.BytesIO(body), read_only=True, data_only=False)
        rows = []
        for sheet, formula_sheet in zip(values.worksheets, formulas.worksheets):
            rows.append(f"[SHEET] {sheet.title}")
            for value_row, formula_row in zip(
                sheet.iter_rows(values_only=True), formula_sheet.iter_rows(values_only=True)
            ):
                cells = []
                for value, formula in zip(value_row, formula_row):
                    if value is None and isinstance(formula, str) and formula.startswith("="):
                        value = formula
                    cells.append("" if value is None else str(value))
                line = " | ".join(cells).strip()
                if line.strip(" |"):
                    rows.append(line)
        return normalize_text("\n".join(rows)), "xlsx", None
    except Exception as exc:
        return None, "xlsx", f"{type(exc).__name__}: {exc}"


def _pptx_shape_text(shape) -> list[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return [text for child in shape.shapes for text in _pptx_shape_text(child)]
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
        return [_table_to_markdown(rows)]
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        return [shape.text_frame.text]
    return []


def _extract_pptx(body: bytes):
    """Slide text including grouped shapes, tables and speaker notes."""
    try:
        from pptx import Presentation

        parts = []
        for index, slide in enumerate(Presentation(io.BytesIO(body)).slides, start=1):
            parts.append(f"[SLIDE] {index}")
            for shape in slide.shapes:
                parts.extend(text for text in _pptx_shape_text(shape) if text.strip())
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame
                if notes is not None and notes.text.strip():
                    parts.append(f"[NOTES] {notes.text}")
        return normalize_text("\n".join(parts)), "pptx", None
    except Exception as exc:
        return None, "pptx", f"{type(exc).__name__}: {exc}"
