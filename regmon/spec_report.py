"""Analyst workbook for an offline spec comparison (python -m regmon.spec_diff ... -o report.xlsx).

The sheets follow the review habit of the team's existing compare tools, so
the output can replace them in the quarterly cycle:

* Summary       headline, both files with size and SHA-256, counts per table
                and category, declared-vs-detected mismatches, colour legend.
* Spec Changes  one row per changed rule/field/element: category, item,
                effective, severity, what changed, before/after, note, the
                regulator's own annotation, a mismatch flag, and an Impact
                drop-down (Impacted / Not impacted / Needs analysis / Out of
                scope) plus Analyst comment for sign-off.
* Cell Changes  one row per changed cell, filterable.
* Patterns      systematic edits collapsed from many rows.

Writing rules that keep the file valid: empty values are real blanks, text
that starts with "=" is stored as text (never a live formula), characters
Excel rejects are removed and long text is cut with a marker.
"""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from regmon.spec_diff import CATEGORIES, CATEGORY_LABELS

FILLS = {
    "added": "E2F0D9", "added_version": "E2F0D9", "removed": "FADBD8", "retired": "FADBD8", "severity": "FCE4D6",
    "logic": "FFF2CC", "message": "FFF2CC", "definition": "FFF2CC", "description": "FFF2CC", "effective": "FFF2CC",
    "reinstated": "E2F0D9", "metadata": "F2F2F2", "cosmetic": "DDEBF7",
}
IMPACT_CHOICES = '"Impacted,Not impacted,Needs analysis,Out of scope"'
MAX_CELL = 32000


def _safe(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return value
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    text = ILLEGAL_CHARACTERS_RE.sub("", str(value))
    if len(text) > MAX_CELL:
        text = text[: MAX_CELL - 12] + " [truncated]"
    return text


def _put(sheet, row: int, column: int, value):
    cell = sheet.cell(row=row, column=column, value=_safe(value))
    if isinstance(cell.value, str) and cell.value.startswith("="):
        cell.data_type = "s"
    return cell


def _header(sheet, row: int, names: list[str]):
    from openpyxl.styles import Font, PatternFill

    for index, name in enumerate(names, start=1):
        cell = _put(sheet, row, index, name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="3A2A1E")


def _what_changed(change: dict) -> str:
    parts = []
    severity = change.get("severity")
    if severity:
        parts.append(f"Severity {severity['before'] or '(blank)'} → {severity['after'] or '(blank)'} ({severity['direction']})")
    for column in change.get("columns") or []:
        if column["kind"] == "systematic":
            continue
        before = column["before"] or "(blank)"
        after = column["after"] or "(blank)"
        if len(before) + len(after) > 160:
            parts.append(f"{column['column']} changed" + (" (cosmetic)" if column["kind"] == "cosmetic" else ""))
        else:
            parts.append(f"{column['column']}: {before} → {after}" + (" (cosmetic)" if column["kind"] == "cosmetic" else ""))
    return "; ".join(parts)


def _before_after(change: dict) -> tuple[str, str]:
    if change.get("values"):
        text = " · ".join(f"{key}: {value}" for key, value in change["values"].items())
        return (text, "") if change["category"] == "removed" else ("", text)
    columns = [column for column in change.get("columns") or [] if column["kind"] != "systematic"]
    main = next((column for column in columns if column["role"] in {"logic", "message", "description"}), columns[0] if columns else None)
    return (main["before"], main["after"]) if main else ("", "")


def write_workbook(outcome: dict, path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    result = outcome.get("result") or {"tables": [], "counts": {}, "headline": "No rule, field or schema table changed."}
    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    _put(summary, 1, 1, "Specification comparison").font = Font(bold=True, size=14)
    _put(summary, 2, 1, "Generated " + datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC") + " by ChangeSentinel (regmon.spec_diff)")
    _put(summary, 4, 1, "In short").font = Font(bold=True)
    headline = _put(summary, 5, 1, result.get("headline") or "No rule, field or schema table changed.")
    headline.alignment = Alignment(wrap_text=True, vertical="top")
    summary.merge_cells("A5:H5")
    summary.row_dimensions[5].height = 60
    row = 7
    _header(summary, row, ["File", "Role", "Bytes", "SHA-256"])
    for item in outcome.get("files", []):
        row += 1
        for column, value in enumerate([item["name"], item["role"], item["bytes"], item["sha256"]], start=1):
            _put(summary, row, column, value)
    row += 2
    _header(summary, row, ["Table", "Kind", "Aligned by", "Rows before", "Rows after"] + [CATEGORY_LABELS["rules"][name] for name in CATEGORIES] + ["Mismatches"])
    for table in result.get("tables", []):
        row += 1
        values = [table["name"], table["kind"], ", ".join(table.get("key") or []), table.get("rows_before"), table.get("rows_after")]
        values += [table["counts"].get(name) for name in CATEGORIES]
        reconciliation = table.get("reconciliation") or {}
        values.append((reconciliation.get("changed_not_declared_count", 0) + reconciliation.get("declared_not_changed_count", 0)) or None)
        for column, value in enumerate(values, start=1):
            _put(summary, row, column, value)
    row += 2
    _put(summary, row, 1, "Colours").font = Font(bold=True)
    for label, colour in (("Added / new version / reinstated", "E2F0D9"), ("Removed / retired", "FADBD8"), ("Severity changed", "FCE4D6"),
                          ("Substantive change", "FFF2CC"), ("Metadata only", "F2F2F2"), ("Cosmetic only", "DDEBF7")):
        row += 1
        _put(summary, row, 1, label).fill = PatternFill("solid", fgColor=colour)
    row += 2
    _put(summary, row, 1, "How to read: Spec Changes has one row per changed item with an Impact drop-down for sign-off. "
                          "“Declared” is the regulator's own annotation in the file; “Mismatch” flags a row whose annotation and actual change disagree.")
    summary.column_dimensions["A"].width = 46
    for letter in "BCDEFGH":
        summary.column_dimensions[letter].width = 16

    changes_sheet = book.create_sheet("Spec Changes")
    columns = ["Category", "Table", "Item", "Field", "Effective", "Quarter", "Severity", "What changed", "Before", "After", "Note",
               "Declared in file", "Mismatch", "Also", "Impact", "Analyst comment"]
    _header(changes_sheet, 1, columns)
    cells_sheet = book.create_sheet("Cell Changes")
    _header(cells_sheet, 1, ["Table", "Item", "Column", "Role", "Before", "After", "Kind"])
    patterns_sheet = book.create_sheet("Patterns")
    _header(patterns_sheet, 1, ["Table", "Column", "From", "To", "Rows"])
    row = cell_row = pattern_row = 1
    for table in result.get("tables", []):
        labels = CATEGORY_LABELS.get(table["kind"], CATEGORY_LABELS["fields"])
        for pattern in table.get("patterns") or []:
            pattern_row += 1
            for column, value in enumerate([table["name"], pattern["column"], pattern["from"], pattern["to"], pattern["rows"]], start=1):
                _put(patterns_sheet, pattern_row, column, value)
        if table.get("table_added") or table.get("table_removed"):
            row += 1
            _put(changes_sheet, row, 1, "New table" if table.get("table_added") else "Table removed")
            _put(changes_sheet, row, 2, table["name"])
            continue
        for change in table.get("changes", []):
            row += 1
            before, after = _before_after(change)
            severity = change.get("severity")
            values = [labels.get(change["category"], change["category"]), table["name"], change["label"], change.get("field"),
                      change.get("effective"), change.get("quarter"),
                      f"{severity['before']} → {severity['after']} ({severity['direction']})" if severity else None,
                      _what_changed(change), before, after, change.get("note"), change.get("declared"), change.get("mismatch"),
                      ", ".join(labels.get(name, name) for name in change.get("also") or []), None, None]
            fill = PatternFill("solid", fgColor=FILLS.get(change["category"], "FFFFFF"))
            for column, value in enumerate(values, start=1):
                cell = _put(changes_sheet, row, column, value)
                if column == 1:
                    cell.fill = fill
                if column in (8, 9, 10, 11):
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            if change.get("mismatch"):
                changes_sheet.cell(row=row, column=13).font = Font(bold=True, color="B5480E")
            for column in change.get("columns") or []:
                cell_row += 1
                for index, value in enumerate([table["name"], change["label"], column["column"], column["role"], column["before"],
                                               column["after"], column["kind"]], start=1):
                    _put(cells_sheet, cell_row, index, value)
    if row > 1:
        validation = DataValidation(type="list", formula1=IMPACT_CHOICES, allow_blank=True)
        validation.add(f"O2:O{row}")
        changes_sheet.add_data_validation(validation)
    for sheet, widths in ((changes_sheet, [24, 26, 24, 18, 26, 10, 22, 60, 50, 50, 36, 30, 36, 20, 16, 40]),
                          (cells_sheet, [26, 24, 22, 12, 60, 60, 12]), (patterns_sheet, [26, 22, 30, 30, 8])):
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
        sheet.freeze_panes = "A2"
        if sheet.max_row > 1:
            sheet.auto_filter.ref = sheet.dimensions
    path = Path(path)
    try:
        book.save(path)
    except PermissionError:
        # The report is open in Excel: keep going with a time-stamped copy.
        path = path.with_name(f"{path.stem}_{datetime.now().strftime('%Y%m%d-%H%M%S')}{path.suffix}")
        book.save(path)
    return path
